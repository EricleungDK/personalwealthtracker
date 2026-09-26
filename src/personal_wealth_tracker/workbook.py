from __future__ import annotations

import shutil
from collections import defaultdict
from copy import copy
from dataclasses import dataclass, replace
from datetime import datetime
from decimal import Decimal
from pathlib import Path

from .config import AppConfig
from .models import (
    CategorizedTransaction,
    TrackerUpdate,
    WorkbookPlan,
    WorkbookStructureChange,
)
from .row_insertion import insert_row_with_formulas, simple_sum_range

DERIVED_WORKBOOK_ROWS = frozenset(
    {
        "Income (net)",
        "Labour market contribution",
        "Taxes",
        "Cashflow",
        "Recurring payments",
        "Living expenses",
        "Services",
        "Insurance",
        "Investments",
        "Assets",
        "Total net worth",
    }
)
NET_SALARY_ROW = "Full-time job (net)"
EXPENSE_CLAIMS_ROW = "Expense claims"
MASTERCARD_REFUND_ROW = "Mastercard refund"
INCOME_LIKE_ROWS = frozenset({NET_SALARY_ROW, EXPENSE_CLAIMS_ROW, MASTERCARD_REFUND_ROW})
MONTHS = ("Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec")


@dataclass(frozen=True)
class WorkbookCategoryOption:
    row_number: int
    category: str
    learnable: bool
    status: str
    category_type: str
    allows_new_children: bool


def plan_updates(
    tracker_path: Path,
    categorized: list[CategorizedTransaction],
    year: int,
    month: str,
    config: AppConfig,
) -> list[TrackerUpdate]:
    return plan_workbook_changes(tracker_path, categorized, year, month, config).updates


def plan_workbook_changes(
    tracker_path: Path,
    categorized: list[CategorizedTransaction],
    year: int,
    month: str,
    config: AppConfig,
) -> WorkbookPlan:
    workbook, sheet = _load_sheet(tracker_path, config.sheet_name)
    target_column = _find_month_column(sheet, year, month, config)
    structure_changes: list[WorkbookStructureChange] = []
    if target_column is None:
        structure_change = _plan_missing_period(sheet, year, month, config)
        structure_changes.append(structure_change)
        target_column = _first_target_column(structure_change)
        if structure_change.write_action == "write":
            _apply_structure_change(sheet, structure_change, config)

    grouped: dict[str, list[CategorizedTransaction]] = defaultdict(list)
    for item in categorized:
        if not item.suggested_category:
            continue
        grouped[item.suggested_category].append(item)

    # Structure changes are applied to the in-memory planning workbook in order, exactly as
    # commit will replay them, so later placements and all update rows see the final layout.
    category_rows = _category_rows(sheet, config.category_column)
    for category in sorted(grouped):
        if category in category_rows:
            continue
        if not config.category_registry.is_leaf_category(category):
            continue
        structure_change = _plan_missing_leaf_category(
            sheet,
            category_rows,
            category,
            year,
            month,
            target_column,
            config,
        )
        if structure_change.write_action == "write":
            failure = _apply_leaf_category_insertion(sheet, structure_change, config)
            if failure is not None:
                structure_change = replace(structure_change, write_action="review", reason=failure)
                workbook.close()
                workbook, sheet = _load_sheet_with_changes(tracker_path, structure_changes, config)
        structure_changes.append(structure_change)
        category_rows = _category_rows(sheet, config.category_column)

    updates: list[TrackerUpdate] = []
    for category, items in sorted(grouped.items()):
        amount = _category_amount(category, items)
        row = category_rows.get(category)
        cell = sheet.cell(row=row, column=target_column) if row and target_column else None
        existing_value = cell.value if cell else None
        action, reason = _write_decision(category, row, target_column, existing_value, items, config)
        updates.append(
            TrackerUpdate(
                year=year,
                month=month,
                category=category,
                amount=amount,
                source_transactions=tuple(item.transaction.transaction_id for item in items),
                target_row=row,
                target_column=target_column,
                target_cell=cell.coordinate if cell else None,
                existing_value=existing_value,
                write_action=action,
                reason=reason,
            )
        )

    workbook.close()
    return WorkbookPlan(updates=updates, structure_changes=structure_changes)


def _load_sheet_with_changes(
    tracker_path: Path, structure_changes: list[WorkbookStructureChange], config: AppConfig
):
    workbook, sheet = _load_sheet(tracker_path, config.sheet_name)
    try:
        for change in structure_changes:
            if change.write_action == "write":
                _apply_structure_change(sheet, change, config)
    except ValueError:
        workbook.close()
        raise
    return workbook, sheet


def rows_in_review(
    categorized: list[CategorizedTransaction], updates: list[TrackerUpdate]
) -> list[CategorizedTransaction]:
    blocked_ids = {
        transaction_id
        for update in updates
        if update.write_action == "review"
        for transaction_id in update.source_transactions
    }
    return [
        item
        for item in categorized
        if item.review_required or item.transaction.transaction_id in blocked_ids
    ]


def workbook_category_options(
    tracker_path: Path,
    year: int,
    month: str,
    config: AppConfig,
) -> list[WorkbookCategoryOption]:
    workbook, sheet = _load_sheet(tracker_path, config.sheet_name)
    target_column = _find_month_column(sheet, year, month, config)
    options: list[WorkbookCategoryOption] = []
    for row in range(1, sheet.max_row + 1):
        category = _row_category(sheet, row, config)
        if category is None:
            continue
        category_type = _category_option_type(category, config)
        learnable = (
            category_type == "leaf"
            and category not in DERIVED_WORKBOOK_ROWS
            and category not in config.fixed_rows
        )
        allows_new_children = config.category_registry.allows_new_leaf_children(category)
        status = _category_option_status(
            sheet,
            row,
            target_column,
            category,
            config,
            category_type,
        )
        options.append(
            WorkbookCategoryOption(
                row_number=row,
                category=category,
                learnable=learnable,
                status=status,
                category_type=category_type,
                allows_new_children=allows_new_children,
            )
        )
    workbook.close()
    return options


def _category_option_type(category: str, config: AppConfig) -> str:
    registered_type = config.category_registry.category_type_by_label.get(category)
    if registered_type:
        return registered_type
    if category in DERIVED_WORKBOOK_ROWS:
        return "derived"
    if config.category_registry.category_type_by_label:
        return "unregistered"
    return "leaf"


def _category_option_status(
    sheet,
    row: int,
    target_column: int | None,
    category: str,
    config: AppConfig,
    category_type: str,
) -> str:
    if category_type == "parent":
        return "parent/section row; not a manual transaction category"
    if category_type == "derived" or category in DERIVED_WORKBOOK_ROWS:
        return "derived/formula-owned; not learnable or writable"
    if category_type == "unregistered":
        return "not in category registry; not learnable or writable"
    if category in config.fixed_rows:
        return "fixed/protected; not learnable or writable"
    if target_column is None:
        return "target month missing; write safety unresolved"

    existing_value = sheet.cell(row=row, column=target_column).value
    if isinstance(existing_value, str) and existing_value.startswith("="):
        return "current-month formula; not writable"
    if existing_value not in (None, ""):
        return "current-month manual value; review before writing"
    return "write eligible"


def _category_amount(category: str, items: list[CategorizedTransaction]) -> Decimal:
    signed_total = sum((item.transaction.amount for item in items), Decimal("0"))
    if category in INCOME_LIKE_ROWS:
        return signed_total
    return -signed_total


def commit_updates(
    tracker_path: Path,
    updates: list[TrackerUpdate],
    config: AppConfig,
    output_dir: Path,
    structure_changes: list[WorkbookStructureChange] | None = None,
) -> Path:
    workbook, sheet = _load_sheet_with_changes(tracker_path, structure_changes or [], config)
    output_dir.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_path = output_dir / f"{tracker_path.stem}_auto_{timestamp}{tracker_path.suffix}"

    for update in updates:
        if update.write_action != "write" or update.target_row is None or update.target_column is None:
            continue
        sheet.cell(row=update.target_row, column=update.target_column).value = float(update.amount)

    workbook.save(output_path)
    workbook.close()
    return output_path


def _apply_structure_change(sheet, change: WorkbookStructureChange, config: AppConfig) -> None:
    if change.change_type == "create_period":
        if change.source_range is None or change.target_range is None:
            return
        source_column = _column_index(change.source_range.split(":", 1)[0])
        target_column = _column_index(change.target_range.split(":", 1)[0])
        if target_column <= sheet.max_column:
            sheet.insert_cols(target_column)
        _copy_period_column(
            sheet,
            source_column=source_column,
            target_column=target_column,
            target_year=change.target_year,
            target_month=change.target_months[0],
            config=config,
        )
        _extend_merged_year_header(
            sheet,
            source_column=source_column,
            target_column=target_column,
            target_year=change.target_year,
            config=config,
        )
    elif change.change_type == "create_year":
        if change.source_range is None or change.target_range is None:
            return
        source_start = _column_index(change.source_range.split(":", 1)[0])
        target_start = _column_index(change.target_range.split(":", 1)[0])
        if target_start <= sheet.max_column:
            sheet.insert_cols(target_start, len(change.target_months))
        for offset, target_month in enumerate(change.target_months):
            _copy_period_column(
                sheet,
                source_column=source_start + offset,
                target_column=target_start + offset,
                target_year=change.target_year,
                target_month=target_month,
                config=config,
            )
        _copy_merged_year_header(
            sheet,
            source_start=source_start,
            source_end=source_start + len(change.target_months) - 1,
            target_start=target_start,
            target_end=target_start + len(change.target_months) - 1,
            target_year=change.target_year,
            config=config,
        )
    elif change.change_type == "insert_leaf_category":
        failure = _apply_leaf_category_insertion(sheet, change, config)
        if failure is not None:
            raise ValueError(failure)


def _apply_leaf_category_insertion(
    sheet, change: WorkbookStructureChange, config: AppConfig
) -> str | None:
    """Insert the planned leaf row; return a failure reason if the safety check fails."""
    if (
        change.source_range is None
        or change.target_range is None
        or change.leaf_category is None
        or change.parent_category is None
    ):
        return f"Leaf insertion for {change.leaf_category} is incomplete."

    source_row = _row_index(change.source_range.split(":", 1)[0])
    target_row = _row_index(change.target_range.split(":", 1)[0])
    parent_row = _category_rows(sheet, config.category_column).get(change.parent_category)
    if parent_row is None:
        return f"Parent row for {change.leaf_category} was not found."

    failure = insert_row_with_formulas(sheet.parent, sheet.title, target_row, parent_row)
    if failure is not None:
        return failure
    _copy_row_format(sheet, source_row, target_row)
    sheet.cell(row=target_row, column=config.category_column).value = change.leaf_category
    return None


def _copy_row_format(sheet, source_row: int, target_row: int) -> None:
    sheet.row_dimensions[target_row].height = sheet.row_dimensions[source_row].height
    sheet.row_dimensions[target_row].hidden = sheet.row_dimensions[source_row].hidden
    sheet.row_dimensions[target_row].outline_level = sheet.row_dimensions[source_row].outline_level
    for column in range(1, sheet.max_column + 1):
        source_cell = sheet.cell(row=source_row, column=column)
        target_cell = sheet.cell(row=target_row, column=column)
        if source_cell.has_style:
            target_cell._style = copy(source_cell._style)
        if source_cell.number_format:
            target_cell.number_format = source_cell.number_format
        if source_cell.font:
            target_cell.font = copy(source_cell.font)
        if source_cell.fill:
            target_cell.fill = copy(source_cell.fill)
        if source_cell.border:
            target_cell.border = copy(source_cell.border)
        if source_cell.alignment:
            target_cell.alignment = copy(source_cell.alignment)
        if source_cell.protection:
            target_cell.protection = copy(source_cell.protection)


def _copy_period_column(
    sheet,
    source_column: int,
    target_column: int,
    target_year: int,
    target_month: str,
    config: AppConfig,
) -> None:
    source_letter = _column_letter(source_column)
    target_letter = _column_letter(target_column)
    sheet.column_dimensions[target_letter].width = sheet.column_dimensions[source_letter].width
    sheet.column_dimensions[target_letter].hidden = sheet.column_dimensions[source_letter].hidden

    for row in range(1, sheet.max_row + 1):
        source_cell = sheet.cell(row=row, column=source_column)
        target_cell = sheet.cell(row=row, column=target_column)
        if source_cell.has_style:
            target_cell._style = copy(source_cell._style)
        if source_cell.number_format:
            target_cell.number_format = source_cell.number_format
        if source_cell.font:
            target_cell.font = copy(source_cell.font)
        if source_cell.fill:
            target_cell.fill = copy(source_cell.fill)
        if source_cell.border:
            target_cell.border = copy(source_cell.border)
        if source_cell.alignment:
            target_cell.alignment = copy(source_cell.alignment)
        if source_cell.protection:
            target_cell.protection = copy(source_cell.protection)

        if row == config.year_header_row:
            target_cell.value = target_year
        elif row == config.month_header_row:
            target_cell.value = target_month
        elif isinstance(source_cell.value, str) and source_cell.value.startswith("="):
            target_cell.value = _translate_formula(
                source_cell.value,
                source_coordinate=source_cell.coordinate,
                target_coordinate=target_cell.coordinate,
            )
        elif _row_category(sheet, row, config) in config.carry_forward_rows:
            target_cell.value = source_cell.value
        else:
            target_cell.value = None


def _copy_merged_year_header(
    sheet,
    source_start: int,
    source_end: int,
    target_start: int,
    target_end: int,
    target_year: int,
    config: AppConfig,
) -> None:
    source_range = (
        f"{_column_letter(source_start)}{config.year_header_row}:"
        f"{_column_letter(source_end)}{config.year_header_row}"
    )
    for merged_range in sheet.merged_cells.ranges:
        if str(merged_range) == source_range:
            target_range = (
                f"{_column_letter(target_start)}{config.year_header_row}:"
                f"{_column_letter(target_end)}{config.year_header_row}"
            )
            sheet.merge_cells(target_range)
            sheet.cell(row=config.year_header_row, column=target_start).value = target_year
            return


def _extend_merged_year_header(
    sheet,
    source_column: int,
    target_column: int,
    target_year: int,
    config: AppConfig,
) -> None:
    for merged_range in tuple(sheet.merged_cells.ranges):
        if (
            merged_range.min_row == config.year_header_row
            and merged_range.max_row == config.year_header_row
            and merged_range.max_col == source_column
            and target_column == source_column + 1
        ):
            start_column = merged_range.min_col
            sheet.unmerge_cells(str(merged_range))
            target_range = (
                f"{_column_letter(start_column)}{config.year_header_row}:"
                f"{_column_letter(target_column)}{config.year_header_row}"
            )
            sheet.merge_cells(target_range)
            sheet.cell(row=config.year_header_row, column=start_column).value = target_year
            return


def _translate_formula(formula: str, source_coordinate: str, target_coordinate: str) -> str:
    try:
        from openpyxl.formula.translate import Translator
    except ImportError as exc:
        raise RuntimeError("openpyxl is required for workbook access. Install with `uv sync`.") from exc

    return Translator(formula, origin=source_coordinate).translate_formula(target_coordinate)


def _row_category(sheet, row: int, config: AppConfig) -> str | None:
    value = sheet.cell(row=row, column=config.category_column).value
    if isinstance(value, str) and value.strip():
        return value.strip()
    return None


def create_backup(tracker_path: Path, backup_dir: Path) -> Path:
    backup_dir.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_path = backup_dir / f"{tracker_path.stem}_backup_{timestamp}{tracker_path.suffix}"
    shutil.copy2(tracker_path, backup_path)
    return backup_path


def _load_sheet(tracker_path: Path, sheet_name: str):
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError("openpyxl is required for workbook access. Install with `uv sync`.") from exc

    workbook = openpyxl.load_workbook(tracker_path)
    if sheet_name not in workbook.sheetnames:
        workbook.close()
        raise ValueError(f"Workbook does not contain sheet {sheet_name!r}.")
    return workbook, workbook[sheet_name]


def _category_rows(sheet, category_column: int) -> dict[str, int]:
    rows: dict[str, int] = {}
    for row in range(1, sheet.max_row + 1):
        value = sheet.cell(row=row, column=category_column).value
        if isinstance(value, str) and value.strip():
            rows[value.strip()] = row
    return rows


def _target_month_column(sheet, year: int, month: str, config: AppConfig) -> int:
    target_column = _find_month_column(sheet, year, month, config)
    if target_column is None:
        raise ValueError(f"Could not find target column for {month} {year}.")
    return target_column


def _find_month_column(sheet, year: int, month: str, config: AppConfig) -> int | None:
    current_year: int | None = None
    for column in range(1, sheet.max_column + 1):
        year_value = sheet.cell(row=config.year_header_row, column=column).value
        if year_value not in (None, ""):
            try:
                current_year = int(year_value)
            except (TypeError, ValueError):
                current_year = None

        month_value = sheet.cell(row=config.month_header_row, column=column).value
        if current_year == year and str(month_value).strip().lower() == month.lower():
            return column

    return None


def _plan_missing_leaf_category(
    sheet,
    category_rows: dict[str, int],
    leaf_category: str,
    year: int,
    month: str,
    target_column: int | None,
    config: AppConfig,
) -> WorkbookStructureChange:
    parent_category = _parent_for_leaf(leaf_category, config)
    parent_row = category_rows.get(parent_category) if parent_category else None
    if parent_category is None or parent_row is None:
        return WorkbookStructureChange(
            change_type="insert_leaf_category",
            target_year=year,
            target_months=(month,),
            source_range=None,
            target_range=None,
            write_action="review",
            reason=f"Parent row for missing leaf category {leaf_category} was not found.",
            parent_category=parent_category,
            leaf_category=leaf_category,
        )

    def review(reason: str) -> WorkbookStructureChange:
        return WorkbookStructureChange(
            change_type="insert_leaf_category",
            target_year=year,
            target_months=(month,),
            source_range=None,
            target_range=None,
            write_action="review",
            reason=reason,
            parent_category=parent_category,
            leaf_category=leaf_category,
        )

    if target_column is None:
        return review("Target month column not found.")

    boundary_row = _next_category_boundary_row(sheet, parent_row, config)
    parent_formula = sheet.cell(row=parent_row, column=target_column).value
    if isinstance(parent_formula, str) and parent_formula.startswith("="):
        range_info = simple_sum_range(parent_formula)
        if range_info is None:
            return review(
                "Parent formula is not a simple SUM range; manual workbook adjustment required."
            )
        start_column, start_row, end_column, end_row = range_info
        target_column_letter = _column_letter(target_column)
        if (
            start_column != target_column_letter
            or end_column != target_column_letter
            or start_row <= parent_row
            or start_row > end_row
        ):
            return review(f"Parent formula SUM range does not match the {parent_category} section.")
        if boundary_row is not None and end_row >= boundary_row:
            return review(
                f"Parent formula SUM range for {parent_category} extends past the next "
                "parent/section boundary."
            )
        target_row = end_row + 1
    elif boundary_row is not None:
        target_row = boundary_row
    else:
        target_row = _last_labeled_row(sheet, parent_row, config) + 1

    source_row = target_row - 1
    if source_row == parent_row:
        return review(f"No existing leaf row in the {parent_category} section to copy formatting from.")

    return WorkbookStructureChange(
        change_type="insert_leaf_category",
        target_year=year,
        target_months=(month,),
        source_range=f"{source_row}:{source_row}",
        target_range=f"{target_row}:{target_row}",
        write_action="write",
        reason=(
            f"Insert missing leaf category {leaf_category} under {parent_category} at the end "
            f"of its section, row {target_row}."
        ),
        parent_category=parent_category,
        leaf_category=leaf_category,
    )


def _parent_for_leaf(leaf_category: str, config: AppConfig) -> str | None:
    for parent_category, children in config.category_registry.children_by_parent.items():
        if leaf_category in children:
            return parent_category
    return None


def _next_category_boundary_row(sheet, parent_row: int, config: AppConfig) -> int | None:
    for row in range(parent_row + 1, sheet.max_row + 1):
        category = _row_category(sheet, row, config)
        if category is None:
            continue
        if config.category_registry.is_parent_category(category) or category in DERIVED_WORKBOOK_ROWS:
            return row
    return None


def _last_labeled_row(sheet, parent_row: int, config: AppConfig) -> int:
    last_row = parent_row
    for row in range(parent_row + 1, sheet.max_row + 1):
        if _row_category(sheet, row, config) is not None:
            last_row = row
    return last_row


def _plan_missing_period(sheet, year: int, month: str, config: AppConfig) -> WorkbookStructureChange:
    if month not in MONTHS:
        return WorkbookStructureChange(
            change_type="review",
            target_year=year,
            target_months=(month,),
            source_range=None,
            target_range=None,
            write_action="review",
            reason=f"Unknown target month {month!r}.",
        )

    if month == "Jan":
        prior_year_columns = _unambiguous_year_block(sheet, year - 1, config)
        if prior_year_columns is None:
            return WorkbookStructureChange(
                change_type="review",
                target_year=year,
                target_months=MONTHS,
                source_range=None,
                target_range=None,
                write_action="review",
                reason=(
                    "Target year is missing and the prior year does not contain an "
                    "unambiguous Jan-Dec period block."
                ),
            )
        source_start = prior_year_columns[0]
        source_end = prior_year_columns[-1]
        target_start = source_end + 1
        target_end = target_start + len(MONTHS) - 1
        return WorkbookStructureChange(
            change_type="create_year",
            target_year=year,
            target_months=MONTHS,
            source_range=f"{_column_letter(source_start)}:{_column_letter(source_end)}",
            target_range=f"{_column_letter(target_start)}:{_column_letter(target_end)}",
            write_action="write",
            reason=(
                f"Create missing {year} year block from "
                f"{_column_letter(source_start)}:{_column_letter(source_end)}."
            ),
        )

    previous_month = MONTHS[MONTHS.index(month) - 1] if month != "Jan" else None
    source_column = _find_month_column(sheet, year, previous_month, config)
    if source_column is None:
        return WorkbookStructureChange(
            change_type="review",
            target_year=year,
            target_months=(month,),
            source_range=None,
            target_range=None,
            write_action="review",
            reason="Target month is missing and the immediately previous period was not found.",
        )

    target_column = source_column + 1
    return WorkbookStructureChange(
        change_type="create_period",
        target_year=year,
        target_months=(month,),
        source_range=f"{_column_letter(source_column)}:{_column_letter(source_column)}",
        target_range=f"{_column_letter(target_column)}:{_column_letter(target_column)}",
        write_action="write",
        reason=(
            f"Create missing {month} {year} period from "
            f"{_column_letter(source_column)}:{_column_letter(source_column)}."
        ),
    )


def _unambiguous_year_block(sheet, year: int, config: AppConfig) -> list[int] | None:
    columns: list[int] = []
    current_year: int | None = None
    for column in range(1, sheet.max_column + 1):
        year_value = sheet.cell(row=config.year_header_row, column=column).value
        if year_value not in (None, ""):
            try:
                current_year = int(year_value)
            except (TypeError, ValueError):
                current_year = None

        if current_year == year:
            columns.append(column)

    if len(columns) != len(MONTHS):
        return None
    if columns != list(range(columns[0], columns[0] + len(MONTHS))):
        return None
    month_values = tuple(
        str(sheet.cell(row=config.month_header_row, column=column).value).strip()
        for column in columns
    )
    if month_values != MONTHS:
        return None
    return columns


def _first_target_column(structure_change: WorkbookStructureChange) -> int | None:
    if structure_change.write_action != "write" or structure_change.target_range is None:
        return None
    first_column = structure_change.target_range.split(":", 1)[0]
    return _column_index(first_column)


def _row_index(row: str) -> int:
    return int(row)


def _column_letter(column: int) -> str:
    try:
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError("openpyxl is required for workbook access. Install with `uv sync`.") from exc

    return get_column_letter(column)


def _column_index(column: str) -> int:
    try:
        from openpyxl.utils import column_index_from_string
    except ImportError as exc:
        raise RuntimeError("openpyxl is required for workbook access. Install with `uv sync`.") from exc

    return column_index_from_string(column)


def _write_decision(
    category: str,
    row: int | None,
    column: int | None,
    existing_value: object,
    items: list[CategorizedTransaction],
    config: AppConfig,
) -> tuple[str, str]:
    if row is None:
        return "review", "Target category row not found."
    if column is None:
        return "review", "Target month column not found."
    if category in DERIVED_WORKBOOK_ROWS:
        return "skip", "Derived workbook row is formula-owned and not writable."
    if category in config.fixed_rows and not config.overwrite_fixed_rows:
        return "skip", "Fixed row is protected by config."
    if isinstance(existing_value, str) and existing_value.startswith("="):
        return "skip", "Target cell contains a formula."
    if existing_value not in (None, ""):
        return "review", "Target cell already contains a manual value."
    if category == NET_SALARY_ROW and len(items) > 1:
        return "review", "Multiple salary deposits matched; review before writing."
    return "write", "Eligible for commit mode write."
