from __future__ import annotations

import csv
import json
import re
import warnings
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path

from .models import Transaction
from .utils import TRANSACTION_ID_SCHEME_VERSION, normalize_month, normalize_text


MEMORY_FILE_NAME = "category_memory.json"


@dataclass(frozen=True)
class CategoryMemoryImportResult:
    imported_count: int
    skipped_unconfirmed_count: int
    memory_path: Path
    skipped_unlearned_count: int = 0
    skipped_non_learnable_count: int = 0


@dataclass(frozen=True)
class RecurringMatchHint:
    amount: Decimal
    amount_tolerance: Decimal
    day_min: int | None
    day_max: int | None


@dataclass(frozen=True)
class CategoryMemoryMapping:
    merchant_identity: str
    category: str
    source_transaction_ids: tuple[str, ...]
    recurring_hint: RecurringMatchHint | None = None


@dataclass(frozen=True)
class CategoryMemory:
    mappings: tuple[CategoryMemoryMapping, ...]


def import_reviewed_decisions(decisions_path: Path, memory_dir: Path) -> CategoryMemoryImportResult:
    if decisions_path.suffix.lower() == ".xlsx":
        return _import_review_workbook_decisions(decisions_path, memory_dir)

    mappings = [_mapping_payload(mapping) for mapping in load_category_memory(memory_dir).mappings]
    skipped_unconfirmed = 0
    imported_count = 0
    with decisions_path.open("r", encoding="utf-8", newline="") as handle:
        reader = csv.DictReader(handle)
        for row_number, row in enumerate(reader, start=2):
            if not _is_confirmed(row.get("confirmed", "")):
                skipped_unconfirmed += 1
                continue
            transaction_id = _required(row, "transaction_id", row_number)
            description = _required(row, "description", row_number)
            category = _required(row, "confirmed_category", row_number)
            _upsert_mapping(
                mappings,
                {
                    "merchant_identity": normalize_merchant_identity(description),
                    "category": category,
                    "source_transaction_ids": [transaction_id],
                    **_recurring_hint_payload(row, row_number),
                },
            )
            imported_count += 1

    memory_dir.mkdir(parents=True, exist_ok=True)
    memory_path = memory_dir / MEMORY_FILE_NAME
    memory_path.write_text(
        json.dumps({"mappings": mappings}, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    return CategoryMemoryImportResult(
        imported_count=imported_count,
        skipped_unconfirmed_count=skipped_unconfirmed,
        memory_path=memory_path,
    )


def load_category_memory(memory_dir: Path) -> CategoryMemory:
    memory_path = memory_dir / MEMORY_FILE_NAME
    if not memory_path.exists():
        return CategoryMemory(mappings=())

    payload = json.loads(memory_path.read_text(encoding="utf-8"))
    return CategoryMemory(
        mappings=tuple(
            CategoryMemoryMapping(
                merchant_identity=str(item["merchant_identity"]),
                category=str(item["category"]),
                source_transaction_ids=tuple(
                    str(value) for value in item["source_transaction_ids"]
                ),
                recurring_hint=_load_recurring_hint(item.get("recurring_hint")),
            )
            for item in payload.get("mappings", [])
        )
    )


def match_category_memory(
    transaction: Transaction, memory: CategoryMemory
) -> CategoryMemoryMapping | None:
    merchant_identity = normalize_merchant_identity(transaction.description)
    for mapping in memory.mappings:
        if mapping.merchant_identity == merchant_identity:
            if mapping.recurring_hint and not _matches_recurring_hint(
                transaction, mapping.recurring_hint
            ):
                continue
            return mapping
    return None


def normalize_merchant_identity(description: str) -> str:
    normalized = normalize_text(description)
    without_numbers = re.sub(r"\b\d+\b", " ", normalized)
    return re.sub(r"\s+", " ", without_numbers).strip()


def _is_confirmed(value: str) -> bool:
    return normalize_text(value) in {"1", "TRUE", "YES", "Y"}


def _import_review_workbook_decisions(
    decisions_path: Path,
    memory_dir: Path,
) -> CategoryMemoryImportResult:
    mappings = [_mapping_payload(mapping) for mapping in load_category_memory(memory_dir).mappings]
    imported_count = 0
    skipped_unlearned_count = 0
    skipped_non_learnable_count = 0

    workbook = _load_review_workbook_without_extension_warning(decisions_path)
    try:
        if "Review Required" not in workbook.sheetnames:
            raise ValueError("Reviewed workbook is missing 'Review Required'.")
        if "Category Options" not in workbook.sheetnames:
            raise ValueError("Reviewed workbook is missing 'Category Options'.")
        if "Run Metadata" not in workbook.sheetnames:
            raise ValueError("Reviewed workbook is missing 'Run Metadata'.")

        _validate_review_workbook_metadata(workbook["Run Metadata"])
        learnable_categories = _learnable_categories(workbook["Category Options"])
        for row_number, row in _worksheet_dicts(workbook["Review Required"]):
            manual_category = (row.get("manual_category") or "").strip()
            if not manual_category:
                skipped_unlearned_count += 1
                continue
            if not _is_confirmed(row.get("learn_to_memory", "")):
                skipped_unlearned_count += 1
                continue
            if manual_category not in learnable_categories:
                skipped_non_learnable_count += 1
                continue

            transaction_id = _required(row, "transaction_id", row_number)
            description = _required(row, "description", row_number)
            _upsert_mapping(
                mappings,
                {
                    "merchant_identity": normalize_merchant_identity(description),
                    "category": manual_category,
                    "source_transaction_ids": [transaction_id],
                    **_recurring_hint_payload(row, row_number),
                },
            )
            imported_count += 1
    finally:
        workbook.close()

    memory_dir.mkdir(parents=True, exist_ok=True)
    memory_path = memory_dir / MEMORY_FILE_NAME
    memory_path.write_text(
        json.dumps({"mappings": mappings}, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    return CategoryMemoryImportResult(
        imported_count=imported_count,
        skipped_unconfirmed_count=0,
        memory_path=memory_path,
        skipped_unlearned_count=skipped_unlearned_count,
        skipped_non_learnable_count=skipped_non_learnable_count,
    )


def _load_review_workbook_without_extension_warning(path: Path):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl is required for reviewed workbook import.") from exc

    with warnings.catch_warnings():
        warnings.filterwarnings(
            "ignore",
            message="Data Validation extension is not supported and will be removed",
            category=UserWarning,
        )
        return load_workbook(path, data_only=True)


def _validate_review_workbook_metadata(sheet) -> None:
    metadata = {
        str(sheet.cell(row=row, column=1).value): sheet.cell(row=row, column=2).value
        for row in range(1, sheet.max_row + 1)
        if sheet.cell(row=row, column=1).value not in (None, "")
    }

    try:
        reporting_year = int(metadata.get("reporting_year", 0))
    except (TypeError, ValueError) as exc:
        raise ValueError("Reviewed workbook has invalid reporting_year metadata.") from exc
    if reporting_year <= 0:
        raise ValueError("Reviewed workbook is missing reporting_year metadata.")

    try:
        normalize_month(str(metadata.get("reporting_month", "")))
    except ValueError as exc:
        raise ValueError("Reviewed workbook has invalid reporting_month metadata.") from exc

    scheme = str(metadata.get("transaction_id_scheme", ""))
    if scheme != TRANSACTION_ID_SCHEME_VERSION:
        raise ValueError(
            "Unsupported reviewed workbook transaction ID scheme "
            f"{scheme!r}; expected {TRANSACTION_ID_SCHEME_VERSION!r}."
        )


def _learnable_categories(sheet) -> frozenset[str]:
    headers = _worksheet_headers(sheet)
    category_column = _required_column(headers, "category", sheet.title)
    learnable_column = _required_column(headers, "learnable", sheet.title)
    categories: set[str] = set()
    for row in range(2, sheet.max_row + 1):
        category = str(sheet.cell(row=row, column=category_column).value or "").strip()
        if not category:
            continue
        if _is_truthy(sheet.cell(row=row, column=learnable_column).value):
            categories.add(category)
    return frozenset(categories)


def _worksheet_dicts(sheet) -> list[tuple[int, dict[str, str]]]:
    headers = _worksheet_headers(sheet)
    return [
        (
            row,
            {
                header: str(sheet.cell(row=row, column=column).value or "").strip()
                for header, column in headers.items()
            },
        )
        for row in range(2, sheet.max_row + 1)
        if any(sheet.cell(row=row, column=column).value not in (None, "") for column in headers.values())
    ]


def _worksheet_headers(sheet) -> dict[str, int]:
    return {
        str(cell.value).strip(): index
        for index, cell in enumerate(sheet[1], start=1)
        if cell.value not in (None, "")
    }


def _required_column(headers: dict[str, int], column: str, sheet_title: str) -> int:
    if column not in headers:
        raise ValueError(f"Reviewed workbook sheet {sheet_title!r} is missing {column!r}.")
    return headers[column]


def _is_truthy(value: object) -> bool:
    if isinstance(value, bool):
        return value
    return _is_confirmed(str(value or ""))


def _required(row: dict[str, str], field: str, row_number: int) -> str:
    value = row.get(field, "").strip()
    if not value:
        raise ValueError(f"Reviewed decision row {row_number} is missing {field!r}.")
    return value


def _recurring_hint_payload(row: dict[str, str], row_number: int) -> dict[str, object]:
    if not _is_confirmed(row.get("recurring", "")):
        return {}

    amount = abs(_decimal(_required(row, "amount", row_number), "amount", row_number))
    amount_tolerance = _decimal(
        row.get("amount_tolerance", "").strip() or "0.00",
        "amount_tolerance",
        row_number,
    )
    if amount_tolerance < 0:
        raise ValueError(f"Reviewed decision row {row_number} has negative amount_tolerance.")
    day_min = _optional_day(row.get("day_min", ""), "day_min", row_number)
    day_max = _optional_day(row.get("day_max", ""), "day_max", row_number)
    if day_min is not None and day_max is not None and day_min > day_max:
        raise ValueError(
            f"Reviewed decision row {row_number} has day_min greater than day_max."
        )
    return {
        "recurring_hint": {
            "amount": str(amount),
            "amount_tolerance": str(amount_tolerance),
            "day_min": day_min,
            "day_max": day_max,
        }
    }


def _load_recurring_hint(payload: object) -> RecurringMatchHint | None:
    if not isinstance(payload, dict):
        return None
    return RecurringMatchHint(
        amount=Decimal(str(payload["amount"])),
        amount_tolerance=Decimal(str(payload["amount_tolerance"])),
        day_min=payload["day_min"],
        day_max=payload["day_max"],
    )


def _matches_recurring_hint(transaction: Transaction, hint: RecurringMatchHint) -> bool:
    if abs(abs(transaction.amount) - hint.amount) > hint.amount_tolerance:
        return False
    if hint.day_min is not None and transaction.date.day < hint.day_min:
        return False
    if hint.day_max is not None and transaction.date.day > hint.day_max:
        return False
    return True


def _decimal(value: str, field: str, row_number: int) -> Decimal:
    try:
        return Decimal(value)
    except InvalidOperation as exc:
        raise ValueError(f"Reviewed decision row {row_number} has invalid {field!r}.") from exc


def _optional_day(value: str, field: str, row_number: int) -> int | None:
    value = value.strip()
    if not value:
        return None
    try:
        day = int(value)
    except ValueError as exc:
        raise ValueError(f"Reviewed decision row {row_number} has invalid {field!r}.") from exc
    if not 1 <= day <= 31:
        raise ValueError(f"Reviewed decision row {row_number} has {field!r} outside 1-31.")
    return day


def _mapping_payload(mapping: CategoryMemoryMapping) -> dict[str, object]:
    payload: dict[str, object] = {
        "merchant_identity": mapping.merchant_identity,
        "category": mapping.category,
        "source_transaction_ids": list(mapping.source_transaction_ids),
    }
    if mapping.recurring_hint:
        payload["recurring_hint"] = {
            "amount": str(mapping.recurring_hint.amount),
            "amount_tolerance": str(mapping.recurring_hint.amount_tolerance),
            "day_min": mapping.recurring_hint.day_min,
            "day_max": mapping.recurring_hint.day_max,
        }
    return payload


def _upsert_mapping(mappings: list[dict[str, object]], new_mapping: dict[str, object]) -> None:
    for index, existing in enumerate(mappings):
        if existing.get("merchant_identity") != new_mapping.get("merchant_identity"):
            continue
        if existing.get("recurring_hint") != new_mapping.get("recurring_hint"):
            continue
        existing_ids = [str(value) for value in existing.get("source_transaction_ids", [])]
        new_ids = [str(value) for value in new_mapping.get("source_transaction_ids", [])]
        mappings[index] = {
            **new_mapping,
            "source_transaction_ids": [
                *existing_ids,
                *[value for value in new_ids if value not in existing_ids],
            ],
        }
        return
    mappings.append(new_mapping)
