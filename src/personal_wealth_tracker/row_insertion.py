"""Formula-aware single-row insertion with a fail-closed safety check.

openpyxl's ``insert_rows`` only moves cells; it leaves every formula, merged range,
conditional format, data validation, defined name, print area and row dimension
pointing at the old rows. ``insert_row_with_formulas`` moves all of them, grows the
parent SUM that ends directly above the new row, then proves the result: every
pre-existing cell must sit at its shifted position with its expected formula, and
every evaluable formula must produce the same value as before (the new row is empty).
"""

from __future__ import annotations

import math
import re

from openpyxl.formula.tokenizer import Token, Tokenizer

_SIMPLE_SUM = re.compile(r"=SUM\((\$?[A-Z]+)(\$?\d+):(\$?[A-Z]+)(\$?\d+)\)")
_POSITION_DEPENDENT_FUNCTIONS = ("INDIRECT(", "OFFSET(")
_SHEET_PREFIX = re.compile(r"^(?:(?P<sheet>'(?:[^']|'')+'|[^'!]+)!)?(?P<ref>[^!]+)$")
_REF_PART = re.compile(r"^(?P<col>\$?[A-Z]{1,3})?(?P<row_abs>\$?)(?P<row>\d+)?$")


def simple_sum_range(formula: str) -> tuple[str, int, str, int] | None:
    """Return (start column, start row, end column, end row) for ``=SUM(A1:A9)``."""
    match = _SIMPLE_SUM.fullmatch(formula)
    if match is None:
        return None
    return tuple(
        int(part.replace("$", "")) if index % 2 else part.replace("$", "")
        for index, part in enumerate(match.groups())
    )


def insert_row_with_formulas(workbook, sheet_title: str, insert_row: int, parent_row: int) -> str | None:
    """Insert one empty row above ``insert_row`` on ``sheet_title``.

    Every formula on ``parent_row`` must be a same-column ``=SUM(Xa:Xb)`` with
    ``b == insert_row - 1``; each grows to include the new row. Returns a failure reason when the insertion is unsupported
    or fails the safety check; the workbook must then be discarded.
    """
    blocker = _insertion_blocker(workbook, sheet_title) or _parent_row_blocker(
        workbook[sheet_title], parent_row, insert_row
    )
    if blocker is not None:
        return blocker

    before_cells = {
        sheet.title: {key: cell.value for key, cell in sheet._cells.items() if cell.value is not None}
        for sheet in workbook.worksheets
    }
    before_values = {sheet.title: evaluate_sheet(workbook, sheet.title) for sheet in workbook.worksheets}

    sheet = workbook[sheet_title]
    merged_ranges = [str(merged) for merged in sheet.merged_cells.ranges]
    for merged in merged_ranges:
        sheet.unmerge_cells(merged)
    sheet.insert_rows(insert_row)
    for merged in merged_ranges:
        sheet.merge_cells(_shift_sqref(merged, sheet_title, insert_row))
    _shift_row_dimensions(sheet, insert_row)
    _shift_sheet_rules(sheet, insert_row)
    _shift_defined_names(workbook, sheet_title, insert_row)
    for other_sheet in workbook.worksheets:
        for cell in other_sheet._cells.values():
            if _is_formula(cell.value):
                cell.value = shift_formula_rows(cell.value, other_sheet.title, sheet_title, insert_row)
    grown = _grow_parent_sums(sheet, parent_row, insert_row)

    return _verify_insertion(workbook, sheet_title, insert_row, before_cells, before_values, grown)


def _insertion_blocker(workbook, sheet_title: str) -> str | None:
    if workbook[sheet_title].tables:
        return "Sheet contains Excel tables; automatic row insertion is not supported."
    for sheet in workbook.worksheets:
        for cell in sheet._cells.values():
            value = cell.value
            if value is None or isinstance(value, (int, float, bool)):
                continue
            if not isinstance(value, str):
                if getattr(cell, "data_type", None) == "f":
                    return (
                        f"{sheet.title}!{cell.coordinate} holds an array/shared formula; "
                        "automatic row insertion is not supported."
                    )
                continue
            if _is_formula(value) and any(name in value.upper() for name in _POSITION_DEPENDENT_FUNCTIONS):
                return (
                    f"{sheet.title}!{cell.coordinate} uses a position-dependent function; "
                    "automatic row insertion is not supported."
                )
    return None


def _shift_bare_formula(formula: str, formula_sheet: str, target_sheet: str, insert_row: int) -> str:
    """Shift a formula stored without its leading ``=`` (sqref, rule, defined name)."""
    return shift_formula_rows(f"={formula}", formula_sheet, target_sheet, insert_row)[1:]


def _shift_sqref(sqref: str, sheet_title: str, insert_row: int) -> str:
    return " ".join(_shift_bare_formula(part, sheet_title, sheet_title, insert_row) for part in sqref.split())


def _shift_row_dimensions(sheet, insert_row: int) -> None:
    for row in sorted((row for row in sheet.row_dimensions if row >= insert_row), reverse=True):
        dimension = sheet.row_dimensions.pop(row)
        dimension.index = row + 1
        sheet.row_dimensions[row + 1] = dimension


def _shift_sheet_rules(sheet, insert_row: int) -> None:
    from openpyxl.formatting.formatting import ConditionalFormattingList
    from openpyxl.worksheet.cell_range import MultiCellRange

    def shift(formula: str) -> str:
        return _shift_bare_formula(formula, sheet.title, sheet.title, insert_row)

    shifted_formatting = ConditionalFormattingList()
    for formatting in sheet.conditional_formatting:
        sqref = _shift_sqref(str(formatting.sqref), sheet.title, insert_row)
        for rule in formatting.rules:
            rule.formula = [shift(formula) for formula in rule.formula]
            shifted_formatting.add(sqref, rule)
    sheet.conditional_formatting = shifted_formatting

    for validation in sheet.data_validations.dataValidation:
        validation.sqref = MultiCellRange(_shift_sqref(str(validation.sqref), sheet.title, insert_row))
        if validation.formula1 is not None:
            validation.formula1 = shift(validation.formula1)
        if validation.formula2 is not None:
            validation.formula2 = shift(validation.formula2)

    if sheet.print_area:
        sheet.print_area = shift(sheet.print_area)


def _shift_defined_names(workbook, sheet_title: str, insert_row: int) -> None:
    scopes = [("", workbook.defined_names)]
    scopes.extend((sheet.title, sheet.defined_names) for sheet in workbook.worksheets)
    for scope_sheet, names in scopes:
        for defined_name in names.values():
            if defined_name.attr_text:
                defined_name.attr_text = _shift_bare_formula(
                    defined_name.attr_text, scope_sheet, sheet_title, insert_row
                )


def _parent_row_blocker(sheet, parent_row: int, insert_row: int) -> str | None:
    for (row, _column), cell in sheet._cells.items():
        if row == parent_row and _is_formula(cell.value) and _grown_sum(cell, insert_row) is None:
            return (
                f"Parent formula {sheet.title}!{cell.coordinate} is not a same-column SUM ending "
                f"at row {insert_row - 1}; manual workbook adjustment required."
            )
    return None


def _grown_sum(cell, insert_row: int) -> str | None:
    match = _SIMPLE_SUM.fullmatch(cell.value)
    range_info = simple_sum_range(cell.value)
    if match is None or range_info is None:
        return None
    start_column, start_row, end_column, end_row = range_info
    if start_column != cell.column_letter or end_column != cell.column_letter:
        return None
    if end_row != insert_row - 1 or start_row > end_row:
        return None
    raw_start_column, raw_start_row, raw_end_column, raw_end_row = match.groups()
    end_absolute = "$" if raw_end_row.startswith("$") else ""
    return f"=SUM({raw_start_column}{raw_start_row}:{raw_end_column}{end_absolute}{insert_row})"


def _grow_parent_sums(sheet, parent_row: int, insert_row: int) -> dict[tuple[int, int], str]:
    grown: dict[tuple[int, int], str] = {}
    for (row, column), cell in sheet._cells.items():
        if row == parent_row and _is_formula(cell.value):
            cell.value = _grown_sum(cell, insert_row)
            grown[(row, column)] = cell.value
    return grown


def _verify_insertion(
    workbook,
    sheet_title: str,
    insert_row: int,
    before_cells: dict[str, dict[tuple[int, int], object]],
    before_values: dict[str, dict[tuple[int, int], float | str]],
    grown: dict[tuple[int, int], str],
) -> str | None:
    failure = "Row insertion safety check failed: "

    def shifted_row(title: str, row: int) -> int:
        return row + 1 if title == sheet_title and row >= insert_row else row

    for title, cells in before_cells.items():
        sheet = workbook[title]
        for (row, column), value in cells.items():
            new_row = shifted_row(title, row)
            if title == sheet_title and (new_row, column) in grown:
                expected = grown[(new_row, column)]
            elif _is_formula(value):
                expected = shift_formula_rows(value, title, sheet_title, insert_row)
            else:
                expected = value
            actual = sheet._cells.get((new_row, column))
            if actual is None or actual.value != expected:
                return failure + f"{title}!{_coordinate(new_row, column)} does not hold its shifted content."
        non_empty = sum(1 for cell in sheet._cells.values() if cell.value is not None)
        if non_empty != len(cells):
            return failure + f"{title} gained or lost cells during insertion."

    for row, column in grown:
        if (row, column) not in before_values[sheet_title]:
            return failure + f"parent total {sheet_title}!{_coordinate(row, column)} could not be evaluated."

    for title, values in before_values.items():
        after_values = evaluate_sheet(workbook, title)
        for (row, column), before in values.items():
            new_row = shifted_row(title, row)
            after = after_values.get((new_row, column))
            if not _same_value(before, after):
                return failure + f"{title}!{_coordinate(new_row, column)} total changed after insertion."
    return None


def _same_value(before: float | str, after: float | str | None) -> bool:
    if isinstance(before, float) and isinstance(after, float):
        return math.isclose(before, after, rel_tol=1e-12, abs_tol=1e-9)
    return before == after


def _coordinate(row: int, column: int) -> str:
    from openpyxl.utils import get_column_letter

    return f"{get_column_letter(column)}{row}"


def shift_formula_rows(formula: str, formula_sheet: str, target_sheet: str, insert_row: int) -> str:
    """Rewrite references to rows >= insert_row on target_sheet as if one row was inserted."""
    tokenizer = Tokenizer(formula)
    for token in tokenizer.items:
        if token.type == Token.OPERAND and token.subtype == Token.RANGE:
            token.value = _shift_reference(token.value, formula_sheet, target_sheet, insert_row)
    return tokenizer.render()


def _shift_reference(reference: str, formula_sheet: str, target_sheet: str, insert_row: int) -> str:
    match = _SHEET_PREFIX.match(reference)
    if match is None:
        return reference
    sheet_prefix = match.group("sheet")
    referenced_sheet = _unquote_sheet(sheet_prefix) if sheet_prefix else formula_sheet
    if referenced_sheet != target_sheet:
        return reference
    parts = match.group("ref").split(":")
    if len(parts) > 2:
        return reference
    shifted_parts = [_shift_reference_part(part, insert_row) for part in parts]
    if any(part is None for part in shifted_parts):
        return reference
    prefix = f"{sheet_prefix}!" if sheet_prefix else ""
    return prefix + ":".join(shifted_parts)


def _shift_reference_part(part: str, insert_row: int) -> str | None:
    match = _REF_PART.match(part)
    if match is None or (match.group("col") is None and match.group("row") is None):
        return None
    if match.group("row") is None:
        return part
    row = int(match.group("row"))
    if row >= insert_row:
        row += 1
    return f"{match.group('col') or ''}{match.group('row_abs')}{row}"


def _unquote_sheet(sheet: str) -> str:
    if sheet.startswith("'") and sheet.endswith("'"):
        return sheet[1:-1].replace("''", "'")
    return sheet


class _Unsupported(Exception):
    """Formula uses a construct the evaluator does not model."""


class _FormulaError(Exception):
    """Formula evaluates to an Excel error value."""


def evaluate_sheet(workbook, sheet_title: str) -> dict[tuple[int, int], float | str]:
    """Evaluate every formula cell the evaluator supports; unsupported ones are omitted.

    Supports numbers, cell references (including other sheets), SUM over ranges,
    + - * / ^ and parentheses. Errors evaluate to Excel error strings such as "#DIV/0!".
    """
    evaluator = _Evaluator(workbook)
    values: dict[tuple[int, int], float | str] = {}
    for (row, column), cell in workbook[sheet_title]._cells.items():
        if not _is_formula(cell.value):
            continue
        try:
            values[(row, column)] = evaluator.formula_value(sheet_title, row, column)
        except _Unsupported:
            continue
    return values


def _is_formula(value: object) -> bool:
    return isinstance(value, str) and value.startswith("=")


class _Evaluator:
    def __init__(self, workbook) -> None:
        self._workbook = workbook
        self._cache: dict[tuple[str, int, int], float | str | _Unsupported] = {}
        self._in_progress: set[tuple[str, int, int]] = set()

    def formula_value(self, sheet_title: str, row: int, column: int) -> float | str:
        try:
            return self._cell_value(sheet_title, row, column)
        except _FormulaError as error:
            return str(error)

    def _cell_value(self, sheet_title: str, row: int, column: int) -> float:
        key = (sheet_title, row, column)
        if key not in self._cache:
            self._cache[key] = self._compute(key)
        result = self._cache[key]
        if isinstance(result, _Unsupported):
            raise result
        if isinstance(result, str):
            raise _FormulaError(result)
        return result

    def _compute(self, key: tuple[str, int, int]) -> float | str | _Unsupported:
        sheet_title, row, column = key
        if sheet_title not in self._workbook.sheetnames:
            return _Unsupported(f"unknown sheet {sheet_title}")
        cell = self._workbook[sheet_title]._cells.get((row, column))
        value = cell.value if cell is not None else None
        if value is None:
            return 0.0
        if isinstance(value, bool):
            return _Unsupported("boolean cell")
        if isinstance(value, (int, float)):
            return float(value)
        if not _is_formula(value):
            return _Unsupported("text cell")
        if key in self._in_progress:
            return _Unsupported("circular reference")
        self._in_progress.add(key)
        try:
            tokens = [
                token for token in Tokenizer(value).items if token.type != Token.WSPACE
            ]
            parser = _FormulaParser(tokens, sheet_title, self)
            return parser.parse()
        except _FormulaError as error:
            return str(error)
        except _Unsupported as unsupported:
            return unsupported
        finally:
            self._in_progress.discard(key)

    def reference_values(self, reference: str, formula_sheet: str) -> list[float]:
        sheet_title, bounds = _reference_bounds(reference, formula_sheet)
        min_col, min_row, max_col, max_row = bounds
        values: list[float] = []
        sheet = self._workbook[sheet_title] if sheet_title in self._workbook.sheetnames else None
        for row in range(min_row, max_row + 1):
            for column in range(min_col, max_col + 1):
                if sheet is not None:
                    cell = sheet._cells.get((row, column))
                    if cell is not None and isinstance(cell.value, str) and not _is_formula(cell.value):
                        continue
                values.append(self._cell_value(sheet_title, row, column))
        return values


def _reference_bounds(reference: str, formula_sheet: str) -> tuple[str, tuple[int, int, int, int]]:
    from openpyxl.utils.cell import range_boundaries

    match = _SHEET_PREFIX.match(reference)
    if match is None:
        raise _Unsupported(f"reference {reference}")
    sheet_prefix = match.group("sheet")
    sheet_title = _unquote_sheet(sheet_prefix) if sheet_prefix else formula_sheet
    try:
        bounds = range_boundaries(match.group("ref").replace("$", ""))
    except ValueError as exc:
        raise _Unsupported(f"reference {reference}") from exc
    if any(bound is None for bound in bounds):
        raise _Unsupported(f"open-ended reference {reference}")
    return sheet_title, bounds


class _FormulaParser:
    """Recursive-descent evaluation of the arithmetic formula subset."""

    def __init__(self, tokens, sheet_title: str, evaluator: _Evaluator) -> None:
        self._tokens = tokens
        self._position = 0
        self._sheet_title = sheet_title
        self._evaluator = evaluator

    def parse(self) -> float:
        value = self._expression()
        if self._peek() is not None:
            raise _Unsupported("trailing tokens")
        return value

    def _peek(self):
        return self._tokens[self._position] if self._position < len(self._tokens) else None

    def _next(self):
        token = self._peek()
        if token is None:
            raise _Unsupported("unexpected end of formula")
        self._position += 1
        return token

    def _peek_operator(self, token_type: str, operators: str) -> str | None:
        token = self._peek()
        if token is not None and token.type == token_type and token.value in operators:
            return token.value
        return None

    def _expression(self) -> float:
        value = self._term()
        while (operator := self._peek_operator(Token.OP_IN, "+-")) is not None:
            self._next()
            right = self._term()
            value = value + right if operator == "+" else value - right
        return value

    def _term(self) -> float:
        value = self._power()
        while (operator := self._peek_operator(Token.OP_IN, "*/")) is not None:
            self._next()
            right = self._power()
            if operator == "*":
                value *= right
            elif right == 0:
                raise _FormulaError("#DIV/0!")
            else:
                value /= right
        return value

    def _power(self) -> float:
        value = self._unary()
        while self._peek_operator(Token.OP_IN, "^") is not None:
            self._next()
            value = value ** self._unary()
        return value

    def _unary(self) -> float:
        operator = self._peek_operator(Token.OP_PRE, "+-")
        if operator is not None:
            self._next()
            value = self._unary()
            return -value if operator == "-" else value
        return self._primary()

    def _primary(self) -> float:
        token = self._next()
        if token.type == Token.OPERAND and token.subtype == Token.NUMBER:
            return float(token.value)
        if token.type == Token.OPERAND and token.subtype == Token.RANGE:
            values = self._evaluator.reference_values(token.value, self._sheet_title)
            if len(values) != 1:
                raise _Unsupported("multi-cell reference outside SUM")
            return values[0]
        if token.type == Token.PAREN and token.subtype == Token.OPEN:
            value = self._expression()
            closing = self._next()
            if closing.type != Token.PAREN or closing.subtype != Token.CLOSE:
                raise _Unsupported("unbalanced parentheses")
            return value
        if token.type == Token.FUNC and token.subtype == Token.OPEN and token.value.upper() == "SUM(":
            return self._sum_arguments()
        raise _Unsupported(f"token {token.value}")

    def _sum_arguments(self) -> float:
        total = 0.0
        while True:
            token = self._peek()
            if token is not None and token.type == Token.OPERAND and token.subtype == Token.RANGE:
                self._next()
                total += sum(self._evaluator.reference_values(token.value, self._sheet_title))
            else:
                total += self._expression()
            separator = self._next()
            if separator.type == Token.FUNC and separator.subtype == Token.CLOSE:
                return total
            if separator.type != Token.SEP or separator.subtype != Token.ARG:
                raise _Unsupported("unexpected SUM argument")
