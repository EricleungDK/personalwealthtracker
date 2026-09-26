from __future__ import annotations

import csv
import json
from collections import Counter
from datetime import datetime
from pathlib import Path

from .category_memory import normalize_merchant_identity
from .config import AppConfig
from .models import (
    CategorizedTransaction,
    CategoryRegistryAddition,
    LocalLLMDiagnostics,
    TrackerUpdate,
    WorkbookStructureChange,
)
from .review_decisions import REJECT_SUGGESTION
from .utils import TRANSACTION_ID_SCHEME_VERSION
from .workbook import WorkbookCategoryOption, rows_in_review, workbook_category_options

SPLIT_HEADERS = (
    "split_role",
    "split_rule",
    "source_transaction_id",
    "allocated_amount",
    "residual_amount",
)
MANUAL_CATEGORY_NOTE = (
    "Blank accepts suggested_category. No suggestion: pick a category or NONE; "
    "a blank row without a suggestion stays in review. NONE rejects the suggestion."
)


def write_outputs(
    output_dir: Path,
    mode: str,
    year: int,
    month: str,
    source_statement: Path,
    tracker_path: Path,
    categorized: list[CategorizedTransaction],
    updates: list[TrackerUpdate],
    structure_changes: list[WorkbookStructureChange] | None = None,
    statement_parser: str = "nordea-pdf",
    workbook_config: AppConfig | None = None,
    category_registry_additions: tuple[CategoryRegistryAddition, ...] = (),
    local_llm_diagnostics: LocalLLMDiagnostics | None = None,
    review_xlsx_path: Path | None = None,
) -> tuple[Path, Path, Path, Path, Path]:
    structure_changes = structure_changes or []
    output_dir.mkdir(parents=True, exist_ok=True)
    period = f"{year}_{month.lower()}"
    report_path = output_dir / f"report_{period}.md"
    audit_path = output_dir / f"audit_{period}.jsonl"
    categorized_path = output_dir / f"categorized_transactions_{period}.csv"
    review_path = output_dir / f"review_required_{period}.csv"
    review_xlsx_path = review_xlsx_path or output_dir / f"review_required_{period}.xlsx"
    category_options = _load_category_options(tracker_path, year, month, workbook_config)
    local_llm_diagnostics = local_llm_diagnostics or LocalLLMDiagnostics()

    _write_report(
        report_path,
        mode,
        year,
        month,
        statement_parser,
        categorized,
        updates,
        structure_changes,
        category_registry_additions,
        local_llm_diagnostics,
        workbook_config,
    )
    _write_audit(
        audit_path,
        mode,
        year,
        month,
        source_statement,
        statement_parser,
        tracker_path,
        categorized,
        updates,
        structure_changes,
        category_registry_additions,
        local_llm_diagnostics,
    )
    _write_categorized_csv(categorized_path, categorized)
    _write_review_csv(review_path, categorized, updates)
    _write_review_workbook(
        review_xlsx_path,
        year,
        month,
        statement_parser,
        categorized,
        updates,
        category_options,
    )
    return report_path, audit_path, categorized_path, review_path, review_xlsx_path


def _load_category_options(
    tracker_path: Path,
    year: int,
    month: str,
    workbook_config: AppConfig | None,
) -> list[WorkbookCategoryOption]:
    if workbook_config is None or not tracker_path.exists():
        return []
    return workbook_category_options(tracker_path, year, month, workbook_config)


def _write_report(
    path: Path,
    mode: str,
    year: int,
    month: str,
    statement_parser: str,
    categorized: list[CategorizedTransaction],
    updates: list[TrackerUpdate],
    structure_changes: list[WorkbookStructureChange],
    category_registry_additions: tuple[CategoryRegistryAddition, ...],
    local_llm_diagnostics: LocalLLMDiagnostics,
    workbook_config: AppConfig | None,
) -> None:
    total = len(categorized)
    review_count = sum(1 for item in categorized if item.review_required)
    write_count = sum(1 for update in updates if update.write_action == "write")
    skip_count = sum(1 for update in updates if update.write_action == "skip")
    classified_count = sum(1 for item in categorized if item.suggested_category)
    no_review_count = sum(1 for item in categorized if not item.review_required)
    unmatched_count = sum(
        1
        for item in categorized
        if item.categorization_method == "unmatched" or item.suggested_category is None
    )
    method_counts = Counter(item.categorization_method for item in categorized)

    lines = [
        "# Monthly Wealth Tracker Automation Report",
        "",
        f"- Reporting month: {month} {year}",
        f"- Mode: {mode}",
        f"- Statement parser: {statement_parser}",
        *(_currency_assumption_lines(workbook_config)),
        f"- Transactions processed: {total}",
        f"- Transactions requiring review: {review_count}",
        f"- Proposed workbook writes: {write_count}",
        f"- Skipped workbook updates: {skip_count}",
        "- Workbook cleanup tasks: not run during monthly update.",
        *(_commit_block_lines(mode, categorized, updates)),
        "",
        "## Categorization Quality",
        "",
        f"- Classification rate: {_format_rate(classified_count, total)}",
        f"- No-review rate: {_format_rate(no_review_count, total)}",
        f"- Unmatched transactions: {unmatched_count}",
        f"- Review-required transactions: {review_count}",
        "- Categorization method counts:",
    ]
    if method_counts:
        for method, count in sorted(method_counts.items()):
            lines.append(f"- {method}: {count}")
    else:
        lines.append("- none: 0")

    if local_llm_diagnostics.enabled:
        lines.extend(
            [
                "",
                "## Local LLM Mode",
                "",
                "- Status: enabled",
                f"- Provider: {local_llm_diagnostics.provider}",
                f"- Endpoint: {local_llm_diagnostics.endpoint}",
                f"- Model: {local_llm_diagnostics.active_model or local_llm_diagnostics.model}",
                f"- Second model: {local_llm_diagnostics.second_model}",
                f"- Fallback model: {local_llm_diagnostics.fallback_model}",
                f"- Eligible rows: {local_llm_diagnostics.eligible_count}",
                f"- Provider calls attempted: {local_llm_diagnostics.attempted_count}",
                f"- Existing-leaf suggestions: {local_llm_diagnostics.existing_leaf_suggestions}",
                f"- New-leaf proposals: {local_llm_diagnostics.new_leaf_proposal_count}",
                f"- No-suggestion responses: {local_llm_diagnostics.no_suggestion_count}",
                (
                    "- Low-confidence responses ignored: "
                    f"{local_llm_diagnostics.low_confidence_response_count}"
                ),
                f"- Invalid responses: {local_llm_diagnostics.invalid_response_count}",
                f"- Provider failures: {local_llm_diagnostics.provider_failure_count}",
                "- Warnings:",
            ]
        )
        if local_llm_diagnostics.warnings:
            for warning in local_llm_diagnostics.warnings:
                lines.append(f"- {warning}")
        else:
            lines.append("- none")

    lines.extend(
        [
            "",
            "## Planned Structure Changes",
            "",
        ]
    )

    if structure_changes:
        for change in structure_changes:
            target_months = ", ".join(
                f"{month} {change.target_year}" for month in change.target_months
            )
            lines.append(
                f"- {change.change_type}: {target_months}, "
                f"{change.source_range or 'unresolved'} -> {change.target_range or 'unresolved'} "
                f"({change.write_action}; {change.reason})"
            )
    else:
        lines.append("- No workbook structure changes planned.")

    lines.extend(
        [
            "",
            "## Category Registry Updates",
            "",
        ]
    )
    if category_registry_additions:
        for addition in category_registry_additions:
            source_ids = ", ".join(addition.source_transaction_ids)
            lines.append(
                f"- Registered {addition.leaf_category} under "
                f"{addition.parent_category} ({source_ids})."
            )
    else:
        lines.append("- No category registry updates.")

    lines.extend(
        [
            "",
            "## Proposed Updates",
            "",
        ]
    )

    if updates:
        categorized_by_id = {item.transaction.transaction_id: item for item in categorized}
        for update in updates:
            lines.append(
                f"- {update.category}: {update.amount} -> {update.target_cell or 'unresolved'} "
                f"({update.write_action}; {update.reason})"
            )
            for transaction_id in update.source_transactions:
                item = categorized_by_id.get(transaction_id)
                if item is None:
                    continue
                lines.append(
                    f"  - Source {transaction_id}: {item.transaction.amount} DKK, "
                    f"{item.reason}"
                )
    else:
        lines.append("- No categorized updates were produced.")

    lines.extend(["", "## Review Required", ""])
    review_items = [item for item in categorized if item.review_required]
    if review_items:
        for item in review_items:
            category = item.suggested_category or "Unmatched"
            lines.append(
                f"- {item.transaction.date}: {item.transaction.amount} DKK, "
                f"{category}, confidence {item.confidence:.2f}, {item.reason}"
            )
    else:
        lines.append("- No transaction-level review items.")

    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def _commit_block_lines(
    mode: str, categorized: list[CategorizedTransaction], updates: list[TrackerUpdate]
) -> list[str]:
    review_count = len(rows_in_review(categorized, updates))
    if mode != "commit" or not review_count:
        return []
    return [f"- Workbook not written: {review_count} row(s) in review."]


def _currency_assumption_lines(workbook_config: AppConfig | None) -> list[str]:
    if workbook_config is None:
        return []
    return [
        f"- Tracker Currency: {workbook_config.tracker_currency}",
        f"- Statement Currency: {workbook_config.statement_currency}",
    ]


def _format_rate(count: int, total: int) -> str:
    percentage = (count / total * 100) if total else 0.0
    return f"{count}/{total} ({percentage:.1f}%)"


def _write_audit(
    path: Path,
    mode: str,
    year: int,
    month: str,
    source_statement: Path,
    statement_parser: str,
    tracker_path: Path,
    categorized: list[CategorizedTransaction],
    updates: list[TrackerUpdate],
    structure_changes: list[WorkbookStructureChange],
    category_registry_additions: tuple[CategoryRegistryAddition, ...],
    local_llm_diagnostics: LocalLLMDiagnostics,
) -> None:
    update_by_transaction = {
        transaction_id: update
        for update in updates
        for transaction_id in update.source_transactions
    }
    timestamp = datetime.now().isoformat(timespec="seconds")
    with path.open("w", encoding="utf-8") as handle:
        if local_llm_diagnostics.enabled:
            payload = {
                "record_type": "local_llm_summary",
                "run_timestamp": timestamp,
                "mode": mode,
                "statement_parser": statement_parser,
                "target_year": year,
                "target_month": month,
                "provider": local_llm_diagnostics.provider,
                "endpoint": local_llm_diagnostics.endpoint,
                "model": local_llm_diagnostics.model,
                "second_model": local_llm_diagnostics.second_model,
                "fallback_model": local_llm_diagnostics.fallback_model,
                "active_model": local_llm_diagnostics.active_model,
                "eligible_count": local_llm_diagnostics.eligible_count,
                "attempted_count": local_llm_diagnostics.attempted_count,
                "existing_leaf_suggestions": local_llm_diagnostics.existing_leaf_suggestions,
                "new_leaf_proposal_count": local_llm_diagnostics.new_leaf_proposal_count,
                "no_suggestion_count": local_llm_diagnostics.no_suggestion_count,
                "low_confidence_response_count": (
                    local_llm_diagnostics.low_confidence_response_count
                ),
                "invalid_response_count": local_llm_diagnostics.invalid_response_count,
                "provider_failure_count": local_llm_diagnostics.provider_failure_count,
            }
            handle.write(json.dumps(payload, ensure_ascii=False) + "\n")
            for warning in local_llm_diagnostics.warnings:
                payload = {
                    "record_type": "local_llm_warning",
                    "run_timestamp": timestamp,
                    "mode": mode,
                    "statement_parser": statement_parser,
                    "target_year": year,
                    "target_month": month,
                    "warning": warning,
                }
                handle.write(json.dumps(payload, ensure_ascii=False) + "\n")
        for addition in category_registry_additions:
            payload = {
                "record_type": "category_registry_addition",
                "run_timestamp": timestamp,
                "mode": mode,
                "statement_parser": statement_parser,
                "source_statement": str(source_statement),
                "target_workbook": str(tracker_path),
                "target_year": year,
                "target_month": month,
                "parent_category": addition.parent_category,
                "leaf_category": addition.leaf_category,
                "source_transaction_ids": list(addition.source_transaction_ids),
            }
            handle.write(json.dumps(payload, ensure_ascii=False) + "\n")
        for change in structure_changes:
            payload = {
                "record_type": "workbook_structure_change",
                "run_timestamp": timestamp,
                "mode": mode,
                "statement_parser": statement_parser,
                "source_statement": str(source_statement),
                "target_workbook": str(tracker_path),
                "target_year": year,
                "target_month": month,
                "change_type": change.change_type,
                "target_structure_year": change.target_year,
                "target_months": list(change.target_months),
                "source_range": change.source_range,
                "target_range": change.target_range,
                "action": change.write_action,
                "reason": change.reason,
            }
            handle.write(json.dumps(payload, ensure_ascii=False) + "\n")
        for item in categorized:
            transaction = item.transaction
            update = update_by_transaction.get(transaction.transaction_id)
            payload = {
                "record_type": "transaction",
                "run_timestamp": timestamp,
                "mode": mode,
                "statement_parser": statement_parser,
                "source_statement": str(source_statement),
                "target_workbook": str(tracker_path),
                "target_year": year,
                "target_month": month,
                "transaction_id": transaction.transaction_id,
                "transaction_date": transaction.date.isoformat(),
                "description": transaction.description,
                "amount": str(transaction.amount),
                "currency": transaction.currency,
                "suggested_category": item.suggested_category,
                "confidence": item.confidence,
                "categorization_method": item.categorization_method,
                "target_row": update.target_row if update else None,
                "target_cell": update.target_cell if update else None,
                "action": update.write_action if update else "review_required",
                "reason": update.reason if update else item.reason,
            }
            handle.write(json.dumps(payload, ensure_ascii=False) + "\n")


def _write_categorized_csv(path: Path, categorized: list[CategorizedTransaction]) -> None:
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(
            [
                "transaction_id",
                "date",
                "description",
                "amount",
                "currency",
                "direction",
                "category",
                "confidence",
                "method",
                "review_required",
            ]
        )
        for item in categorized:
            transaction = item.transaction
            writer.writerow(
                [
                    transaction.transaction_id,
                    transaction.date.isoformat(),
                    transaction.description,
                    transaction.amount,
                    transaction.currency,
                    transaction.direction,
                    item.suggested_category or "",
                    f"{item.confidence:.2f}",
                    item.categorization_method,
                    item.review_required,
                ]
            )


def _write_review_csv(
    path: Path, categorized: list[CategorizedTransaction], updates: list[TrackerUpdate]
) -> None:
    update_by_transaction = {
        transaction_id: update
        for update in updates
        for transaction_id in update.source_transactions
    }
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["type", "id", "date_or_cell", "amount", "category", "reason"])
        for item in categorized:
            if item.review_required or item.categorization_method.startswith("proxy_split_"):
                transaction = item.transaction
                writer.writerow(
                    [
                        item.categorization_method
                        if item.categorization_method.startswith("proxy_split_")
                        else "transaction",
                        transaction.transaction_id,
                        transaction.date.isoformat(),
                        transaction.amount,
                        item.suggested_category or "",
                        item.reason,
                    ]
                )
        for update in updates:
            if update.write_action != "write":
                writer.writerow(
                    [
                        "workbook_update",
                        ",".join(update.source_transactions),
                        update.target_cell or "",
                        update.amount,
                        update.category,
                        update.reason,
                    ]
                )


def _write_review_workbook(
    path: Path,
    year: int,
    month: str,
    statement_parser: str,
    categorized: list[CategorizedTransaction],
    updates: list[TrackerUpdate],
    category_options: list[WorkbookCategoryOption],
) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.comments import Comment
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.datavalidation import DataValidation
    except ImportError as exc:
        raise RuntimeError("openpyxl is required for review workbook output.") from exc

    update_by_transaction = {
        transaction_id: update
        for update in updates
        for transaction_id in update.source_transactions
    }
    workbook = Workbook()
    review_sheet = workbook.active
    review_sheet.title = "Review Required"
    audit_sheet = workbook.create_sheet("Audit")
    options_sheet = workbook.create_sheet("Category Options")
    decision_options_sheet = workbook.create_sheet("Decision Options")
    metadata_sheet = workbook.create_sheet("Run Metadata")

    review_items = sorted(
        rows_in_review(categorized, updates),
        key=lambda item: (bool(item.suggested_category), item.transaction.date),
    )
    review_ids = {item.transaction.transaction_id for item in review_items}
    audit_items = [item for item in categorized if item.transaction.transaction_id not in review_ids]
    audit_headers = [
        "transaction_id",
        "date",
        "description",
        "amount",
        "currency",
        "direction",
        "merchant_identity",
        "category",
        "corrected_category",
        "confidence",
        "source",
        "votes",
        "reason",
        "evidence",
        *(SPLIT_HEADERS if any(item.split_role for item in audit_items) else ()),
    ]
    options_headers = [
        "row_number",
        "category",
        "learnable",
        "status",
        "category_type",
        "allows_new_children",
        "manual_category_option",
        "new_parent_category_option",
    ]
    audit_sheet.append(audit_headers)
    options_sheet.append(options_headers)

    for item in audit_items:
        row = _audit_row(item)
        audit_sheet.append([row.get(header) for header in audit_headers])

    review_headers = [
        "date",
        "description",
        "amount",
        "suggested_category",
        *(
            ("suggested_parent_category",)
            if any(item.new_leaf_parent for item in review_items)
            else ()
        ),
        "manual_category",
        "reason",
        "confidence",
        "new_parent_category",
        "new_leaf_category",
        "learn_to_memory",
        "blocked",
        *(SPLIT_HEADERS if any(item.split_role for item in review_items) else ()),
        "transaction_id",
    ]
    review_sheet.append(review_headers)
    for item in review_items:
        row = _review_row(item, update_by_transaction.get(item.transaction.transaction_id))
        review_sheet.append([row.get(header) for header in review_headers])
    review_sheet.column_dimensions[
        get_column_letter(review_headers.index("transaction_id") + 1)
    ].hidden = True
    manual_category_column = get_column_letter(review_headers.index("manual_category") + 1)
    review_sheet[f"{manual_category_column}1"].comment = Comment(
        MANUAL_CATEGORY_NOTE, author="wealth-tracker"
    )

    leaf_categories = [
        option.category for option in category_options if option.category_type == "leaf"
    ]
    manual_category_options = [REJECT_SUGGESTION, *leaf_categories]
    new_parent_category_options = [
        option.category for option in category_options if option.allows_new_children
    ]
    option_rows = max(
        len(category_options),
        len(manual_category_options),
        len(new_parent_category_options),
    )
    for index in range(option_rows):
        if index < len(category_options):
            option = category_options[index]
            context_values = [
                option.row_number,
                option.category,
                option.learnable,
                option.status,
                option.category_type,
                option.allows_new_children,
            ]
        else:
            context_values = [None, None, None, None, None, None]
        options_sheet.append(
            [
                *context_values,
                manual_category_options[index] if index < len(manual_category_options) else None,
                new_parent_category_options[index]
                if index < len(new_parent_category_options)
                else None,
            ]
        )

    option_column = get_column_letter(options_headers.index("manual_category_option") + 1)
    option_end_row = len(manual_category_options) + 1
    leaf_list_formula = f"'Category Options'!${option_column}$2:${option_column}${option_end_row}"
    # One Decision Options column per review row: no 255-character inline-list cap.
    for row, item in enumerate(review_items, start=2):
        choices = _decision_choices(item, leaf_categories)
        choice_column = get_column_letter(row - 1)
        for choice_row, choice in enumerate(choices, start=1):
            decision_options_sheet.cell(row=choice_row, column=row - 1, value=choice)
        validation = DataValidation(
            type="list",
            formula1=f"'Decision Options'!${choice_column}$1:${choice_column}${len(choices)}",
        )
        _allow_blank_validation(validation)
        review_sheet.add_data_validation(validation)
        validation.add(f"{manual_category_column}{row}")

    correction_validation = DataValidation(type="list", formula1=leaf_list_formula)
    _allow_blank_validation(correction_validation)
    audit_sheet.add_data_validation(correction_validation)
    correction_column = get_column_letter(audit_headers.index("corrected_category") + 1)
    correction_validation.add(f"{correction_column}2:{correction_column}1048576")

    if new_parent_category_options:
        option_end_row = len(new_parent_category_options) + 1
        option_column = get_column_letter(options_headers.index("new_parent_category_option") + 1)
        new_parent_category_validation = DataValidation(
            type="list",
            formula1=f"'Category Options'!${option_column}$2:${option_column}${option_end_row}",
        )
        _allow_blank_validation(new_parent_category_validation)
        review_sheet.add_data_validation(new_parent_category_validation)
        new_parent_category_column = get_column_letter(
            review_headers.index("new_parent_category") + 1
        )
        new_parent_category_validation.add(
            f"{new_parent_category_column}2:{new_parent_category_column}1048576"
        )

    learn_validation = DataValidation(type="list", formula1='"yes,no"')
    _allow_blank_validation(learn_validation)
    review_sheet.add_data_validation(learn_validation)
    learn_column = get_column_letter(review_headers.index("learn_to_memory") + 1)
    learn_validation.add(f"{learn_column}2:{learn_column}1048576")

    for key, value in [
        ("reporting_year", year),
        ("reporting_month", month),
        ("statement_parser", statement_parser),
        ("generated_timestamp", datetime.now().isoformat(timespec="seconds")),
        ("transaction_id_scheme", TRANSACTION_ID_SCHEME_VERSION),
    ]:
        metadata_sheet.append([key, value])

    options_sheet.sheet_state = "hidden"
    decision_options_sheet.sheet_state = "hidden"
    metadata_sheet.sheet_state = "hidden"
    workbook.save(path)
    workbook.close()


def _audit_row(item: CategorizedTransaction) -> dict[str, object]:
    transaction = item.transaction
    return {
        "transaction_id": transaction.transaction_id,
        "date": transaction.date.isoformat(),
        "description": transaction.description,
        "amount": _format_amount(transaction.amount),
        "currency": transaction.currency,
        "direction": transaction.direction,
        "merchant_identity": _merchant_identity(transaction),
        "category": item.suggested_category or "",
        "confidence": f"{item.confidence:.2f}",
        "source": item.categorization_method,
        "votes": _format_votes(item),
        "reason": item.authority_reason,
        "evidence": item.reason,
        **_split_values(item),
    }


def _review_row(item: CategorizedTransaction, update: TrackerUpdate | None) -> dict[str, object]:
    transaction = item.transaction
    return {
        "date": transaction.date.isoformat(),
        "description": transaction.description,
        "amount": _format_amount(transaction.amount),
        "suggested_category": item.suggested_category,
        "suggested_parent_category": item.new_leaf_parent,
        "reason": f"{item.reason} ({item.categorization_method})",
        "confidence": f"{item.confidence:.2f}",
        "blocked": _blocked_note(update),
        "transaction_id": transaction.transaction_id,
        **_split_values(item),
    }


def _split_values(item: CategorizedTransaction) -> dict[str, str | None]:
    values = (
        item.split_role,
        item.split_rule,
        item.source_transaction_id,
        _format_optional_amount(item.allocated_amount),
        _format_optional_amount(item.residual_amount),
    )
    return dict(zip(SPLIT_HEADERS, values, strict=True))


def _blocked_note(update: TrackerUpdate | None) -> str | None:
    """Only a `review` workbook update blocks the commit; `write` and `skip` do not."""
    if update is None or update.write_action != "review":
        return None
    return f"{update.target_cell}: {update.reason}" if update.target_cell else update.reason


def _format_votes(item: CategorizedTransaction) -> str:
    return "; ".join(
        f"{vote.category or 'NONE'} ({vote.source}, {vote.confidence:.2f})" for vote in item.votes
    )


def _decision_choices(item: CategorizedTransaction, leaf_categories: list[str]) -> list[str]:
    """Suggestion and alternatives, then NONE, then every leaf.

    A proposed new leaf is not offered; it is accepted via the new-leaf cells.
    """
    proposed = item.suggested_category if item.new_leaf_parent else None
    candidates = [
        item.suggested_category,
        *item.alternatives,
        *(vote.category for vote in item.votes),
    ]
    suggestions = [category for category in candidates if category and category != proposed]
    return list(dict.fromkeys([*suggestions, REJECT_SUGGESTION, *leaf_categories]))


def _allow_blank_validation(validation) -> None:
    validation.allowBlank = True
    validation.allow_blank = True


def _format_amount(amount) -> str:
    return format(amount, "f")


def _format_optional_amount(amount) -> str | None:
    if amount is None:
        return None
    return _format_amount(amount)


def _merchant_identity(transaction) -> str:
    return normalize_merchant_identity(transaction.merchant or transaction.description)
