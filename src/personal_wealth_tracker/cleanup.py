from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from pathlib import Path


@dataclass(frozen=True)
class WorkbookCleanupChange:
    sheet_name: str
    cell: str
    before: str
    after: str


@dataclass(frozen=True)
class WorkbookCleanupResult:
    mode: str
    changes: list[WorkbookCleanupChange]
    report_path: Path
    output_workbook_path: Path | None = None


def run_currency_label_cleanup(
    tracker_path: Path,
    tracker_currency: str,
    output_dir: Path,
    commit: bool = False,
) -> WorkbookCleanupResult:
    changes = plan_currency_label_cleanup(tracker_path, tracker_currency)
    mode = "cleanup-commit" if commit else "cleanup-dry-run"
    output_dir.mkdir(parents=True, exist_ok=True)
    output_workbook_path = None
    if commit:
        output_workbook_path = _write_currency_label_cleanup_copy(
            tracker_path,
            output_dir,
            changes,
        )
    report_path = output_dir / "workbook_cleanup_currency_labels.md"
    _write_cleanup_report(report_path, mode, changes)
    return WorkbookCleanupResult(
        mode=mode,
        changes=changes,
        report_path=report_path,
        output_workbook_path=output_workbook_path,
    )


def plan_currency_label_cleanup(
    tracker_path: Path,
    tracker_currency: str,
) -> list[WorkbookCleanupChange]:
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError("openpyxl is required for workbook access. Install with `uv sync`.") from exc

    workbook = openpyxl.load_workbook(tracker_path)
    changes: list[WorkbookCleanupChange] = []
    try:
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if not isinstance(cell.value, str):
                        continue
                    after = cell.value.replace("EUR", tracker_currency)
                    if after != cell.value:
                        changes.append(
                            WorkbookCleanupChange(
                                sheet_name=sheet.title,
                                cell=cell.coordinate,
                                before=cell.value,
                                after=after,
                            )
                        )
    finally:
        workbook.close()
    return changes


def _write_cleanup_report(
    path: Path,
    mode: str,
    changes: list[WorkbookCleanupChange],
) -> None:
    lines = [
        "# Workbook Cleanup Report",
        "",
        f"- Mode: {mode}",
        "- Cleanup task: currency-label text replacement",
        "- Structure changes: none",
        "",
        "## Planned Text Changes",
        "",
    ]
    if changes:
        for change in changes:
            lines.append(
                f'- {change.sheet_name}!{change.cell}: "{change.before}" -> "{change.after}"'
            )
    else:
        lines.append("- No currency-label text changes found.")
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def _write_currency_label_cleanup_copy(
    tracker_path: Path,
    output_dir: Path,
    changes: list[WorkbookCleanupChange],
) -> Path:
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError("openpyxl is required for workbook access. Install with `uv sync`.") from exc

    workbook = openpyxl.load_workbook(tracker_path)
    try:
        for change in changes:
            workbook[change.sheet_name][change.cell] = change.after
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_path = (
            output_dir
            / f"{tracker_path.stem}_currency_labels_cleaned_{timestamp}{tracker_path.suffix}"
        )
        workbook.save(output_path)
    finally:
        workbook.close()
    return output_path
