from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path


TEMPLATE_WORKBOOK_ID = "local-wealth-tracker-template"
SUPPORTED_TEMPLATE_VERSION = "1.0"
TEMPLATE_METADATA_SHEET = "Template Metadata"
TEMPLATE_SCHEMA = "net-worth-v1"
DEFAULT_TEMPLATE_SHEET_NAME = "Net worth"
DEFAULT_TEMPLATE_CURRENCY = "DKK"
DEFAULT_TEMPLATE_YEAR = 2026

MONTHS = ("Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec")

TEMPLATE_ROWS = (
    (5, "Income (net)", "derived"),
    (6, "Full-time job (net)", "leaf"),
    (7, "Other income", "leaf"),
    (8, "Cashflow", "derived"),
    (9, "Living expenses", "parent"),
    (10, "Rent (monthly)", "leaf"),
    (11, "Groceries (monthly)", "leaf"),
    (12, "Transportation", "leaf"),
    (13, "Services", "parent"),
    (14, "Mobile phone (monthly)", "leaf"),
    (15, "Internet (monthly)", "leaf"),
    (16, "Cloud services", "leaf"),
    (17, "Insurance", "parent"),
    (18, "Insurance (monthly)", "leaf"),
    (19, "Investments", "parent"),
    (20, "Brokerage contributions", "leaf"),
    (21, "Pension contributions", "leaf"),
    (22, "Assets", "parent"),
    (23, "Cash savings", "leaf"),
    (24, "Brokerage account", "leaf"),
    (25, "Pension account", "leaf"),
    (26, "Total net worth", "derived"),
)


@dataclass(frozen=True)
class TemplateWorkbookMetadata:
    template_id: str
    template_version: str
    workbook_kind: str
    tracker_currency: str
    template_schema: str


def create_template_workbook(
    output_path: Path,
    *,
    tracker_currency: str = DEFAULT_TEMPLATE_CURRENCY,
    start_year: int = DEFAULT_TEMPLATE_YEAR,
    sheet_name: str = DEFAULT_TEMPLATE_SHEET_NAME,
) -> Path:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError("openpyxl is required for workbook access. Install with `uv sync`.") from exc

    output_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    metadata = workbook.create_sheet(TEMPLATE_METADATA_SHEET)
    metadata.sheet_state = "hidden"

    _write_metadata(metadata, tracker_currency)

    title_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
    section_fill = PatternFill(fill_type="solid", fgColor="E2F0D9")
    derived_fill = PatternFill(fill_type="solid", fgColor="FCE4D6")
    header_font = Font(bold=True)
    section_font = Font(bold=True)

    sheet["A1"] = "Local Wealth-Tracker Agent Template"
    sheet["B1"] = "Category"
    sheet["C1"] = "Tracker Currency"
    sheet["D1"] = tracker_currency
    for cell in sheet[1]:
        cell.font = header_font
        cell.fill = title_fill

    sheet.merge_cells(start_row=2, start_column=3, end_row=2, end_column=14)
    sheet.cell(row=2, column=3).value = start_year
    sheet.cell(row=2, column=3).alignment = Alignment(horizontal="center")
    sheet.cell(row=2, column=3).font = header_font

    for offset, month in enumerate(MONTHS, start=3):
        sheet.cell(row=3, column=offset).value = month
        sheet.cell(row=3, column=offset).font = header_font
        sheet.cell(row=3, column=offset).alignment = Alignment(horizontal="center")
        sheet.column_dimensions[get_column_letter(offset)].width = 14

    sheet.column_dimensions["A"].width = 4
    sheet.column_dimensions["B"].width = 32

    for row, label, kind in TEMPLATE_ROWS:
        label_cell = sheet.cell(row=row, column=2)
        label_cell.value = label
        if kind in {"parent", "derived"}:
            label_cell.font = section_font
            label_cell.fill = derived_fill if kind == "derived" else section_fill

        for column in range(3, 15):
            value_cell = sheet.cell(row=row, column=column)
            value_cell.number_format = "#,##0.00"
            if kind == "derived":
                value_cell.fill = derived_fill
            elif kind == "parent":
                value_cell.fill = section_fill
            formula = _formula_for_row(row, column)
            if formula is not None:
                value_cell.value = formula

    sheet.freeze_panes = "C4"
    workbook.save(output_path)
    workbook.close()
    return output_path


def validate_template_workbook(
    workbook_path: Path,
    *,
    supported_versions: tuple[str, ...] = (SUPPORTED_TEMPLATE_VERSION,),
) -> TemplateWorkbookMetadata:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl is required for workbook access. Install with `uv sync`.") from exc

    workbook = load_workbook(workbook_path, data_only=False)
    try:
        if TEMPLATE_METADATA_SHEET not in workbook.sheetnames:
            raise ValueError(f"Template metadata sheet {TEMPLATE_METADATA_SHEET!r} not found.")
        metadata = _read_metadata(workbook[TEMPLATE_METADATA_SHEET])
    finally:
        workbook.close()

    if metadata.template_id != TEMPLATE_WORKBOOK_ID:
        raise ValueError(f"Unsupported template workbook id {metadata.template_id!r}.")
    if metadata.template_version not in supported_versions:
        raise ValueError(
            f"Unsupported template workbook version {metadata.template_version!r}; "
            f"supported versions: {', '.join(supported_versions)}."
        )
    if metadata.workbook_kind != "synthetic_template":
        raise ValueError(f"Unsupported template workbook kind {metadata.workbook_kind!r}.")
    if metadata.template_schema != TEMPLATE_SCHEMA:
        raise ValueError(f"Unsupported template workbook schema {metadata.template_schema!r}.")
    return metadata


def customize_template_workbook(
    template_path: Path,
    output_path: Path,
    *,
    tracker_currency: str | None = None,
    label_renames: dict[str, str] | None = None,
) -> Path:
    validate_template_workbook(template_path)
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl is required for workbook access. Install with `uv sync`.") from exc

    workbook = load_workbook(template_path, data_only=False)
    try:
        sheet = _template_sheet(workbook)
        _validate_supported_template_layout(sheet)
        renames = label_renames or {}
        _validate_label_renames(sheet, renames)
        for row in range(1, sheet.max_row + 1):
            current = sheet.cell(row=row, column=2).value
            if current in renames:
                sheet.cell(row=row, column=2).value = renames[str(current)]

        if tracker_currency is not None:
            tracker_currency = tracker_currency.upper()
            sheet["D1"] = tracker_currency
            _write_metadata(workbook[TEMPLATE_METADATA_SHEET], tracker_currency)

        output_path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(output_path)
    finally:
        workbook.close()
    return output_path


def _write_metadata(sheet, tracker_currency: str) -> None:
    rows = (
        ("key", "value"),
        ("template_id", TEMPLATE_WORKBOOK_ID),
        ("template_version", SUPPORTED_TEMPLATE_VERSION),
        ("workbook_kind", "synthetic_template"),
        ("tracker_currency", tracker_currency),
        ("template_schema", TEMPLATE_SCHEMA),
    )
    for row_index, (key, value) in enumerate(rows, start=1):
        sheet.cell(row=row_index, column=1).value = key
        sheet.cell(row=row_index, column=2).value = value


def _read_metadata(sheet) -> TemplateWorkbookMetadata:
    values = {}
    for row in range(1, sheet.max_row + 1):
        key = sheet.cell(row=row, column=1).value
        value = sheet.cell(row=row, column=2).value
        if key in (None, ""):
            continue
        values[str(key)] = "" if value is None else str(value)

    required_keys = (
        "template_id",
        "template_version",
        "workbook_kind",
        "tracker_currency",
        "template_schema",
    )
    missing_keys = [key for key in required_keys if not values.get(key)]
    if missing_keys:
        raise ValueError(f"Template metadata missing required keys: {', '.join(missing_keys)}.")

    return TemplateWorkbookMetadata(
        template_id=values["template_id"],
        template_version=values["template_version"],
        workbook_kind=values["workbook_kind"],
        tracker_currency=values["tracker_currency"],
        template_schema=values["template_schema"],
    )


def _template_sheet(workbook):
    sheet_names = [name for name in workbook.sheetnames if name != TEMPLATE_METADATA_SHEET]
    if len(sheet_names) != 1:
        raise ValueError(
            "Unsupported template workbook edit: expected exactly one tracker sheet "
            f"besides {TEMPLATE_METADATA_SHEET!r}."
        )
    return workbook[sheet_names[0]]


def _validate_supported_template_layout(sheet) -> None:
    for row, label, _kind in TEMPLATE_ROWS:
        actual_label = sheet.cell(row=row, column=2).value
        if actual_label != label:
            raise ValueError(
                "Unsupported template workbook edit: expected "
                f"{label!r} in B{row}, found {actual_label!r}."
            )
        for column in range(3, 15):
            expected_formula = _formula_for_row(row, column)
            if expected_formula is None:
                continue
            actual_formula = sheet.cell(row=row, column=column).value
            coordinate = sheet.cell(row=row, column=column).coordinate
            if actual_formula != expected_formula:
                raise ValueError(
                    "Unsupported template workbook edit: expected formula "
                    f"{coordinate} to be {expected_formula!r}, found {actual_formula!r}."
                )


def _validate_label_renames(sheet, label_renames: dict[str, str]) -> None:
    labels = [
        str(sheet.cell(row=row, column=2).value)
        for row, _label, _kind in TEMPLATE_ROWS
        if sheet.cell(row=row, column=2).value
    ]
    known_labels = set(labels)
    unknown_labels = sorted(set(label_renames) - known_labels)
    if unknown_labels:
        raise ValueError(
            "Unsupported template workbook edit: cannot rename unknown template label(s): "
            f"{', '.join(unknown_labels)}."
        )

    renamed_labels = []
    for old_label, new_label in label_renames.items():
        if not new_label or not new_label.strip():
            raise ValueError(
                f"Unsupported template workbook edit: replacement for {old_label!r} is blank."
            )
        renamed_labels.append(new_label.strip())

    final_labels = [label_renames.get(label, label) for label in labels]
    if len(set(final_labels)) != len(final_labels):
        raise ValueError(
            "Unsupported template workbook edit: category and section labels must be unique."
        )


def _formula_for_row(row: int, column: int) -> str | None:
    column_letter = _column_letter(column)
    formulas = {
        5: f"=SUM({column_letter}6:{column_letter}7)",
        8: f"={column_letter}5-SUM({column_letter}9,{column_letter}13,{column_letter}17,{column_letter}19)",
        9: f"=SUM({column_letter}10:{column_letter}12)",
        13: f"=SUM({column_letter}14:{column_letter}16)",
        17: f"=SUM({column_letter}18:{column_letter}18)",
        19: f"=SUM({column_letter}20:{column_letter}21)",
        22: f"=SUM({column_letter}23:{column_letter}25)",
        26: f"={column_letter}22",
    }
    return formulas.get(row)


def _column_letter(column: int) -> str:
    try:
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError("openpyxl is required for workbook access. Install with `uv sync`.") from exc

    return get_column_letter(column)
