from __future__ import annotations

import warnings
from dataclasses import dataclass
from pathlib import Path

from .models import CategorizedTransaction
from .utils import TRANSACTION_ID_SCHEME_VERSION, normalize_month


@dataclass(frozen=True)
class MonthlyReviewDecision:
    transaction_id: str
    manual_category: str = ""
    new_parent_category: str = ""
    new_leaf_category: str = ""

    @property
    def category(self) -> str:
        return self.new_leaf_category or self.manual_category


def load_monthly_review_decisions(
    path: Path,
    year: int,
    month: str,
) -> dict[str, MonthlyReviewDecision]:
    workbook = _load_workbook_without_extension_warning(path)
    try:
        if "Run Metadata" not in workbook.sheetnames:
            raise ValueError("Review decisions workbook is missing 'Run Metadata'.")
        if "Review Required" not in workbook.sheetnames:
            raise ValueError("Review decisions workbook is missing 'Review Required'.")

        metadata = _metadata(workbook["Run Metadata"])
        _validate_metadata(metadata, year, month)
        return _decisions(workbook["Review Required"])
    finally:
        workbook.close()


def _load_workbook_without_extension_warning(path: Path):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl is required for review decision import.") from exc

    with warnings.catch_warnings():
        warnings.filterwarnings(
            "ignore",
            message="Data Validation extension is not supported and will be removed",
            category=UserWarning,
        )
        return load_workbook(path, data_only=True)


def apply_monthly_review_decisions(
    categorized: list[CategorizedTransaction],
    decisions: dict[str, MonthlyReviewDecision],
) -> list[CategorizedTransaction]:
    if not decisions:
        return categorized

    current_ids = {item.transaction.transaction_id for item in categorized}
    missing_ids = sorted(set(decisions) - current_ids)
    if missing_ids:
        raise ValueError(
            "Review decisions contain transaction IDs not present in current statement: "
            + ", ".join(missing_ids)
        )

    return [
        _apply_decision(item, decisions[item.transaction.transaction_id])
        if item.transaction.transaction_id in decisions
        else item
        for item in categorized
    ]


def validate_monthly_review_decision_categories(
    decisions: dict[str, MonthlyReviewDecision],
    valid_categories: set[str],
) -> None:
    missing_categories = sorted(
        {
            decision.manual_category
            for decision in decisions.values()
            if decision.manual_category and decision.manual_category not in valid_categories
        }
    )
    if missing_categories:
        raise ValueError(
            "Review decisions contain manual_category values not found in the current "
            "tracker workbook or YAML leaf category registry: "
            + ", ".join(missing_categories)
        )


def _metadata(sheet) -> dict[str, object]:
    return {
        str(sheet.cell(row=row, column=1).value): sheet.cell(row=row, column=2).value
        for row in range(1, sheet.max_row + 1)
        if sheet.cell(row=row, column=1).value not in (None, "")
    }


def _validate_metadata(metadata: dict[str, object], year: int, month: str) -> None:
    workbook_year = int(metadata.get("reporting_year", 0))
    workbook_month = normalize_month(str(metadata.get("reporting_month", "")))
    target_month = normalize_month(month)
    if workbook_year != year or workbook_month != target_month:
        raise ValueError(
            "review decisions reporting period "
            f"{workbook_month} {workbook_year} does not match requested {target_month} {year}."
        )

    scheme = str(metadata.get("transaction_id_scheme", ""))
    if scheme != TRANSACTION_ID_SCHEME_VERSION:
        raise ValueError(
            "Unsupported review decision transaction ID scheme "
            f"{scheme!r}; expected {TRANSACTION_ID_SCHEME_VERSION!r}. "
            "Regenerate the review workbook from a fresh dry run before applying decisions."
        )


def _decisions(sheet) -> dict[str, MonthlyReviewDecision]:
    headers = {
        str(cell.value): index
        for index, cell in enumerate(sheet[1], start=1)
        if cell.value not in (None, "")
    }
    transaction_id_column = _required_column(headers, "transaction_id")
    manual_category_column = _required_column(headers, "manual_category")
    new_parent_category_column = headers.get("new_parent_category")
    new_leaf_category_column = headers.get("new_leaf_category")

    decisions: dict[str, MonthlyReviewDecision] = {}
    for row in range(2, sheet.max_row + 1):
        manual_category = str(sheet.cell(row=row, column=manual_category_column).value or "").strip()
        new_parent_category = _optional_stripped_cell(sheet, row, new_parent_category_column)
        new_leaf_category = _optional_display_cell(sheet, row, new_leaf_category_column)
        if manual_category and new_leaf_category:
            raise ValueError(
                f"Review decision row {row} cannot fill both manual_category and "
                "new_leaf_category."
            )
        if new_leaf_category and not new_parent_category:
            raise ValueError(
                f"Review decision row {row} with new_leaf_category must include "
                "new_parent_category."
            )
        if not manual_category and not new_leaf_category:
            continue
        transaction_id = str(sheet.cell(row=row, column=transaction_id_column).value or "").strip()
        if not transaction_id:
            raise ValueError(f"Review decision row {row} is missing transaction_id.")
        if transaction_id in decisions:
            raise ValueError(f"Review decision transaction_id {transaction_id!r} appears more than once.")
        decisions[transaction_id] = MonthlyReviewDecision(
            transaction_id=transaction_id,
            manual_category=manual_category,
            new_parent_category=new_parent_category,
            new_leaf_category=new_leaf_category,
        )
    return decisions


def _optional_stripped_cell(sheet, row: int, column: int | None) -> str:
    if column is None:
        return ""
    return str(sheet.cell(row=row, column=column).value or "").strip()


def _optional_display_cell(sheet, row: int, column: int | None) -> str:
    if column is None:
        return ""
    value = str(sheet.cell(row=row, column=column).value or "")
    return value if value.strip() else ""


def _required_column(headers: dict[str, int], column: str) -> int:
    if column not in headers:
        raise ValueError(f"Review decisions workbook is missing {column!r} column.")
    return headers[column]


def _apply_decision(
    item: CategorizedTransaction,
    decision: MonthlyReviewDecision,
) -> CategorizedTransaction:
    return CategorizedTransaction(
        transaction=item.transaction,
        suggested_category=decision.category,
        confidence=1.0,
        categorization_method="monthly_review_decision",
        review_required=False,
        reason="Monthly review decision.",
        source_transaction_id=item.source_transaction_id,
        split_rule=item.split_rule,
        split_role=item.split_role,
        allocated_amount=item.allocated_amount,
        residual_amount=item.residual_amount,
    )
