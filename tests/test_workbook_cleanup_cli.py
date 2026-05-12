from pathlib import Path

from openpyxl import Workbook, load_workbook

from personal_wealth_tracker.cli import main


def test_cleanup_currency_labels_dry_run_reports_text_changes_without_writing(
    tmp_path, capsys
):
    tracker = tmp_path / "tracker.xlsx"
    output_dir = tmp_path / "reports"
    _create_tracker_with_currency_labels(tracker)

    exit_code = main(
        [
            "cleanup-currency-labels",
            "--tracker",
            str(tracker),
            "--tracker-currency",
            "DKK",
            "--output-dir",
            str(output_dir),
        ]
    )

    captured = capsys.readouterr()
    workbook = load_workbook(tracker)
    try:
        sheet = workbook["Net worth"]
        assert sheet["A1"].value == "Tracker currency: EUR"
        assert sheet["A2"].value == "Assumptions use EUR labels"
    finally:
        workbook.close()

    report_path = output_dir / "workbook_cleanup_currency_labels.md"
    assert exit_code == 0
    assert "Mode: cleanup-dry-run" in captured.out
    assert "Planned cleanup changes: 2" in captured.out
    assert f"Cleanup report: {report_path}" in captured.out
    assert report_path.exists()

    report = report_path.read_text(encoding="utf-8")
    assert "# Workbook Cleanup Report" in report
    assert "- Mode: cleanup-dry-run" in report
    assert '- Net worth!A1: "Tracker currency: EUR" -> "Tracker currency: DKK"' in report
    assert '- Net worth!A2: "Assumptions use EUR labels" -> "Assumptions use DKK labels"' in report
    assert "- Structure changes: none" in report


def test_cleanup_currency_labels_commit_writes_only_to_copied_workbook(tmp_path, capsys):
    tracker = tmp_path / "tracker.xlsx"
    output_dir = tmp_path / "reports"
    _create_tracker_with_currency_labels(tracker)

    exit_code = main(
        [
            "cleanup-currency-labels",
            "--tracker",
            str(tracker),
            "--tracker-currency",
            "DKK",
            "--output-dir",
            str(output_dir),
            "--commit",
        ]
    )

    captured = capsys.readouterr()
    copied_paths = list(output_dir.glob("tracker_currency_labels_cleaned_*.xlsx"))
    assert exit_code == 0
    assert len(copied_paths) == 1
    assert f"Output workbook: {copied_paths[0]}" in captured.out

    original = load_workbook(tracker)
    copied = load_workbook(copied_paths[0])
    try:
        original_sheet = original["Net worth"]
        copied_sheet = copied["Net worth"]
        assert original_sheet["A1"].value == "Tracker currency: EUR"
        assert original_sheet["A2"].value == "Assumptions use EUR labels"
        assert copied_sheet["A1"].value == "Tracker currency: DKK"
        assert copied_sheet["A2"].value == "Assumptions use DKK labels"
        assert copied_sheet["A3"].value == "Already DKK"
        assert copied_sheet["A4"].value == 123
    finally:
        original.close()
        copied.close()

    report = (output_dir / "workbook_cleanup_currency_labels.md").read_text(encoding="utf-8")
    assert "- Mode: cleanup-commit" in report


def _create_tracker_with_currency_labels(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Net worth"
    sheet["A1"] = "Tracker currency: EUR"
    sheet["A2"] = "Assumptions use EUR labels"
    sheet["A3"] = "Already DKK"
    sheet["A4"] = 123
    workbook.save(path)
    workbook.close()
