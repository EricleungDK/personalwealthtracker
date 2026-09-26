from datetime import date
from decimal import Decimal

from openpyxl import Workbook, load_workbook

from personal_wealth_tracker.models import Authority, CategorizedTransaction, Transaction
from personal_wealth_tracker.reporting import write_outputs
from personal_wealth_tracker.review_decisions import (
    MonthlyReviewDecision,
    load_monthly_review_decisions,
)
from personal_wealth_tracker.utils import TRANSACTION_ID_SCHEME_VERSION

RESIDUAL_ID = "tx-source:split:example_transfer:residual"

EXPECTED_DECISIONS = {
    "tx-accept": MonthlyReviewDecision("tx-accept", manual_category="Traveling"),
    "tx-reject": MonthlyReviewDecision("tx-reject", manual_category="NONE"),
    "tx-explicit": MonthlyReviewDecision("tx-explicit", manual_category="Apple Cloud"),
    "tx-new-leaf": MonthlyReviewDecision(
        "tx-new-leaf", new_parent_category="Living expenses", new_leaf_category="Pet Supplies"
    ),
    "tx-one-off": MonthlyReviewDecision(
        "tx-one-off", manual_category="Traveling", learn_to_memory=False
    ),
    "tx-proposal": MonthlyReviewDecision(
        "tx-proposal",
        new_parent_category="Services",
        new_leaf_category="Claude subscription",
        accepted_subscription_proposal=True,
    ),
    RESIDUAL_ID: MonthlyReviewDecision(RESIDUAL_ID, manual_category="Traveling"),
    "tx-auto": MonthlyReviewDecision(
        "tx-auto", manual_category="Apple Cloud", audit_correction=True
    ),
}


def test_filled_exception_sheet_loads_every_decision_kind(tmp_path):
    _, _, _, _, sheet_path = write_outputs(
        output_dir=tmp_path,
        mode="dry-run",
        year=2026,
        month="May",
        source_statement=tmp_path / "statement.csv",
        tracker_path=tmp_path / "tracker.xlsx",
        categorized=[
            _review("tx-accept", "Traveling"),
            _review("tx-reject", "Traveling"),
            _review("tx-explicit", None),
            _review("tx-new-leaf", None),
            _review("tx-one-off", None),
            _review("tx-proposal", "Claude subscription", new_leaf_parent="Services"),
            _review(RESIDUAL_ID, None, split_role="residual"),
            _review("tx-open", None),
            _review("tx-auto", "Food& Drinks (monthly)", authority=Authority.auto),
        ],
        updates=[],
    )

    _fill(
        sheet_path,
        "Review Required",
        {
            "tx-reject": {"manual_category": "NONE"},
            "tx-explicit": {"manual_category": "Apple Cloud"},
            "tx-new-leaf": {
                "new_parent_category": "Living expenses",
                "new_leaf_category": "Pet Supplies",
            },
            "tx-one-off": {"manual_category": "Traveling", "learn_to_memory": "no"},
            RESIDUAL_ID: {"manual_category": "Traveling"},
        },
    )
    _fill(sheet_path, "Audit", {"tx-auto": {"corrected_category": "Apple Cloud"}})

    assert load_monthly_review_decisions(sheet_path, 2026, "May") == EXPECTED_DECISIONS


def test_pre_operator_layout_exception_sheet_still_loads(tmp_path):
    sheet_path = tmp_path / "review_required_2026_may.xlsx"
    workbook = Workbook()
    review = workbook.active
    review.title = "Review Required"
    review.append(PRE_OPERATOR_LAYOUT_HEADERS)
    for transaction_id, values in {
        "tx-accept": {"suggested_category": "Traveling"},
        "tx-reject": {"suggested_category": "Traveling", "manual_category": "NONE"},
        "tx-explicit": {"manual_category": "Apple Cloud"},
        "tx-new-leaf": {
            "new_parent_category": "Living expenses",
            "new_leaf_category": "Pet Supplies",
        },
        "tx-one-off": {"manual_category": "Traveling", "learn_to_memory": "no"},
        "tx-proposal": {
            "suggested_category": "Claude subscription",
            "suggested_parent_category": "Services",
        },
        RESIDUAL_ID: {"manual_category": "Traveling", "split_role": "residual"},
        "tx-open": {"method": "unmatched", "workbook_action": "write"},
    }.items():
        row = {"transaction_id": transaction_id, "description": "SHOP", **values}
        review.append([row.get(header) for header in PRE_OPERATOR_LAYOUT_HEADERS])
    audit = workbook.create_sheet("Audit")
    audit.append(["transaction_id", "corrected_category"])
    audit.append(["tx-auto", "Apple Cloud"])
    metadata = workbook.create_sheet("Run Metadata")
    for key, value in [
        ("reporting_year", 2026),
        ("reporting_month", "May"),
        ("transaction_id_scheme", TRANSACTION_ID_SCHEME_VERSION),
    ]:
        metadata.append([key, value])
    workbook.save(sheet_path)
    workbook.close()

    assert load_monthly_review_decisions(sheet_path, 2026, "May") == EXPECTED_DECISIONS


PRE_OPERATOR_LAYOUT_HEADERS = [
    "transaction_id",
    "date",
    "description",
    "amount",
    "manual_category",
    "new_parent_category",
    "new_leaf_category",
    "learn_to_memory",
    "split_role",
    "split_rule",
    "source_transaction_id",
    "allocated_amount",
    "residual_amount",
    "suggested_category",
    "suggested_parent_category",
    "method",
    "reason",
    "workbook_action",
    "target_cell",
    "workbook_reason",
    "merchant_identity",
    "confidence",
    "direction",
]


def _review(
    transaction_id: str,
    category: str | None,
    authority: Authority = Authority.review,
    new_leaf_parent: str | None = None,
    split_role: str | None = None,
) -> CategorizedTransaction:
    return CategorizedTransaction(
        transaction=Transaction(
            transaction_id=transaction_id,
            date=date(2026, 5, 15),
            interest_date=None,
            description=f"SHOP {transaction_id}",
            amount=Decimal("-40.00"),
            currency="DKK",
            direction="expense",
        ),
        suggested_category=category,
        confidence=0.9 if category else 0.0,
        categorization_method="local_llm_gemma",
        reason="Model suggestion.",
        authority=authority,
        new_leaf_parent=new_leaf_parent,
        split_role=split_role,
    )


def _fill(path, sheet_name: str, values_by_id: dict[str, dict[str, str]]) -> None:
    workbook = load_workbook(path)
    try:
        sheet = workbook[sheet_name]
        headers = {cell.value: index for index, cell in enumerate(sheet[1], start=1)}
        for row in range(2, sheet.max_row + 1):
            values = values_by_id.get(sheet.cell(row=row, column=headers["transaction_id"]).value)
            for header, value in (values or {}).items():
                sheet.cell(row=row, column=headers[header]).value = value
        workbook.save(path)
    finally:
        workbook.close()
