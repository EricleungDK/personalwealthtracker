from dataclasses import replace
from datetime import date
from decimal import Decimal

from openpyxl import Workbook, load_workbook

from personal_wealth_tracker.config import AppConfig, CategoryRegistry
from personal_wealth_tracker.models import (
    Authority,
    CategorizedTransaction,
    TrackerUpdate,
    Transaction,
    Vote,
    WorkbookStructureChange,
)
from personal_wealth_tracker.reporting import write_outputs
from personal_wealth_tracker.utils import TRANSACTION_ID_SCHEME_VERSION


def test_report_includes_refund_and_expense_claim_source_reasons(tmp_path):
    categorized = [
        _categorized(
            "tx-refund",
            "NETTO REFUND",
            "25.00",
            "Food& Drinks (monthly)",
            "Refund matched expense keyword rule; nets against category in reporting month.",
        ),
        _categorized(
            "tx-claim",
            "ACME EXPENSE CLAIM",
            "125.00",
            "Expense claims",
            "Expense claim keyword rule match.",
        ),
    ]
    updates = [
        _update(
            "Food& Drinks (monthly)",
            Decimal("-25.00"),
            ("tx-refund",),
            "D5",
        ),
        _update(
            "Expense claims",
            Decimal("125.00"),
            ("tx-claim",),
            "D22",
        ),
    ]

    report_path, _, _, _, _ = write_outputs(
        output_dir=tmp_path,
        mode="dry-run",
        year=2026,
        month="May",
        source_statement=tmp_path / "statement.pdf",
        tracker_path=tmp_path / "tracker.xlsx",
        categorized=categorized,
        updates=updates,
    )

    report = report_path.read_text(encoding="utf-8")
    assert "Refund matched expense keyword rule" in report
    assert "Expense claim keyword rule match." in report


def test_report_and_audit_include_planned_period_creation(tmp_path):
    report_path, audit_path, _, _, _ = write_outputs(
        output_dir=tmp_path,
        mode="dry-run",
        year=2026,
        month="Apr",
        source_statement=tmp_path / "statement.pdf",
        tracker_path=tmp_path / "tracker.xlsx",
        categorized=[],
        updates=[],
        structure_changes=[
            WorkbookStructureChange(
                change_type="create_period",
                target_year=2026,
                target_months=("Apr",),
                source_range="C:C",
                target_range="D:D",
                write_action="write",
                reason="Create missing Apr 2026 period from C:C.",
            )
        ],
    )

    report = report_path.read_text(encoding="utf-8")
    assert "## Planned Structure Changes" in report
    assert "create_period: Apr 2026, C:C -> D:D (write;" in report

    audit = audit_path.read_text(encoding="utf-8")
    assert '"record_type": "workbook_structure_change"' in audit
    assert '"change_type": "create_period"' in audit
    assert '"target_months": ["Apr"]' in audit


def test_report_and_audit_include_statement_parser(tmp_path):
    categorized = [
        _categorized(
            "tx-food",
            "FOETEX SCANNGO",
            "-123.45",
            "Food& Drinks (monthly)",
            "Keyword rule match.",
        )
    ]

    report_path, audit_path, _, _, _ = write_outputs(
        output_dir=tmp_path,
        mode="dry-run",
        year=2026,
        month="Apr",
        source_statement=tmp_path / "statement.csv",
        statement_parser="nordea-csv",
        tracker_path=tmp_path / "tracker.xlsx",
        categorized=categorized,
        updates=[],
        structure_changes=[
            WorkbookStructureChange(
                change_type="create_period",
                target_year=2026,
                target_months=("Apr",),
                source_range="C:C",
                target_range="D:D",
                write_action="write",
                reason="Create missing Apr 2026 period from C:C.",
            )
        ],
    )

    report = report_path.read_text(encoding="utf-8")
    assert "- Statement parser: nordea-csv" in report

    audit_lines = audit_path.read_text(encoding="utf-8").splitlines()
    assert '"statement_parser": "nordea-csv"' in audit_lines[0]
    assert '"statement_parser": "nordea-csv"' in audit_lines[1]


def test_review_workbook_contains_review_queue_audit_options_dropdowns_and_metadata(tmp_path):
    tracker = tmp_path / "tracker.xlsx"
    _create_review_tracker(tracker)
    categorized = [
        _categorized(
            "tx-review",
            "UNKNOWN SHOP",
            "-42.50",
            "",
            "No historical or keyword rule matched.",
            authority=Authority.review,
            authority_reason="No category suggested.",
            method="unmatched",
            confidence=0.0,
            merchant="UNKNOWN SHOP",
        ),
        _categorized(
            "tx-food",
            "NETTO",
            "-125.00",
            "Food& Drinks (monthly)",
            "Keyword rule match.",
            authority=Authority.auto,
            method="rule",
            confidence=0.95,
            merchant="NETTO",
        ),
    ]

    _, _, _, _, review_xlsx_path = write_outputs(
        output_dir=tmp_path,
        mode="dry-run",
        year=2026,
        month="May",
        source_statement=tmp_path / "statement.csv",
        statement_parser="nordea-csv",
        tracker_path=tracker,
        categorized=categorized,
        updates=[],
        workbook_config=_config(fixed_rows=frozenset({"Fixed Category"})),
    )

    assert review_xlsx_path == tmp_path / "review_required_2026_may.xlsx"

    workbook = load_workbook(review_xlsx_path)
    try:
        assert workbook.sheetnames == [
            "Review Required",
            "Audit",
            "Category Options",
            "Run Metadata",
        ]
        assert [sheet.sheet_state for sheet in workbook.worksheets] == [
            "visible",
            "visible",
            "hidden",
            "hidden",
        ]

        review_sheet = workbook["Review Required"]
        assert [cell.value for cell in review_sheet[1]] == [
            "date",
            "description",
            "amount",
            "suggested_category",
            "manual_category",
            "reason",
            "confidence",
            "new_parent_category",
            "new_leaf_category",
            "learn_to_memory",
            "blocked",
            "transaction_id",
        ]
        assert _hidden_columns(review_sheet) == {"L"}
        assert [cell.value for cell in review_sheet[2]] == [
            "2026-05-15",
            "UNKNOWN SHOP",
            "-42.50",
            None,
            None,
            "No historical or keyword rule matched. (unmatched)",
            "0.00",
            None,
            None,
            None,
            None,
            "tx-review",
        ]
        note = review_sheet["E1"].comment.text
        assert "Blank accepts suggested_category" in note
        assert "No suggestion: pick a category or NONE" in note

        audit_sheet = workbook["Audit"]
        assert [cell.value for cell in audit_sheet[1]] == [
            "transaction_id",
            "date",
            "description",
            "amount",
            "currency",
            "direction",
            "merchant_identity",
            "category",
            "corrected_category",
            "confidence",
            "source",
            "votes",
            "reason",
            "evidence",
        ]
        assert audit_sheet.max_row == 2
        assert audit_sheet["A2"].value == "tx-food"
        assert audit_sheet["K2"].value == "rule"
        assert any(
            validation.formula1 == "'Category Options'!$G$2:$G$5"
            and validation.allow_blank
            and "I2" in validation.sqref
            for validation in audit_sheet.data_validations.dataValidation
        )

        options_sheet = workbook["Category Options"]
        assert [cell.value for cell in options_sheet[1]] == [
            "row_number",
            "category",
            "learnable",
            "status",
            "category_type",
            "allows_new_children",
            "manual_category_option",
            "new_parent_category_option",
        ]
        options = {
            options_sheet.cell(row=row, column=2).value: (
                options_sheet.cell(row=row, column=1).value,
                options_sheet.cell(row=row, column=3).value,
                options_sheet.cell(row=row, column=4).value,
                options_sheet.cell(row=row, column=5).value,
                options_sheet.cell(row=row, column=6).value,
            )
            for row in range(2, options_sheet.max_row + 1)
        }
        assert list(options) == [
            "Food& Drinks (monthly)",
            "Income (net)",
            "Fixed Category",
            "Manual Category",
        ]
        assert options["Food& Drinks (monthly)"] == (5, True, "write eligible", "leaf", False)
        assert "derived" in options["Income (net)"][2]
        assert options["Income (net)"][1] is False
        assert "fixed" in options["Fixed Category"][2]
        assert options["Fixed Category"][1] is False
        assert "manual value" in options["Manual Category"][2]
        assert options["Manual Category"][1] is True

        validations = tuple(review_sheet.data_validations.dataValidation)
        assert any(
            validation.formula1 == "'Category Options'!$G$2:$G$5"
            and "E2" in validation.sqref
            for validation in validations
        )
        assert not any("G2" in validation.sqref for validation in validations)
        assert any(
            validation.formula1 == '"yes,no"'
            and validation.allow_blank
            and "J2" in validation.sqref
            for validation in validations
        )

        metadata = {
            workbook["Run Metadata"].cell(row=row, column=1).value: workbook[
                "Run Metadata"
            ].cell(row=row, column=2).value
            for row in range(1, workbook["Run Metadata"].max_row + 1)
        }
        assert metadata["reporting_year"] == 2026
        assert metadata["reporting_month"] == "May"
        assert metadata["statement_parser"] == "nordea-csv"
        assert metadata["transaction_id_scheme"] == TRANSACTION_ID_SCHEME_VERSION
        assert metadata["generated_timestamp"]
    finally:
        workbook.close()


def test_exception_sheet_offers_suggestion_alternatives_and_none_as_dropdown(tmp_path):
    tracker = tmp_path / "tracker.xlsx"
    _create_review_tracker(tracker)
    disputed = replace(
        _categorized(
            "tx-disputed",
            "CAFE",
            "-60.00",
            "Food& Drinks (monthly)",
            "Model suggestion.",
            authority=Authority.review,
            method="local_llm_gemma",
        ),
        votes=(
            Vote("Food& Drinks (monthly)", 0.9, "gemma4:26b"),
            Vote("Manual Category", 0.8, "gemma4:12b"),
        ),
    )
    unmatched = _categorized(
        "tx-unmatched", "UNKNOWN", "-5.00", "", "No match.", authority=Authority.review
    )

    _, _, _, _, review_xlsx_path = write_outputs(
        output_dir=tmp_path,
        mode="dry-run",
        year=2026,
        month="May",
        source_statement=tmp_path / "statement.csv",
        tracker_path=tracker,
        categorized=[disputed, unmatched],
        updates=[],
        workbook_config=_config(),
    )

    workbook = load_workbook(review_xlsx_path)
    try:
        review_sheet = workbook["Review Required"]
        assert [review_sheet["L2"].value, review_sheet["L3"].value] == [
            "tx-unmatched",
            "tx-disputed",
        ]
        assert review_sheet["D3"].value == "Food& Drinks (monthly)"
        validations = {
            str(validation.sqref): validation
            for validation in review_sheet.data_validations.dataValidation
        }
        assert validations["E2"].formula1 == "'Category Options'!$G$2:$G$5"
        assert validations["E3"].formula1 == '"Food& Drinks (monthly),Manual Category,NONE"'
        assert workbook["Category Options"]["G2"].value == "NONE"
    finally:
        workbook.close()


def test_exception_sheet_lists_rows_needing_a_decision_first_then_by_date(tmp_path):
    def review_row(transaction_id, category, day):
        item = _categorized(
            transaction_id, "SHOP", "-5.00", category, "Model.", authority=Authority.review
        )
        return replace(item, transaction=replace(item.transaction, date=date(2026, 5, day)))

    _, _, _, _, review_xlsx_path = write_outputs(
        output_dir=tmp_path,
        mode="dry-run",
        year=2026,
        month="May",
        source_statement=tmp_path / "statement.csv",
        tracker_path=tmp_path / "tracker.xlsx",
        categorized=[
            review_row("tx-suggested-late", "Traveling", 20),
            review_row("tx-suggested-early", "Traveling", 1),
            review_row("tx-open-late", "", 25),
            review_row("tx-open-early", "", 3),
        ],
        updates=[],
    )

    workbook = load_workbook(review_xlsx_path)
    try:
        rows = _sheet_rows(workbook["Review Required"])
    finally:
        workbook.close()
    assert [row["transaction_id"] for row in rows] == [
        "tx-open-early",
        "tx-open-late",
        "tx-suggested-early",
        "tx-suggested-late",
    ]


def test_exception_sheet_offers_alternatives_when_suggester_answered_none(tmp_path):
    tracker = tmp_path / "tracker.xlsx"
    _create_review_tracker(tracker)
    none_answer = replace(
        _categorized(
            "tx-none",
            "CAFE",
            "-60.00",
            "",
            "Local model answered NONE.",
            authority=Authority.review,
            method="local_llm_gemma_no_suggestion",
        ),
        votes=(Vote(None, 0.4, "fake"),),
        alternatives=("Manual Category",),
    )

    _, _, _, _, review_xlsx_path = write_outputs(
        output_dir=tmp_path,
        mode="dry-run",
        year=2026,
        month="May",
        source_statement=tmp_path / "statement.csv",
        tracker_path=tracker,
        categorized=[none_answer],
        updates=[],
        workbook_config=_config(),
    )

    workbook = load_workbook(review_xlsx_path)
    try:
        (validation,) = workbook["Review Required"].data_validations.dataValidation[:1]
        assert "E2" in validation.sqref
        assert validation.formula1 == '"Manual Category,NONE"'
    finally:
        workbook.close()


def test_exception_sheet_falls_back_to_leaf_list_when_choices_cannot_be_inlined(tmp_path):
    tracker = tmp_path / "tracker.xlsx"
    _create_review_tracker(tracker)
    comma_label = _categorized(
        "tx-comma",
        "SHOP",
        "-9.00",
        "Books, Games",
        "Model suggestion.",
        authority=Authority.review,
        method="local_llm_gemma",
    )

    _, _, _, _, review_xlsx_path = write_outputs(
        output_dir=tmp_path,
        mode="dry-run",
        year=2026,
        month="May",
        source_statement=tmp_path / "statement.csv",
        tracker_path=tracker,
        categorized=[comma_label],
        updates=[],
        workbook_config=_config(),
    )

    workbook = load_workbook(review_xlsx_path)
    try:
        (validation,) = [
            validation
            for validation in workbook["Review Required"].data_validations.dataValidation
            if "E2" in validation.sqref
        ]
        assert validation.formula1 == "'Category Options'!$G$2:$G$5"
    finally:
        workbook.close()


def test_review_workbook_uses_registry_leaf_and_parent_dropdown_sources(tmp_path):
    tracker = tmp_path / "tracker.xlsx"
    _create_registry_review_tracker(tracker)

    categorized = [
        _categorized(
            "tx-review",
            "UNKNOWN SHOP",
            "-42.50",
            None,
            "No historical or keyword rule matched.",
            authority=Authority.review,
            method="unmatched",
            confidence=0.0,
        )
    ]

    _, _, _, _, review_xlsx_path = write_outputs(
        output_dir=tmp_path,
        mode="dry-run",
        year=2026,
        month="May",
        source_statement=tmp_path / "statement.csv",
        statement_parser="nordea-csv",
        tracker_path=tracker,
        categorized=categorized,
        updates=[],
        workbook_config=_config(
            category_registry=CategoryRegistry(
                leaf_categories=("Food& Drinks (monthly)", "Apple Cloud"),
                parent_categories=("Living expenses", "Services", "Income (net)"),
                new_leaf_parent_categories=("Living expenses", "Services"),
                children_by_parent={
                    "Living expenses": ("Food& Drinks (monthly)",),
                    "Services": ("Apple Cloud",),
                },
                category_type_by_label={
                    "Living expenses": "parent",
                    "Food& Drinks (monthly)": "leaf",
                    "Services": "parent",
                    "Apple Cloud": "leaf",
                    "Income (net)": "derived",
                },
            )
        ),
    )

    workbook = load_workbook(review_xlsx_path)
    try:
        review_sheet = workbook["Review Required"]
        review_headers = {cell.value: index for index, cell in enumerate(review_sheet[1], start=1)}
        assert "new_parent_category" in review_headers
        assert "new_leaf_category" in review_headers

        options_sheet = workbook["Category Options"]
        option_headers = {
            cell.value: index for index, cell in enumerate(options_sheet[1], start=1)
        }
        assert "category_type" in option_headers
        assert "allows_new_children" in option_headers
        assert "manual_category_option" in option_headers
        assert "new_parent_category_option" in option_headers

        context = {
            options_sheet.cell(row=row, column=2).value: {
                "category_type": options_sheet.cell(
                    row=row, column=option_headers["category_type"]
                ).value,
                "allows_new_children": options_sheet.cell(
                    row=row, column=option_headers["allows_new_children"]
                ).value,
            }
            for row in range(2, options_sheet.max_row + 1)
            if options_sheet.cell(row=row, column=2).value
        }
        assert context["Living expenses"] == {
            "category_type": "parent",
            "allows_new_children": True,
        }
        assert context["Food& Drinks (monthly)"]["category_type"] == "leaf"
        assert context["Food& Drinks (monthly)"]["allows_new_children"] is False
        assert context["Income (net)"]["category_type"] == "derived"
        assert context["Income (net)"]["allows_new_children"] is False

        manual_options = [
            options_sheet.cell(row=row, column=option_headers["manual_category_option"]).value
            for row in range(2, options_sheet.max_row + 1)
            if options_sheet.cell(row=row, column=option_headers["manual_category_option"]).value
        ]
        parent_options = [
            options_sheet.cell(row=row, column=option_headers["new_parent_category_option"]).value
            for row in range(2, options_sheet.max_row + 1)
            if options_sheet.cell(row=row, column=option_headers["new_parent_category_option"]).value
        ]
        assert manual_options == ["NONE", "Food& Drinks (monthly)", "Apple Cloud"]
        assert parent_options == ["Living expenses", "Services"]

        validations = tuple(review_sheet.data_validations.dataValidation)
        assert any(
            validation.formula1 == "'Category Options'!$G$2:$G$4"
            and "E2" in validation.sqref
            for validation in validations
        )
        assert any(
            validation.formula1 == "'Category Options'!$H$2:$H$3"
            and "H2" in validation.sqref
            for validation in validations
        )
        assert not any("G2" in validation.sqref for validation in validations)
    finally:
        workbook.close()


def test_review_workbook_exposes_proxy_split_metadata_and_leaves_normal_rows_blank(tmp_path):
    categorized = [
        _categorized(
            "tx-normal",
            "UNKNOWN SHOP",
            "-42.50",
            "",
            "No historical or keyword rule matched.",
            authority=Authority.review,
            method="unmatched",
            confidence=0.0,
        ),
        _categorized(
            "tx-source:split:revolut_family_transfer:residual",
            "REVOLUT [revolut_family_transfer:residual]",
            "-2160.00",
            "",
            "Proxy split residual needs current-month review.",
            authority=Authority.review,
            method="proxy_split_residual",
            confidence=0.0,
            source_transaction_id="tx-source",
            split_rule="revolut_family_transfer",
            split_role="residual",
            residual_amount=Decimal("2160.00"),
        ),
        _categorized(
            "tx-source:split:revolut_family_transfer:dad",
            "REVOLUT [revolut_family_transfer:dad]",
            "-6560.00",
            "Dad",
            "Proxy split allocation.",
            method="proxy_split_allocation",
            source_transaction_id="tx-source",
            split_rule="revolut_family_transfer",
            split_role="dad",
            allocated_amount=Decimal("6560.00"),
        ),
    ]

    _, _, _, _, review_xlsx_path = write_outputs(
        output_dir=tmp_path,
        mode="dry-run",
        year=2026,
        month="May",
        source_statement=tmp_path / "statement.csv",
        tracker_path=tmp_path / "tracker.xlsx",
        categorized=categorized,
        updates=[],
    )

    workbook = load_workbook(review_xlsx_path, data_only=True)
    try:
        review_sheet = workbook["Review Required"]
        headers = {cell.value: index for index, cell in enumerate(review_sheet[1], start=1)}
        assert list(headers)[-6:] == [
            "split_role",
            "split_rule",
            "source_transaction_id",
            "allocated_amount",
            "residual_amount",
            "transaction_id",
        ]
        normal = {
            header: review_sheet.cell(row=2, column=column).value
            for header, column in headers.items()
        }
        residual = {
            header: review_sheet.cell(row=3, column=column).value
            for header, column in headers.items()
        }
        assert normal["split_role"] is None
        assert normal["split_rule"] is None
        assert normal["source_transaction_id"] is None
        assert normal["allocated_amount"] is None
        assert normal["residual_amount"] is None
        assert residual["split_role"] == "residual"
        assert residual["split_rule"] == "revolut_family_transfer"
        assert residual["source_transaction_id"] == "tx-source"
        assert residual["residual_amount"] == "2160.00"
        (allocation,) = _sheet_rows(workbook["Audit"])
        assert allocation["split_role"] == "dad"
        assert allocation["allocated_amount"] == "6560.00"
    finally:
        workbook.close()


def test_exception_sheet_lists_blocked_workbook_updates_but_not_skipped_rows(tmp_path):
    categorized = [
        _categorized(
            "tx-manual-cell",
            "KNOWN SHOP",
            "-25.00",
            "Manual Category",
            "Keyword rule match.",
            authority=Authority.auto,
            method="rule",
            confidence=0.95,
        ),
        _categorized(
            "tx-fixed",
            "FIXED SHOP",
            "-30.00",
            "Fixed Category",
            "Keyword rule match.",
            authority=Authority.auto,
            method="rule",
            confidence=0.95,
        ),
        _categorized(
            "tx-write",
            "NETTO",
            "-40.00",
            "Food& Drinks (monthly)",
            "Keyword rule match.",
            authority=Authority.auto,
            method="rule",
            confidence=0.95,
        ),        _categorized(
            "tx-review-writable",
            "CAFE",
            "-12.00",
            "Food& Drinks (monthly)",
            "Model suggestion.",
            authority=Authority.review,
            method="local_llm_gemma",
        ),
    ]
    updates = [
        _update(
            "Manual Category",
            Decimal("25.00"),
            ("tx-manual-cell",),
            "D8",
            write_action="review",
            reason="Target cell already contains a manual value.",
        ),
        _update(
            "Fixed Category",
            Decimal("30.00"),
            ("tx-fixed",),
            "D7",
            write_action="skip",
            reason="Fixed row is protected by config.",
        ),
        _update(
            "Food& Drinks (monthly)",
            Decimal("40.00"),
            ("tx-write", "tx-review-writable"),
            "D5",
        ),
    ]

    _, _, _, _, review_xlsx_path = write_outputs(
        output_dir=tmp_path,
        mode="dry-run",
        year=2026,
        month="May",
        source_statement=tmp_path / "statement.csv",
        tracker_path=tmp_path / "tracker.xlsx",
        categorized=categorized,
        updates=updates,
        statement_parser="nordea-csv",
    )

    workbook = load_workbook(review_xlsx_path)
    try:
        review_sheet = workbook["Review Required"]
        headers = {cell.value: index for index, cell in enumerate(review_sheet[1], start=1)}
        rows = {
            review_sheet.cell(row=row, column=headers["transaction_id"]).value: {
                header: review_sheet.cell(row=row, column=column).value
                for header, column in headers.items()
            }
            for row in range(2, review_sheet.max_row + 1)
        }

        assert set(rows) == {"tx-manual-cell", "tx-review-writable"}
        assert rows["tx-manual-cell"]["blocked"] == "D8: Target cell already contains a manual value."
        assert rows["tx-review-writable"]["blocked"] is None
    finally:
        workbook.close()


def test_report_includes_categorization_quality_diagnostics(tmp_path):
    categorized = [
        _categorized(
            "tx-food",
            "NETTO",
            "-125.00",
            "Food& Drinks (monthly)",
            "Keyword rule match.",
            authority=Authority.auto,
            method="rule",
        ),
        _categorized(
            "tx-recurring-review",
            "GYM",
            "-250.00",
            "Fitness",
            "Recurring amount/date rule match.",
            authority=Authority.review,
            method="recurring",
            confidence=0.75,
        ),
        _categorized(
            "tx-unmatched",
            "UNKNOWN SHOP",
            "-42.50",
            "",
            "No historical or keyword rule matched.",
            authority=Authority.review,
            method="unmatched",
            confidence=0.0,
        ),
    ]

    report_path, _, _, _, _ = write_outputs(
        output_dir=tmp_path,
        mode="dry-run",
        year=2026,
        month="May",
        source_statement=tmp_path / "statement.csv",
        tracker_path=tmp_path / "tracker.xlsx",
        categorized=categorized,
        updates=[],
        statement_parser="nordea-csv",
    )

    report = report_path.read_text(encoding="utf-8")
    assert "## Categorization Quality" in report
    assert "- Classification rate: 2/3 (66.7%)" in report
    assert "- No-review rate: 1/3 (33.3%)" in report
    assert "- Unmatched transactions: 1" in report
    assert "- Review-required transactions: 2" in report
    assert "- rule: 1" in report
    assert "- recurring: 1" in report
    assert "- unmatched: 1" in report


def _categorized(
    transaction_id: str,
    description: str,
    amount: str,
    category: str,
    reason: str,
    authority: Authority = Authority.auto,
    authority_reason: str = "",
    method: str = "rule",
    confidence: float = 0.95,
    merchant: str | None = None,
    source_transaction_id: str | None = None,
    split_rule: str | None = None,
    split_role: str | None = None,
    allocated_amount: Decimal | None = None,
    residual_amount: Decimal | None = None,
) -> CategorizedTransaction:
    amount_value = Decimal(amount)
    return CategorizedTransaction(
        transaction=Transaction(
            transaction_id=transaction_id,
            date=date(2026, 5, 15),
            interest_date=None,
            description=description,
            amount=amount_value,
            currency="DKK",
            direction="income" if amount_value >= Decimal("0") else "expense",
            merchant=merchant,
        ),
        suggested_category=category or None,
        confidence=confidence,
        categorization_method=method,
        authority=authority,
        authority_reason=authority_reason,
        reason=reason,
        source_transaction_id=source_transaction_id,
        split_rule=split_rule,
        split_role=split_role,
        allocated_amount=allocated_amount,
        residual_amount=residual_amount,
    )


def _sheet_rows(sheet) -> list[dict[str, object]]:
    headers = [cell.value for cell in sheet[1]]
    return [dict(zip(headers, row)) for row in sheet.iter_rows(min_row=2, values_only=True)]


def _hidden_columns(sheet) -> set[str]:
    return {letter for letter, dimension in sheet.column_dimensions.items() if dimension.hidden}


def _update(
    category: str,
    amount: Decimal,
    source_transactions: tuple[str, ...],
    target_cell: str,
    write_action: str = "write",
    reason: str = "Eligible for commit mode write.",
) -> TrackerUpdate:
    return TrackerUpdate(
        year=2026,
        month="May",
        category=category,
        amount=amount,
        source_transactions=source_transactions,
        target_row=5,
        target_column=4,
        target_cell=target_cell,
        existing_value=None,
        write_action=write_action,
        reason=reason,
    )


def _create_review_tracker(path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Net worth"
    sheet["C2"] = 2026
    sheet["C3"] = "May"
    sheet["B5"] = "Food& Drinks (monthly)"
    sheet["B6"] = "Income (net)"
    sheet["C6"] = "=C5"
    sheet["B7"] = "Fixed Category"
    sheet["B8"] = "Manual Category"
    sheet["C8"] = 100
    workbook.save(path)
    workbook.close()


def _create_registry_review_tracker(path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Net worth"
    sheet["C2"] = 2026
    sheet["C3"] = "May"
    sheet["B5"] = "Living expenses"
    sheet["B6"] = "Food& Drinks (monthly)"
    sheet["B7"] = "Services"
    sheet["B8"] = "Apple Cloud"
    sheet["B9"] = "Income (net)"
    workbook.save(path)
    workbook.close()


def _config(
    fixed_rows: frozenset[str] = frozenset(),
    category_registry: CategoryRegistry = CategoryRegistry(),
) -> AppConfig:
    return AppConfig(
        sheet_name="Net worth",
        tracker_currency="DKK",
        category_column=2,
        year_header_row=2,
        month_header_row=3,
        statement_currency="DKK",
        review_threshold=0.60,
        reject_threshold=0.60,
        overwrite_fixed_rows=False,
        highlight_auto_filled_cells=False,
        categories=(),
        aliases={},
        historical_mappings={},
        rules=(),
        fixed_rows=fixed_rows,
        category_registry=category_registry,
    )
