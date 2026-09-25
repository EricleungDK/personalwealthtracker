from decimal import Decimal

import pytest
from openpyxl import load_workbook

from personal_wealth_tracker.config import AppConfig, CategoryRegistry
from personal_wealth_tracker.cli import main
from personal_wealth_tracker.models import Authority, CategorizedTransaction, Transaction
from personal_wealth_tracker.template_workbook import (
    TEMPLATE_METADATA_SHEET,
    create_template_workbook,
    customize_template_workbook,
    validate_template_workbook,
)
from personal_wealth_tracker.workbook import commit_updates, plan_workbook_changes


def test_customize_template_workbook_renames_known_labels_and_currency(tmp_path):
    source = create_template_workbook(tmp_path / "template.xlsx")
    customized = customize_template_workbook(
        source,
        tmp_path / "customized.xlsx",
        tracker_currency="EUR",
        label_renames={
            "Living expenses": "Household costs",
            "Groceries (monthly)": "Groceries",
        },
    )

    metadata = validate_template_workbook(customized)
    assert metadata.tracker_currency == "EUR"

    workbook = load_workbook(customized, data_only=False)
    try:
        sheet = workbook["Net worth"]
        labels = {
            sheet.cell(row=row, column=2).value
            for row in range(1, sheet.max_row + 1)
            if sheet.cell(row=row, column=2).value
        }
        assert "Household costs" in labels
        assert "Groceries" in labels
        assert "Living expenses" not in labels
        assert "Groceries (monthly)" not in labels
        assert sheet["D1"].value == "EUR"
        assert workbook[TEMPLATE_METADATA_SHEET]["B5"].value == "EUR"
        assert sheet["F5"].value == "=SUM(F6:F7)"
        assert sheet["F9"].value == "=SUM(F10:F12)"
    finally:
        workbook.close()

    plan = plan_workbook_changes(
        customized,
        [_categorized("tx-groceries", "Groceries", "-25.00")],
        2026,
        "Apr",
        _customized_config(),
    )

    assert plan.structure_changes == []
    assert plan.updates[0].target_cell == "F11"
    assert plan.updates[0].write_action == "write"


def test_template_workbook_cli_customizes_supported_template(tmp_path, capsys):
    source = create_template_workbook(tmp_path / "template.xlsx")
    output = tmp_path / "customized.xlsx"

    exit_code = main(
        [
            "template-workbook",
            "customize",
            "--template",
            str(source),
            "--output",
            str(output),
            "--tracker-currency",
            "EUR",
            "--rename",
            "Groceries (monthly)=Groceries",
        ]
    )

    captured = capsys.readouterr()
    assert exit_code == 0
    assert f"Customized template workbook: {output}" in captured.out
    workbook = load_workbook(output, data_only=False)
    try:
        assert workbook["Net worth"]["B11"].value == "Groceries"
        assert workbook[TEMPLATE_METADATA_SHEET]["B5"].value == "EUR"
    finally:
        workbook.close()


def test_customize_template_workbook_rejects_unsupported_formula_edits(tmp_path):
    source = create_template_workbook(tmp_path / "template.xlsx")
    workbook = load_workbook(source)
    try:
        workbook["Net worth"]["F9"] = "=F10+F11+F12"
        workbook.save(source)
    finally:
        workbook.close()

    with pytest.raises(ValueError, match="Unsupported template workbook edit"):
        customize_template_workbook(
            source,
            tmp_path / "customized.xlsx",
            label_renames={"Groceries (monthly)": "Groceries"},
        )


def test_supported_template_period_creation_preserves_formulas_and_clears_values(tmp_path):
    template = create_template_workbook(tmp_path / "template.xlsx")
    workbook = load_workbook(template)
    try:
        sheet = workbook["Net worth"]
        sheet["C11"] = 200
        sheet["C12"] = 80
        workbook.save(template)
    finally:
        workbook.close()

    config = _template_config()
    plan = plan_workbook_changes(
        template,
        [_categorized("tx-groceries", "Groceries (monthly)", "-42.50")],
        2027,
        "Jan",
        config,
    )

    assert plan.structure_changes[0].change_type == "create_year"
    assert plan.structure_changes[0].write_action == "write"

    output_path = commit_updates(
        template,
        plan.updates,
        config,
        tmp_path / "processed",
        structure_changes=plan.structure_changes,
    )

    copied = load_workbook(output_path, data_only=False)
    try:
        sheet = copied["Net worth"]
        assert sheet["O2"].value == 2027
        assert sheet["O3"].value == "Jan"
        assert sheet["O5"].value == "=SUM(O6:O7)"
        assert sheet["O9"].value == "=SUM(O10:O12)"
        assert sheet["O11"].value == 42.5
        assert sheet["O12"].value is None
        assert sheet["P11"].value is None
    finally:
        copied.close()


def _template_config() -> AppConfig:
    return AppConfig(
        sheet_name="Net worth",
        tracker_currency="DKK",
        category_column=2,
        year_header_row=2,
        month_header_row=3,
        statement_currency="DKK",
        review_threshold=0.60,
        reject_threshold=0.60,
        overwrite_fixed_rows=False,
        highlight_auto_filled_cells=False,
        categories=(),
        aliases={},
        historical_mappings={},
        rules=(),
        fixed_rows=frozenset(),
        category_registry=CategoryRegistry(
            leaf_categories=("Groceries (monthly)",),
            parent_categories=("Living expenses",),
            category_type_by_label={
                "Living expenses": "parent",
                "Groceries (monthly)": "leaf",
            },
        ),
    )


def _customized_config() -> AppConfig:
    config = _template_config()
    return AppConfig(
        sheet_name=config.sheet_name,
        tracker_currency="EUR",
        category_column=config.category_column,
        year_header_row=config.year_header_row,
        month_header_row=config.month_header_row,
        statement_currency="EUR",
        review_threshold=config.review_threshold,
        reject_threshold=config.reject_threshold,
        overwrite_fixed_rows=config.overwrite_fixed_rows,
        highlight_auto_filled_cells=config.highlight_auto_filled_cells,
        categories=(),
        aliases={},
        historical_mappings={},
        rules=(),
        fixed_rows=frozenset(),
        category_registry=CategoryRegistry(
            leaf_categories=("Groceries",),
            parent_categories=("Household costs",),
            category_type_by_label={
                "Household costs": "parent",
                "Groceries": "leaf",
            },
        ),
    )


def _categorized(
    transaction_id: str,
    category: str,
    amount: str,
) -> CategorizedTransaction:
    from datetime import date

    amount_value = Decimal(amount)
    return CategorizedTransaction(
        transaction=Transaction(
            transaction_id=transaction_id,
            date=date(2026, 4, 30),
            interest_date=None,
            description="SYNTHETIC",
            amount=amount_value,
            currency="DKK",
            direction="income" if amount_value >= Decimal("0") else "expense",
        ),
        suggested_category=category,
        confidence=0.9,
        categorization_method="rule",
        authority=Authority.auto,
        reason="Synthetic template transaction",
    )
