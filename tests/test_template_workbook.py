from decimal import Decimal

import pytest

from personal_wealth_tracker.config import AppConfig, CategoryRegistry
from personal_wealth_tracker.models import CategorizedTransaction, Transaction
from personal_wealth_tracker.template_workbook import (
    SUPPORTED_TEMPLATE_VERSION,
    TEMPLATE_METADATA_SHEET,
    TEMPLATE_WORKBOOK_ID,
    create_template_workbook,
    validate_template_workbook,
)
from personal_wealth_tracker.workbook import plan_updates, workbook_category_options


def test_template_workbook_cli_generates_public_template(monkeypatch, tmp_path, capsys):
    import personal_wealth_tracker.cli as cli

    captured = {}

    def fake_create_template_workbook(
        output_path,
        *,
        tracker_currency,
        start_year,
        sheet_name,
    ):
        captured["output_path"] = output_path
        captured["tracker_currency"] = tracker_currency
        captured["start_year"] = start_year
        captured["sheet_name"] = sheet_name
        output_path.write_text("synthetic template placeholder", encoding="utf-8")
        return output_path

    monkeypatch.setattr(cli, "create_template_workbook", fake_create_template_workbook)

    template_path = tmp_path / "local-wealth-tracker-template.xlsx"
    exit_code = cli.main(
        [
            "template-workbook",
            "create",
            "--output",
            str(template_path),
            "--tracker-currency",
            "SEK",
            "--start-year",
            "2027",
            "--sheet-name",
            "Public Template",
        ]
    )

    captured_output = capsys.readouterr()
    assert exit_code == 0
    assert captured == {
        "output_path": template_path,
        "tracker_currency": "SEK",
        "start_year": 2027,
        "sheet_name": "Public Template",
    }
    assert "Template workbook: " in captured_output.out
    assert str(template_path) in captured_output.out


def test_create_template_workbook_writes_versioned_synthetic_workbook(tmp_path):
    from openpyxl import load_workbook

    template_path = create_template_workbook(tmp_path / "local-wealth-tracker-template.xlsx")

    metadata = validate_template_workbook(template_path)
    assert metadata.template_id == TEMPLATE_WORKBOOK_ID
    assert metadata.template_version == SUPPORTED_TEMPLATE_VERSION
    assert metadata.workbook_kind == "synthetic_template"
    assert metadata.tracker_currency == "DKK"

    workbook = load_workbook(template_path, data_only=False)
    try:
        assert workbook.sheetnames == ["Net worth", TEMPLATE_METADATA_SHEET]
        assert workbook[TEMPLATE_METADATA_SHEET].sheet_state == "hidden"

        sheet = workbook["Net worth"]
        assert sheet["B1"].value == "Category"
        assert sheet["C2"].value == 2026
        assert [sheet.cell(row=3, column=column).value for column in range(3, 15)] == [
            "Jan",
            "Feb",
            "Mar",
            "Apr",
            "May",
            "Jun",
            "Jul",
            "Aug",
            "Sep",
            "Oct",
            "Nov",
            "Dec",
        ]
        assert "C2:N2" in {str(merged_range) for merged_range in sheet.merged_cells.ranges}

        labels = {
            sheet.cell(row=row, column=2).value
            for row in range(1, sheet.max_row + 1)
            if sheet.cell(row=row, column=2).value
        }
        assert {
            "Income (net)",
            "Full-time job (net)",
            "Living expenses",
            "Groceries (monthly)",
            "Services",
            "Insurance",
            "Investments",
            "Assets",
            "Total net worth",
        }.issubset(labels)
        assert {"Mom", "Dad", "JEPI", "OXY", "CRYPTO"}.isdisjoint(labels)

        assert sheet["F5"].value == "=SUM(F6:F7)"
        assert sheet["F9"].value == "=SUM(F10:F12)"
        assert sheet["F26"].value == "=F22"
    finally:
        workbook.close()


def test_validate_template_workbook_rejects_missing_metadata(tmp_path):
    from openpyxl import Workbook

    workbook_path = tmp_path / "ordinary.xlsx"
    workbook = Workbook()
    workbook.save(workbook_path)
    workbook.close()

    with pytest.raises(ValueError, match="Template metadata"):
        validate_template_workbook(workbook_path)


def test_validate_template_workbook_rejects_unsupported_version(tmp_path):
    from openpyxl import load_workbook

    template_path = create_template_workbook(tmp_path / "template.xlsx")
    workbook = load_workbook(template_path)
    try:
        metadata = workbook[TEMPLATE_METADATA_SHEET]
        metadata["B3"] = "999.0"
        workbook.save(template_path)
    finally:
        workbook.close()

    with pytest.raises(ValueError, match="Unsupported template workbook version"):
        validate_template_workbook(template_path)


def test_template_workbook_works_with_existing_safety_checks(tmp_path):
    template_path = create_template_workbook(tmp_path / "template.xlsx")
    config = _template_config()

    options = workbook_category_options(template_path, 2026, "Apr", config)
    option_by_label = {option.category: option for option in options}

    assert option_by_label["Living expenses"].category_type == "parent"
    assert "parent/section row" in option_by_label["Living expenses"].status
    assert option_by_label["Groceries (monthly)"].learnable is True
    assert option_by_label["Groceries (monthly)"].status == "write eligible"
    assert option_by_label["Income (net)"].category_type == "derived"

    updates = plan_updates(
        template_path,
        [_categorized("tx-groceries", "Groceries (monthly)", "-52.40")],
        2026,
        "Apr",
        config,
    )

    assert len(updates) == 1
    assert updates[0].category == "Groceries (monthly)"
    assert updates[0].amount == Decimal("52.40")
    assert updates[0].target_cell == "F11"
    assert updates[0].write_action == "write"


def _template_config() -> AppConfig:
    leaf_categories = (
        "Full-time job (net)",
        "Other income",
        "Rent (monthly)",
        "Groceries (monthly)",
        "Transportation",
        "Mobile phone (monthly)",
        "Internet (monthly)",
        "Cloud services",
        "Insurance (monthly)",
        "Brokerage contributions",
        "Pension contributions",
        "Cash savings",
        "Brokerage account",
        "Pension account",
    )
    parent_categories = (
        "Living expenses",
        "Services",
        "Insurance",
        "Investments",
        "Assets",
    )
    return AppConfig(
        sheet_name="Net worth",
        tracker_currency="DKK",
        category_column=2,
        year_header_row=2,
        month_header_row=3,
        statement_currency="DKK",
        auto_write_threshold=0.85,
        review_threshold=0.60,
        reject_threshold=0.60,
        overwrite_fixed_rows=False,
        highlight_auto_filled_cells=False,
        categories=(),
        aliases={},
        historical_mappings={},
        rules=(),
        fixed_rows=frozenset(),
        category_registry=CategoryRegistry(
            leaf_categories=leaf_categories,
            parent_categories=parent_categories,
            category_type_by_label={
                **{category: "leaf" for category in leaf_categories},
                **{category: "parent" for category in parent_categories},
            },
        ),
    )


def _categorized(
    transaction_id: str,
    category: str,
    amount: str,
) -> CategorizedTransaction:
    from datetime import date

    amount_value = Decimal(amount)
    return CategorizedTransaction(
        transaction=Transaction(
            transaction_id=transaction_id,
            date=date(2026, 4, 30),
            interest_date=None,
            description="SYNTHETIC",
            amount=amount_value,
            currency="DKK",
            direction="income" if amount_value >= Decimal("0") else "expense",
        ),
        suggested_category=category,
        confidence=0.9,
        categorization_method="rule",
        review_required=False,
        reason="Synthetic template transaction",
    )
