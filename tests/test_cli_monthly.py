import os
from pathlib import Path
from urllib.error import URLError

from openpyxl import Workbook, load_workbook

from personal_wealth_tracker.cli import main
from personal_wealth_tracker.suggester import FakeSuggester, ScriptedVote

CSV_HEADER = "﻿Booking date;Amount;Balance;Currency;Name;Title;Sender;Recipient;Reconciled\n"


def test_monthly_picks_newest_csv_and_infers_month(tmp_path, monkeypatch, capsys):
    _workspace(tmp_path, monkeypatch)
    older = _statement(
        tmp_path, "older.csv", "2026/03/05;-10,00;990,00;DKK;OLD SHOP;Card purchase;1;2;Yes\n"
    )
    newer = _statement(
        tmp_path, "newer.csv", "2026/04/01;-42,50;957,50;DKK;UNKNOWN SHOP;Card purchase;1;2;Yes\n"
    )
    os.utime(older, (1_000, 1_000))
    os.utime(newer, (2_000, 2_000))
    _use_suggester(monkeypatch, FakeSuggester({}))

    exit_code = main(["monthly", "--dry-run"])

    output = capsys.readouterr().out
    assert exit_code == 0
    assert f"Statement: {Path('data/raw_statements/newer.csv')}" in output
    assert "Month: Apr 2026" in output


def test_monthly_rejects_statement_spanning_several_months(tmp_path, monkeypatch, capsys):
    _workspace(tmp_path, monkeypatch)
    _statement(
        tmp_path,
        "export.csv",
        "2026/03/31;-10,00;990,00;DKK;SHOP ONE;Card purchase;1;2;Yes\n"
        "2026/04/01;-42,50;947,50;DKK;SHOP TWO;Card purchase;1;2;Yes\n",
    )

    exit_code = main(["monthly"])

    assert exit_code == 1
    assert (
        "Statement 'export.csv' must cover exactly one month; found: 2026-03, 2026-04."
        in capsys.readouterr().err
    )


def test_monthly_summary_shows_auto_review_pending_amount_and_next_action(
    tmp_path, monkeypatch, capsys
):
    _workspace(tmp_path, monkeypatch, trust_policy=ONE_VOTE_TRUST_POLICY)
    _statement(
        tmp_path,
        "april.csv",
        "2026/04/01;-25,00;975,00;DKK;APPLE.COM/BILL;Card purchase;1;2;Yes\n"
        "2026/04/02;-42,50;932,50;DKK;UNKNOWN SHOP;Card purchase;1;2;Yes\n"
        "2026/04/03;-2500,00;-1567,50;DKK;BIG SHOP;Card purchase;1;2;Yes\n"
        "2026/04/04;-7,50;-1575,00;DKK;MYSTERY SHOP;Card purchase;1;2;Yes\n",
    )
    _use_suggester(
        monkeypatch,
        FakeSuggester(
            {
                "UNKNOWN SHOP": ScriptedVote("Traveling", 0.9),
                "BIG SHOP": ScriptedVote("Traveling", 0.9),
            }
        ),
    )

    exit_code = main(["monthly"])

    output = capsys.readouterr().out
    assert exit_code == 0
    assert "Auto rows: 2" in output
    assert "Rows in review: 2" in output
    assert "Pending amount: 2507.50 DKK" in output
    assert (
        f"Next action: fill {Path('reports/review_required_2026_apr.xlsx')} "
        "(blank accepts, NONE rejects), then re-run `wealth-tracker monthly`."
    ) in output


def test_monthly_dry_run_previews_exception_and_audit_without_memory_or_workbook_writes(
    tmp_path, monkeypatch, capsys
):
    _workspace(tmp_path, monkeypatch, trust_policy=ONE_VOTE_TRUST_POLICY)
    _statement(
        tmp_path,
        "april.csv",
        "2026/04/01;-42,50;957,50;DKK;UNKNOWN SHOP;Card purchase;1;2;Yes\n"
        "2026/04/02;-2500,00;-1542,50;DKK;BIG SHOP;Card purchase;1;2;Yes\n",
    )
    _use_suggester(
        monkeypatch,
        FakeSuggester(
            {merchant: ScriptedVote("Traveling", 0.9) for merchant in ("UNKNOWN SHOP", "BIG SHOP")}
        ),
    )

    assert main(["monthly", "--dry-run"]) == 0
    capsys.readouterr()
    preview = load_workbook(tmp_path / "reports" / "review_required_2026_apr.xlsx")
    try:
        review_rows = _sheet_rows(preview["Review Required"])
        audit_rows = _sheet_rows(preview["Audit"])
    finally:
        preview.close()
    assert ["BIG SHOP" in row["description"] for row in review_rows] == [True]
    assert ["UNKNOWN SHOP" in row["description"] for row in audit_rows] == [True]
    _save_unchanged(tmp_path / "reports" / "review_required_2026_apr.xlsx")

    assert main(["monthly", "--dry-run"]) == 0

    output = capsys.readouterr().out
    assert "Rows in review: 0" in output
    assert "Next action: re-run `wealth-tracker monthly` without --dry-run to commit." in output
    assert "Output workbook" not in output
    for private_dir in ("data/processed", "data/backups", "data/category_memory"):
        assert not (tmp_path / private_dir).exists()
    tracker = load_workbook(tmp_path / "Net Worth Tracker.xlsx")
    try:
        assert tracker["Net worth"]["C6"].value is None
    finally:
        tracker.close()


def test_monthly_dry_run_previews_new_leaf_request_without_writing_category_registry(
    tmp_path, monkeypatch, capsys
):
    _workspace(tmp_path, monkeypatch)
    _statement(
        tmp_path,
        "april.csv",
        "2026/04/01;-42,50;957,50;DKK;PET SHOP;Card purchase;1;2;Yes\n",
    )
    _use_suggester(monkeypatch, FakeSuggester({}))
    assert main(["monthly", "--dry-run"]) == 0
    exception_sheet = tmp_path / "reports" / "review_required_2026_apr.xlsx"
    workbook = load_workbook(exception_sheet)
    try:
        sheet = workbook["Review Required"]
        headers = {cell.value: index for index, cell in enumerate(sheet[1], start=1)}
        sheet.cell(row=2, column=headers["new_parent_category"]).value = "Living expenses"
        sheet.cell(row=2, column=headers["new_leaf_category"]).value = "Pet Supplies"
        workbook.save(exception_sheet)
    finally:
        workbook.close()
    categories_yaml = tmp_path / "config" / "categories.yaml"
    registry_before = categories_yaml.read_text(encoding="utf-8")
    capsys.readouterr()

    assert main(["monthly", "--dry-run"]) == 0

    assert categories_yaml.read_text(encoding="utf-8") == registry_before
    report = (tmp_path / "reports" / "report_2026_apr.md").read_text(encoding="utf-8")
    assert "Pet Supplies under Living expenses" in report


def test_monthly_without_local_model_completes_with_unmatched_rows_as_exceptions(
    tmp_path, monkeypatch, capsys
):
    _workspace(tmp_path, monkeypatch)
    _statement(
        tmp_path,
        "april.csv",
        "2026/04/01;-25,00;975,00;DKK;APPLE.COM/BILL;Card purchase;1;2;Yes\n"
        "2026/04/02;-42,50;932,50;DKK;UNKNOWN SHOP;Card purchase;1;2;Yes\n",
    )

    def ollama_down(*_args, **_kwargs):
        raise URLError("connection refused")

    monkeypatch.setattr("personal_wealth_tracker.local_llm.urlopen", ollama_down)

    exit_code = main(["monthly"])

    output = capsys.readouterr().out
    assert exit_code == 0
    assert "Auto rows: 1" in output
    assert "Rows in review: 1" in output
    assert "Pending amount: 42.50 DKK" in output
    exception_sheet = load_workbook(tmp_path / "reports" / "review_required_2026_apr.xlsx")
    try:
        (review_row,) = _sheet_rows(exception_sheet["Review Required"])
    finally:
        exception_sheet.close()
    assert "UNKNOWN SHOP" in review_row["description"]


def test_monthly_rerun_with_untouched_exception_sheet_commits_nothing(
    tmp_path, monkeypatch, capsys
):
    _workspace(tmp_path, monkeypatch, trust_policy=ONE_VOTE_TRUST_POLICY)
    _statement(
        tmp_path,
        "april.csv",
        "2026/04/01;-42,50;957,50;DKK;UNKNOWN SHOP;Card purchase;1;2;Yes\n"
        "2026/04/02;-2500,00;-1542,50;DKK;BIG SHOP;Card purchase;1;2;Yes\n",
    )
    _use_suggester(
        monkeypatch,
        FakeSuggester(
            {merchant: ScriptedVote("Traveling", 0.9) for merchant in ("UNKNOWN SHOP", "BIG SHOP")}
        ),
    )
    assert main(["monthly"]) == 0
    assert main(["monthly", "--dry-run"]) == 0
    capsys.readouterr()

    exit_code = main(["monthly"])

    output = capsys.readouterr().out
    assert exit_code == 0
    assert "Rows in review: 1" in output
    assert (
        "Next action: Exception sheet not reviewed yet: "
        f"{Path('reports/review_required_2026_apr.xlsx')}"
    ) in output
    for private_dir in ("data/processed", "data/backups", "data/category_memory"):
        assert not (tmp_path / private_dir).exists()


def test_monthly_rerun_after_saving_blank_exception_sheet_commits_suggestions(
    tmp_path, monkeypatch, capsys
):
    _workspace(tmp_path, monkeypatch, trust_policy=ONE_VOTE_TRUST_POLICY)
    _statement(
        tmp_path,
        "april.csv",
        "2026/04/01;-2500,00;-1500,00;DKK;BIG SHOP;Card purchase;1;2;Yes\n",
    )
    _use_suggester(monkeypatch, FakeSuggester({"BIG SHOP": ScriptedVote("Traveling", 0.9)}))
    assert main(["monthly"]) == 0
    assert main(["monthly"]) == 0
    capsys.readouterr()
    _save_unchanged(tmp_path / "reports" / "review_required_2026_apr.xlsx")

    exit_code = main(["monthly"])

    output = capsys.readouterr().out
    assert exit_code == 0
    assert "Next action: none; month committed." in output
    (copied_path,) = (tmp_path / "data" / "processed").glob("*.xlsx")
    copied = load_workbook(copied_path)
    try:
        assert copied["Net worth"]["C6"].value == 2500
    finally:
        copied.close()


def test_monthly_rerun_after_filling_exception_sheet_commits_month(
    tmp_path, monkeypatch, capsys
):
    _workspace(tmp_path, monkeypatch, trust_policy=ONE_VOTE_TRUST_POLICY)
    _statement(
        tmp_path,
        "april.csv",
        "2026/04/01;-42,50;957,50;DKK;UNKNOWN SHOP;Card purchase;1;2;Yes\n"
        "2026/04/02;-2500,00;-1542,50;DKK;BIG SHOP;Card purchase;1;2;Yes\n"
        "2026/04/03;-7,50;-1550,00;DKK;MYSTERY SHOP;Card purchase;1;2;Yes\n",
    )
    _use_suggester(
        monkeypatch,
        FakeSuggester(
            {merchant: ScriptedVote("Traveling", 0.9) for merchant in ("UNKNOWN SHOP", "BIG SHOP")}
        ),
    )
    assert main(["monthly"]) == 0
    assert "Output workbook" not in capsys.readouterr().out
    assert main(["monthly"]) == 0
    unfilled_rerun = capsys.readouterr().out
    assert "Rows in review: 2" in unfilled_rerun
    assert f"Exception sheet: {Path('reports/review_required_2026_apr.xlsx')}\n" in unfilled_rerun
    exception_sheet = tmp_path / "reports" / "review_required_2026_apr.xlsx"
    workbook = load_workbook(exception_sheet)
    try:
        sheet = workbook["Review Required"]
        headers = {cell.value: index for index, cell in enumerate(sheet[1], start=1)}
        for row in range(2, sheet.max_row + 1):
            if "MYSTERY SHOP" in sheet.cell(row=row, column=headers["description"]).value:
                sheet.cell(row=row, column=headers["manual_category"]).value = "Apple Cloud"
        workbook.save(exception_sheet)
    finally:
        workbook.close()

    exit_code = main(["monthly"])

    output = capsys.readouterr().out
    assert exit_code == 0
    assert "Rows in review: 0" in output
    assert "Next action: none; month committed." in output
    (copied_path,) = (tmp_path / "data" / "processed").glob("*.xlsx")
    copied = load_workbook(copied_path)
    try:
        assert copied["Net worth"]["C5"].value == 7.5
        assert copied["Net worth"]["C6"].value == 2542.5
    finally:
        copied.close()
    assert list((tmp_path / "data" / "backups").glob("*.xlsx"))


def test_monthly_rewrites_exception_sheet_with_decisions_and_remaining_rows_first(
    tmp_path, monkeypatch, capsys
):
    exception_sheet = _partly_filled_exception_sheet(tmp_path, monkeypatch)

    exit_code = main(["monthly"])

    output = capsys.readouterr().out
    assert exit_code == 0
    assert "Rows in review: 1" in output
    assert (
        f"Next action: fill {Path('reports/review_required_2026_apr.xlsx')} "
        "(blank accepts, NONE rejects), then re-run `wealth-tracker monthly`."
    ) in output
    assert not list((tmp_path / "reports").glob("*_after_decisions.xlsx"))
    assert _decision_cells(exception_sheet) == [
        ("ODD SHOP", None, None),
        ("BIG SHOP", "Traveling", None),
        ("MYSTERY SHOP", "Apple Cloud", "no"),
    ]
    workbook = load_workbook(exception_sheet)
    try:
        sheet = workbook["Review Required"]
        headers = {cell.value: index for index, cell in enumerate(sheet[1], start=1)}
        manual_column = sheet.cell(row=1, column=headers["manual_category"]).column_letter
        validated = {
            str(cell_range)
            for validation in sheet.data_validations.dataValidation
            for cell_range in validation.sqref.ranges
        }
    finally:
        workbook.close()
    assert {f"{manual_column}{row}" for row in (2, 3, 4)} <= validated


def test_monthly_commits_after_filling_rest_of_rewritten_exception_sheet(
    tmp_path, monkeypatch, capsys
):
    exception_sheet = _partly_filled_exception_sheet(tmp_path, monkeypatch)
    assert main(["monthly"]) == 0
    _fill(exception_sheet, "ODD SHOP", manual_category="Traveling")
    capsys.readouterr()

    exit_code = main(["monthly"])

    output = capsys.readouterr().out
    assert exit_code == 0
    assert "Next action: none; month committed." in output
    (copied_path,) = (tmp_path / "data" / "processed").glob("*.xlsx")
    copied = load_workbook(copied_path)
    try:
        assert copied["Net worth"]["C5"].value == 7.5
        assert copied["Net worth"]["C6"].value == 4042.5
    finally:
        copied.close()
    memory = "".join(
        path.read_text(encoding="utf-8")
        for path in (tmp_path / "data" / "category_memory").rglob("*")
        if path.is_file()
    )
    assert "ODD" in memory
    assert "MYSTERY" not in memory


def test_monthly_rerun_of_unsaved_rewrite_keeps_carried_decisions_and_accepts_nothing_new(
    tmp_path, monkeypatch, capsys
):
    exception_sheet = _partly_filled_exception_sheet(tmp_path, monkeypatch)
    assert main(["monthly"]) == 0
    _use_suggester(
        monkeypatch,
        FakeSuggester(
            {
                merchant: ScriptedVote("Traveling", 0.9)
                for merchant in ("UNKNOWN SHOP", "BIG SHOP", "ODD SHOP")
            }
        ),
    )
    assert main(["monthly"]) == 0
    capsys.readouterr()

    exit_code = main(["monthly"])

    output = capsys.readouterr().out
    assert exit_code == 0
    assert "Rows in review: 1" in output
    assert (
        "Next action: Exception sheet not reviewed yet: "
        f"{Path('reports/review_required_2026_apr.xlsx')}"
    ) in output
    assert not (tmp_path / "data" / "processed").exists()
    assert _decision_cells(exception_sheet) == [
        ("ODD SHOP", None, None),
        ("BIG SHOP", "Traveling", None),
        ("MYSTERY SHOP", "Apple Cloud", "no"),
    ]
    assert _sheet_value(exception_sheet, "ODD SHOP", "suggested_category") == "Traveling"
    _save_unchanged(exception_sheet)

    assert main(["monthly"]) == 0

    assert "Next action: none; month committed." in capsys.readouterr().out


def test_monthly_stops_on_unreadable_exception_sheet_and_leaves_it_untouched(
    tmp_path, monkeypatch, capsys
):
    exception_sheet = _partly_filled_exception_sheet(tmp_path, monkeypatch)
    exception_sheet.write_bytes(b"not an xlsx workbook")
    capsys.readouterr()

    exit_code = main(["monthly"])

    assert exit_code == 1
    assert f"Cannot read {exception_sheet.name}" in capsys.readouterr().err
    assert exception_sheet.read_bytes() == b"not an xlsx workbook"


def test_monthly_rewrite_keeps_audit_correction_on_audit_sheet(tmp_path, monkeypatch, capsys):
    _workspace(tmp_path, monkeypatch, trust_policy=ONE_VOTE_TRUST_POLICY)
    _statement(
        tmp_path,
        "april.csv",
        "2026/04/01;-42,50;957,50;DKK;UNKNOWN SHOP;Card purchase;1;2;Yes\n",
    )
    _use_suggester(monkeypatch, FakeSuggester({"UNKNOWN SHOP": ScriptedVote("Traveling", 0.9)}))
    assert main(["monthly"]) == 0
    exception_sheet = tmp_path / "reports" / "review_required_2026_apr.xlsx"
    workbook = load_workbook(exception_sheet)
    try:
        sheet = workbook["Audit"]
        headers = {cell.value: index for index, cell in enumerate(sheet[1], start=1)}
        sheet.cell(row=2, column=headers["corrected_category"]).value = "Apple Cloud"
        workbook.save(exception_sheet)
    finally:
        workbook.close()
    assert main(["monthly"]) == 0
    capsys.readouterr()

    assert main(["monthly"]) == 0

    assert "Next action: none; month committed." in capsys.readouterr().out
    rewritten = load_workbook(exception_sheet)
    try:
        assert _sheet_rows(rewritten["Review Required"]) == []
        (audit_row,) = _sheet_rows(rewritten["Audit"])
    finally:
        rewritten.close()
    assert audit_row["corrected_category"] == "Apple Cloud"
    newest = max((tmp_path / "data" / "processed").glob("*.xlsx"), key=os.path.getmtime)
    copied = load_workbook(newest)
    try:
        assert copied["Net worth"]["C5"].value == 42.5
    finally:
        copied.close()


def test_monthly_clearing_a_carried_decision_puts_the_row_back_in_review(
    tmp_path, monkeypatch, capsys
):
    exception_sheet = _partly_filled_exception_sheet(tmp_path, monkeypatch)
    assert main(["monthly"]) == 0
    _fill(exception_sheet, "BIG SHOP", manual_category=None)
    capsys.readouterr()

    assert main(["monthly"]) == 0

    assert "Rows in review: 2" in capsys.readouterr().out


def test_monthly_stops_before_commit_when_exception_sheet_cannot_be_rewritten(
    tmp_path, monkeypatch, capsys
):
    exception_sheet = _partly_filled_exception_sheet(tmp_path, monkeypatch)
    _fill(exception_sheet, "ODD SHOP", manual_category="Traveling")
    before = exception_sheet.read_bytes()
    exception_sheet.chmod(0o444)
    capsys.readouterr()
    try:
        exit_code = main(["monthly"])
    finally:
        exception_sheet.chmod(0o644)

    assert exit_code == 1
    assert f"Cannot write {exception_sheet.name}" in capsys.readouterr().err
    assert exception_sheet.read_bytes() == before
    assert not (tmp_path / "data" / "processed").exists()


ONE_VOTE_TRUST_POLICY = """
trust_policy:
  min_agreement: 1
"""


def _workspace(tmp_path: Path, monkeypatch, trust_policy: str = "") -> None:
    monkeypatch.chdir(tmp_path)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Net worth"
    sheet["C2"] = 2026
    sheet["C3"] = "Apr"
    sheet["B5"] = "Apple Cloud"
    sheet["B6"] = "Traveling"
    workbook.save(tmp_path / "Net Worth Tracker.xlsx")
    workbook.close()
    config_dir = tmp_path / "config"
    config_dir.mkdir()
    _write(
        config_dir / "settings.yaml",
        """
tracker:
  sheet_name: "Net worth"
  currency: "DKK"
  category_column: 2
  year_header_row: 2
  month_header_row: 3
statement:
  currency: "DKK"
confidence_thresholds:
  auto_write: 0.85
  review_required: 0.60
  reject_below: 0.60
writer:
  overwrite_fixed_rows: false
  highlight_auto_filled_cells: false
"""
        + trust_policy,
    )
    _write(
        config_dir / "categories.yaml",
        """
category_registry:
  - label: "Living expenses"
    type: "parent"
    allow_new_children: true
    children:
      - "Apple Cloud"
      - "Traveling"
aliases: {}
""",
    )
    _write(
        config_dir / "rules.yaml",
        """
historical_mappings:
  "APPLE.COM/BILL": "Apple Cloud"
rules: []
fixed_rows: []
""",
    )


def _statement(tmp_path: Path, name: str, rows: str) -> Path:
    path = tmp_path / "data" / "raw_statements" / name
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(CSV_HEADER + rows, encoding="utf-8")
    return path


def _use_suggester(monkeypatch, suggester: FakeSuggester) -> None:
    monkeypatch.setattr(
        "personal_wealth_tracker.pipeline.local_consensus", lambda _settings: suggester
    )


def _partly_filled_exception_sheet(tmp_path: Path, monkeypatch) -> Path:
    """First run lists BIG (suggested), MYSTERY and ODD (no suggestion); MYSTERY is filled."""
    _workspace(tmp_path, monkeypatch, trust_policy=ONE_VOTE_TRUST_POLICY)
    _statement(
        tmp_path,
        "april.csv",
        "2026/04/01;-42,50;957,50;DKK;UNKNOWN SHOP;Card purchase;1;2;Yes\n"
        "2026/04/02;-2500,00;-1542,50;DKK;BIG SHOP;Card purchase;1;2;Yes\n"
        "2026/04/03;-7,50;-1550,00;DKK;MYSTERY SHOP;Card purchase;1;2;Yes\n"
        "2026/04/04;-1500,00;-3050,00;DKK;ODD SHOP;Card purchase;1;2;Yes\n",
    )
    _use_suggester(
        monkeypatch,
        FakeSuggester(
            {merchant: ScriptedVote("Traveling", 0.9) for merchant in ("UNKNOWN SHOP", "BIG SHOP")}
        ),
    )
    assert main(["monthly"]) == 0
    exception_sheet = tmp_path / "reports" / "review_required_2026_apr.xlsx"
    _fill(exception_sheet, "MYSTERY SHOP", manual_category="Apple Cloud", learn_to_memory="no")
    return exception_sheet


def _fill(exception_sheet: Path, merchant: str, **values: str) -> None:
    workbook = load_workbook(exception_sheet)
    try:
        sheet = workbook["Review Required"]
        headers = {cell.value: index for index, cell in enumerate(sheet[1], start=1)}
        for row in range(2, sheet.max_row + 1):
            if merchant in sheet.cell(row=row, column=headers["description"]).value:
                for header, value in values.items():
                    sheet.cell(row=row, column=headers[header]).value = value
        workbook.save(exception_sheet)
    finally:
        workbook.close()


def _sheet_value(exception_sheet: Path, merchant: str, header: str) -> object:
    workbook = load_workbook(exception_sheet)
    try:
        (row,) = (
            row
            for row in _sheet_rows(workbook["Review Required"])
            if merchant in row["description"]
        )
    finally:
        workbook.close()
    return row[header]


def _decision_cells(exception_sheet: Path) -> list[tuple[str, object, object]]:
    workbook = load_workbook(exception_sheet)
    try:
        rows = _sheet_rows(workbook["Review Required"])
    finally:
        workbook.close()
    return [
        (row["description"].split()[0] + " SHOP", row["manual_category"], row["learn_to_memory"])
        for row in rows
    ]


def _save_unchanged(path: Path) -> None:
    workbook = load_workbook(path)
    try:
        workbook.save(path)
    finally:
        workbook.close()


def _sheet_rows(sheet) -> list[dict[str, object]]:
    headers = [cell.value for cell in sheet[1]]
    return [
        dict(zip(headers, (cell.value for cell in row)))
        for row in sheet.iter_rows(min_row=2)
    ]


def _write(path: Path, content: str) -> None:
    path.write_text(content.lstrip(), encoding="utf-8")
