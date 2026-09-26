import hashlib
from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import PatternFill
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from personal_wealth_tracker import row_insertion
from personal_wealth_tracker.config import AppConfig, CategoryRegistry
from personal_wealth_tracker.models import Authority, CategorizedTransaction, Transaction
from personal_wealth_tracker.row_insertion import evaluate_sheet
from personal_wealth_tracker.workbook import commit_updates, plan_workbook_changes

MONTH_COLUMNS = ("C", "D")


def test_plan_places_missing_leaf_at_end_of_parent_sum_section(tmp_path):
    tracker = _create_sectioned_workbook(tmp_path / "tracker.xlsx")

    plan = plan_workbook_changes(
        tracker,
        [
            _categorized("tx-restaurant", "Restaurants", "-100"),
            _categorized("tx-apple", "Apple Cloud", "-25"),
        ],
        2026,
        "Apr",
        _config(),
    )

    [change] = plan.structure_changes
    assert change.change_type == "insert_leaf_category"
    assert change.write_action == "write", change.reason
    assert change.parent_category == "Living expenses"
    assert change.source_range == "18:18"
    assert change.target_range == "19:19"
    updates = {update.category: update for update in plan.updates}
    assert updates["Restaurants"].target_cell == "D19"
    assert updates["Restaurants"].write_action == "write"
    assert updates["Apple Cloud"].target_cell == "D21"


def test_commit_inserts_leaf_with_every_reference_shifted_only_in_copy(tmp_path):
    tracker = _create_sectioned_workbook(tmp_path / "tracker.xlsx")
    original_digest = hashlib.sha256(tracker.read_bytes()).hexdigest()
    config = _config()
    plan = plan_workbook_changes(
        tracker, [_categorized("tx-restaurant", "Restaurants", "-100")], 2026, "Apr", config
    )

    output_path = commit_updates(
        tracker, plan.updates, config, tmp_path / "processed", structure_changes=plan.structure_changes
    )

    assert hashlib.sha256(tracker.read_bytes()).hexdigest() == original_digest
    copied = load_workbook(output_path)
    try:
        sheet = copied["Net worth"]
        assert sheet["B19"].value == "Restaurants"
        assert sheet["D19"].value == 100
        assert sheet["C19"].value is None
        assert sheet["B19"].fill.fgColor.rgb == "00FFFF00"
        assert sheet.row_dimensions[19].outline_level == 2
        assert sheet.row_dimensions[20].outline_level == 1
        assert sheet["B20"].value == "Services"
        for col in MONTH_COLUMNS:
            assert sheet[f"{col}15"].value == f"=SUM({col}16:{col}19)"
            assert sheet[f"{col}14"].value == f"={col}15+{col}20+{col}23"
            assert sheet[f"{col}20"].value == f"=SUM({col}21:{col}22)"
            assert sheet[f"{col}23"].value == f"=SUM({col}24:{col}24)"
            assert sheet[f"{col}25"].value == f"=SUM({col}26:{col}26)"
            assert sheet[f"{col}10"].value == f"={col}12-{col}14-{col}25"
            assert sheet[f"{col}11"].value == f"=SUM($C$13:{col}13)*-1"
            assert sheet[f"{col}5"].value == f"={col}8-{col}10"
        assert sheet["D6"].value == "=IFERROR((D5-C5)/C5,0)"
        assert copied["Summary"]["A1"].value == "='Net worth'!D25"
        assert copied["Summary"]["A2"].value == "='Net worth'!D14"
        assert {str(merged) for merged in sheet.merged_cells.ranges} == {"F20:G20", "F16:F20"}
        assert [str(cf.sqref) for cf in sheet.conditional_formatting] == ["C16:D26"]
        assert [str(dv.sqref) for dv in sheet.data_validations.dataValidation] == ["B16:B26"]
        assert copied.defined_names["InvestApr"].attr_text == "'Net worth'!$D$25"
        assert copied.defined_names["RecurringApr"].attr_text == "'Net worth'!$D$14"
        assert sheet.print_area == "'Net worth'!$A$1:$D$26"
    finally:
        copied.close()


def test_committed_totals_match_original_when_new_row_is_empty(tmp_path):
    tracker = _create_sectioned_workbook(tmp_path / "tracker.xlsx")
    config = _config()
    plan = plan_workbook_changes(
        tracker, [_categorized("tx-restaurant", "Restaurants", "-100")], 2026, "Apr", config
    )
    output_path = commit_updates(
        tracker, plan.updates, config, tmp_path / "processed", structure_changes=plan.structure_changes
    )

    original = load_workbook(tracker)
    copied = load_workbook(output_path)
    try:
        before = evaluate_sheet(original, "Net worth")
        copied["Net worth"]["D19"].value = None
        after = evaluate_sheet(copied, "Net worth")
    finally:
        original.close()
        copied.close()

    assert before[(15, 4)] == 9540  # Living expenses Apr: rent + lunch
    assert before[(10, 4)] == 30000 - (9540 + 400) - 2000
    shifted_before = {(row + 1 if row >= 19 else row, column): value for (row, column), value in before.items()}
    assert after == shifted_before


def test_multiple_leaves_in_same_and_different_sections_stay_consistent(tmp_path):
    tracker = _create_sectioned_workbook(tmp_path / "tracker.xlsx")
    config = _config()
    plan = plan_workbook_changes(
        tracker,
        [
            _categorized("tx-restaurant", "Restaurants", "-100"),
            _categorized("tx-bus", "transportation", "-30"),
            _categorized("tx-doctor", "healthcare", "-200"),
            _categorized("tx-apple", "Apple Cloud", "-25"),
        ],
        2026,
        "Apr",
        config,
    )

    assert [(change.leaf_category, change.target_range, change.write_action) for change in plan.structure_changes] == [
        ("Restaurants", "19:19", "write"),
        ("healthcare", "23:23", "write"),
        ("transportation", "20:20", "write"),
    ]
    assert {update.category: update.target_cell for update in plan.updates} == {
        "Apple Cloud": "D22",
        "Restaurants": "D19",
        "healthcare": "D24",
        "transportation": "D20",
    }

    output_path = commit_updates(
        tracker, plan.updates, config, tmp_path / "processed", structure_changes=plan.structure_changes
    )
    copied = load_workbook(output_path)
    try:
        sheet = copied["Net worth"]
        assert [sheet.cell(row=row, column=2).value for row in range(19, 25)] == [
            "Restaurants",
            "transportation",
            "Services",
            "Apple Cloud",
            "Metro (monthly)",
            "healthcare",
        ]
        assert [sheet.cell(row=row, column=4).value for row in (19, 20, 22, 24)] == [100, 30, 25, 200]
        for col in MONTH_COLUMNS:
            assert sheet[f"{col}15"].value == f"=SUM({col}16:{col}20)"
            assert sheet[f"{col}21"].value == f"=SUM({col}22:{col}24)"
            assert sheet[f"{col}14"].value == f"={col}15+{col}21+{col}25"
            assert sheet[f"{col}10"].value == f"={col}12-{col}14-{col}27"
        values = evaluate_sheet(copied, "Net worth")
        assert values[(15, 4)] == 9540 + 100 + 30
        assert values[(21, 4)] == 400 + 25 + 200
    finally:
        copied.close()


def test_insertion_that_changes_totals_fails_closed(tmp_path, monkeypatch):
    tracker = _create_sectioned_workbook(tmp_path / "tracker.xlsx")
    config = _config()
    # Simulate a reference-shifting defect: formulas keep pointing at the old rows.
    monkeypatch.setattr(row_insertion, "shift_formula_rows", lambda formula, *args, **kwargs: formula)

    plan = plan_workbook_changes(
        tracker, [_categorized("tx-restaurant", "Restaurants", "-100")], 2026, "Apr", config
    )

    [change] = plan.structure_changes
    assert change.write_action == "review"
    assert "safety check failed" in change.reason
    assert "total changed" in change.reason
    [update] = plan.updates
    assert update.write_action == "review"


def test_parent_sum_in_other_month_not_ending_at_section_end_fails_closed(tmp_path):
    tracker = _create_sectioned_workbook(tmp_path / "tracker.xlsx")
    workbook = load_workbook(tracker)
    workbook["Net worth"]["C15"] = "=SUM(C16:C17)"
    workbook.save(tracker)
    workbook.close()

    plan = plan_workbook_changes(
        tracker, [_categorized("tx-restaurant", "Restaurants", "-100")], 2026, "Apr", _config()
    )

    [change] = plan.structure_changes
    assert change.write_action == "review"
    assert "Net worth!C15" in change.reason


def test_absolute_parent_sum_keeps_its_absolute_markers_when_grown(tmp_path):
    tracker = _create_sectioned_workbook(tmp_path / "tracker.xlsx")
    workbook = load_workbook(tracker)
    workbook["Net worth"]["D15"] = "=SUM($D$16:$D$18)"
    workbook.save(tracker)
    workbook.close()
    config = _config()
    plan = plan_workbook_changes(
        tracker, [_categorized("tx-restaurant", "Restaurants", "-100")], 2026, "Apr", config
    )

    output_path = commit_updates(
        tracker, plan.updates, config, tmp_path / "processed", structure_changes=plan.structure_changes
    )

    copied = load_workbook(output_path)
    try:
        assert copied["Net worth"]["D15"].value == "=SUM($D$16:$D$19)"
    finally:
        copied.close()


def test_position_dependent_formula_blocks_insertion(tmp_path):
    tracker = _create_sectioned_workbook(tmp_path / "tracker.xlsx")
    workbook = load_workbook(tracker)
    workbook["Net worth"]["E5"] = '=INDIRECT("D20")'
    workbook.save(tracker)
    workbook.close()

    plan = plan_workbook_changes(
        tracker, [_categorized("tx-restaurant", "Restaurants", "-100")], 2026, "Apr", _config()
    )

    [change] = plan.structure_changes
    assert change.write_action == "review"
    assert "position-dependent function" in change.reason


def _create_sectioned_workbook(path: Path) -> Path:
    """Mirror the real tracker shape: parent SUM sections, sum-of-parents, derived rows,
    and a Living expenses leaf (Traveling) that sits in a different section (Others)."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Net worth"
    sheet["C2"] = 2026
    sheet["C3"] = "Mar"
    sheet["D3"] = "Apr"
    labels = {
        5: "Total net worth",
        6: "Growth",
        8: "Assets",
        9: "Saving Account",
        10: "Cashflow",
        11: "Accrued taxes",
        12: "Income (net)",
        13: "Full-time job (net)",
        14: "Recurring payments",
        15: "Living expenses",
        16: "Rent (monthly)",
        18: "Lunch (monthly)",
        19: "Services",
        20: "Apple Cloud",
        21: "Metro (monthly)",
        22: "Others",
        23: "Traveling",
        24: "Investments",
        25: "Stock investment plan",
    }
    for row, label in labels.items():
        sheet.cell(row=row, column=2).value = label
    for col in MONTH_COLUMNS:
        sheet[f"{col}5"] = f"={col}8-{col}10"
        sheet[f"{col}8"] = f"={col}9"
        sheet[f"{col}10"] = f"={col}12-{col}14-{col}24"
        sheet[f"{col}11"] = f"=SUM($C$13:{col}13)*-1"
        sheet[f"{col}12"] = f"=SUM({col}13:{col}13)"
        sheet[f"{col}14"] = f"={col}15+{col}19+{col}22"
        sheet[f"{col}15"] = f"=SUM({col}16:{col}18)"
        sheet[f"{col}18"] = "=27*20"
        sheet[f"{col}19"] = f"=SUM({col}20:{col}21)"
        sheet[f"{col}22"] = f"=SUM({col}23:{col}23)"
        sheet[f"{col}24"] = f"=SUM({col}25:{col}25)"
        sheet[f"{col}9"] = 1000
        sheet[f"{col}13"] = 30000
        sheet[f"{col}16"] = 9000
        sheet[f"{col}21"] = 400
        sheet[f"{col}25"] = 2000
    sheet["C6"] = 0
    sheet["D6"] = "=IFERROR((D5-C5)/C5,0)"
    sheet["C20"] = 25
    sheet["C23"] = 500
    sheet["B18"].fill = PatternFill(fill_type="solid", fgColor="FFFF00")
    for row in (16, 17, 18, 20, 21, 23, 25):
        sheet.row_dimensions[row].outline_level = 2
    for row in (15, 19, 22, 24):
        sheet.row_dimensions[row].outline_level = 1
    sheet.merge_cells("F19:G19")
    sheet.merge_cells("F16:F19")
    sheet.conditional_formatting.add(
        "C16:D25", CellIsRule(operator="lessThan", formula=["0"], fill=PatternFill(bgColor="FF0000"))
    )
    validation = DataValidation(type="list", formula1='"a,b"')
    validation.add("B16:B25")
    sheet.add_data_validation(validation)
    sheet.print_area = "A1:D25"
    workbook.defined_names["InvestApr"] = DefinedName("InvestApr", attr_text="'Net worth'!$D$24")
    workbook.defined_names["RecurringApr"] = DefinedName("RecurringApr", attr_text="'Net worth'!$D$14")
    summary = workbook.create_sheet("Summary")
    summary["A1"] = "='Net worth'!D24"
    summary["A2"] = "='Net worth'!D14"
    workbook.save(path)
    workbook.close()
    return path


def _config() -> AppConfig:
    children_by_parent = {
        "Living expenses": (
            "Rent (monthly)",
            "Lunch (monthly)",
            "Traveling",
            "Restaurants",
            "transportation",
        ),
        "Services": ("Apple Cloud", "Metro (monthly)", "healthcare"),
        "Investments": ("Stock investment plan",),
    }
    leaves = tuple(leaf for children in children_by_parent.values() for leaf in children)
    category_types = {parent: "parent" for parent in children_by_parent}
    category_types.update({leaf: "leaf" for leaf in leaves})
    return AppConfig(
        sheet_name="Net worth",
        tracker_currency="DKK",
        category_column=2,
        year_header_row=2,
        month_header_row=3,
        statement_currency="DKK",
        review_threshold=0.60,
        reject_threshold=0.60,
        overwrite_fixed_rows=False,
        highlight_auto_filled_cells=False,
        categories=leaves,
        aliases={},
        historical_mappings={},
        rules=(),
        fixed_rows=frozenset(),
        category_registry=CategoryRegistry(
            leaf_categories=leaves,
            parent_categories=tuple(children_by_parent),
            new_leaf_parent_categories=tuple(children_by_parent),
            children_by_parent=children_by_parent,
            category_type_by_label=category_types,
        ),
    )


def _categorized(transaction_id: str, category: str, amount: str) -> CategorizedTransaction:
    amount_value = Decimal(amount)
    return CategorizedTransaction(
        transaction=Transaction(
            transaction_id=transaction_id,
            date=date(2026, 4, 30),
            interest_date=None,
            description="REDACTED",
            amount=amount_value,
            currency="DKK",
            direction="income" if amount_value >= Decimal(0) else "expense",
        ),
        suggested_category=category,
        confidence=0.9,
        categorization_method="rule",
        authority=Authority.auto,
        reason="Test transaction",
    )
