from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from personal_wealth_tracker.config import AppConfig
from personal_wealth_tracker.models import CategorizedTransaction, Transaction
from personal_wealth_tracker.workbook import commit_updates, plan_updates, plan_workbook_changes


def test_plan_updates_aggregates_and_marks_writeable_empty_cells(tmp_path):
    tracker = _workbook_path(tmp_path)
    _create_workbook(tracker)

    updates = plan_updates(
        tracker,
        [
            _categorized("tx-food-1", "Food& Drinks (monthly)", "-10.50"),
            _categorized("tx-food-2", "Food& Drinks (monthly)", "-20.25"),
        ],
        2026,
        "Apr",
        _config(),
    )

    assert len(updates) == 1
    assert updates[0].category == "Food& Drinks (monthly)"
    assert updates[0].amount == Decimal("30.75")
    assert updates[0].target_cell == "C5"
    assert updates[0].write_action == "write"


def test_plan_workbook_changes_reports_missing_month_as_structure_change(tmp_path):
    tracker = _workbook_path(tmp_path)
    _create_workbook(tracker)
    workbook = load_workbook(tracker)
    try:
        sheet = workbook["Net worth"]
        sheet["C3"] = "Mar"
        workbook.save(tracker)
    finally:
        workbook.close()

    plan = plan_workbook_changes(
        tracker,
        [_categorized("tx-food", "Food& Drinks (monthly)", "-30.00")],
        2026,
        "Apr",
        _config(),
    )

    assert len(plan.structure_changes) == 1
    assert plan.structure_changes[0].change_type == "create_period"
    assert plan.structure_changes[0].write_action == "write"
    assert plan.structure_changes[0].target_year == 2026
    assert plan.structure_changes[0].target_months == ("Apr",)
    assert plan.structure_changes[0].source_range == "C:C"
    assert plan.structure_changes[0].target_range == "D:D"

    assert len(plan.updates) == 1
    assert plan.updates[0].target_cell == "D5"
    assert plan.updates[0].write_action == "write"


def test_plan_workbook_changes_creates_full_year_block_from_unambiguous_prior_year(tmp_path):
    tracker = _workbook_path(tmp_path)
    _create_year_workbook(tracker)

    plan = plan_workbook_changes(
        tracker,
        [_categorized("tx-food", "Food& Drinks (monthly)", "-30.00")],
        2027,
        "Jan",
        _config(),
    )

    assert len(plan.structure_changes) == 1
    assert plan.structure_changes[0].change_type == "create_year"
    assert plan.structure_changes[0].write_action == "write"
    assert plan.structure_changes[0].target_year == 2027
    assert plan.structure_changes[0].target_months == (
        "Jan",
        "Feb",
        "Mar",
        "Apr",
        "May",
        "Jun",
        "Jul",
        "Aug",
        "Sep",
        "Oct",
        "Nov",
        "Dec",
    )
    assert plan.structure_changes[0].source_range == "C:N"
    assert plan.structure_changes[0].target_range == "O:Z"
    assert plan.updates[0].target_cell == "O5"


def test_plan_workbook_changes_requires_review_for_ambiguous_missing_year_pattern(tmp_path):
    tracker = _workbook_path(tmp_path)
    _create_year_workbook(tracker)
    workbook = load_workbook(tracker)
    try:
        sheet = workbook["Net worth"]
        sheet["H3"] = "Not Jun"
        workbook.save(tracker)
    finally:
        workbook.close()

    plan = plan_workbook_changes(
        tracker,
        [_categorized("tx-food", "Food& Drinks (monthly)", "-30.00")],
        2027,
        "Jan",
        _config(),
    )

    assert len(plan.structure_changes) == 1
    assert plan.structure_changes[0].write_action == "review"
    assert plan.structure_changes[0].target_range is None
    assert "prior year does not contain an unambiguous Jan-Dec period block" in (
        plan.structure_changes[0].reason
    )
    assert plan.updates[0].target_cell is None
    assert plan.updates[0].write_action == "review"


def test_plan_updates_nets_deterministic_refunds_against_category_total(tmp_path):
    tracker = _workbook_path(tmp_path)
    _create_workbook(tracker)

    updates = plan_updates(
        tracker,
        [
            _categorized("tx-food-purchase", "Food& Drinks (monthly)", "-100.00"),
            _categorized("tx-food-refund", "Food& Drinks (monthly)", "25.00"),
        ],
        2026,
        "Apr",
        _config(),
    )

    assert len(updates) == 1
    assert updates[0].category == "Food& Drinks (monthly)"
    assert updates[0].amount == Decimal("75.00")
    assert updates[0].source_transactions == ("tx-food-purchase", "tx-food-refund")
    assert updates[0].write_action == "write"


def test_later_month_refund_targets_selected_month_without_reopening_prior_period(
    tmp_path,
):
    tracker = _workbook_path(tmp_path)
    _create_workbook(tracker)
    workbook = load_workbook(tracker)
    try:
        sheet = workbook["Net worth"]
        sheet["D2"] = 2026
        sheet["D3"] = "May"
        sheet["C5"] = 100
        workbook.save(tracker)
    finally:
        workbook.close()

    config = _config()
    updates = plan_updates(
        tracker,
        [_categorized("tx-food-refund", "Food& Drinks (monthly)", "25.00")],
        2026,
        "May",
        config,
    )

    assert len(updates) == 1
    assert updates[0].amount == Decimal("-25.00")
    assert updates[0].target_cell == "D5"
    assert updates[0].write_action == "write"

    output_path = commit_updates(tracker, updates, config, tmp_path / "processed")
    copied = load_workbook(output_path)
    try:
        sheet = copied["Net worth"]
        assert sheet["C5"].value == 100
        assert sheet["D5"].value == -25
    finally:
        copied.close()


def test_plan_updates_protects_formula_manual_fixed_and_review_rows(tmp_path):
    tracker = _workbook_path(tmp_path)
    _create_workbook(tracker)

    updates = {
        update.category: update
        for update in plan_updates(
            tracker,
            [
                _categorized("tx-formula", "Formula Category", "-1.00"),
                _categorized("tx-manual", "Manual Category", "-2.00"),
                _categorized("tx-fixed", "Fixed Category", "-3.00"),
                _categorized("tx-review", "Review Category", "-4.00", review_required=True),
                _categorized("tx-missing", "Missing Category", "-5.00"),
            ],
            2026,
            "Apr",
            _config(fixed_rows=frozenset({"Fixed Category"})),
        )
    }

    assert updates["Formula Category"].write_action == "skip"
    assert updates["Formula Category"].reason == "Target cell contains a formula."
    assert updates["Manual Category"].write_action == "review"
    assert updates["Manual Category"].reason == "Target cell already contains a manual value."
    assert updates["Fixed Category"].write_action == "skip"
    assert updates["Fixed Category"].reason == "Fixed row is protected by config."
    assert updates["Review Category"].write_action == "review"
    assert updates["Review Category"].reason == "One or more source transactions require review."
    assert updates["Missing Category"].write_action == "review"
    assert updates["Missing Category"].reason == "Target category row not found."


def test_plan_updates_skips_derived_income_row_even_when_cell_is_empty(tmp_path):
    tracker = _workbook_path(tmp_path)
    _create_workbook(tracker)

    updates = plan_updates(
        tracker,
        [_categorized("tx-income", "Income (net)", "10000.00")],
        2026,
        "Apr",
        _config(),
    )

    assert len(updates) == 1
    assert updates[0].category == "Income (net)"
    assert updates[0].target_cell == "C10"
    assert updates[0].write_action == "skip"
    assert updates[0].reason == "Derived workbook row is formula-owned and not writable."


@pytest.mark.parametrize(
    "section_total",
    [
        "Cashflow",
        "Recurring payments",
        "Living expenses",
        "Services",
        "Insurance",
        "Investments",
        "Assets",
        "Total net worth",
    ],
)
def test_plan_updates_skips_section_total_rows_even_when_cell_is_empty(
    tmp_path, section_total
):
    tracker = _workbook_path(tmp_path)
    _create_workbook(tracker)

    updates = plan_updates(
        tracker,
        [_categorized("tx-section-total", section_total, "100.00")],
        2026,
        "Apr",
        _config(),
    )

    assert len(updates) == 1
    assert updates[0].category == section_total
    assert updates[0].write_action == "skip"
    assert updates[0].reason == "Derived workbook row is formula-owned and not writable."


@pytest.mark.parametrize("payroll_row", ["Labour market contribution", "Taxes"])
def test_plan_updates_skips_payroll_deduction_rows_even_when_cell_is_empty(
    tmp_path, payroll_row
):
    tracker = _workbook_path(tmp_path)
    _create_workbook(tracker)

    updates = plan_updates(
        tracker,
        [_categorized("tx-payroll", payroll_row, "100.00")],
        2026,
        "Apr",
        _config(),
    )

    assert len(updates) == 1
    assert updates[0].category == payroll_row
    assert updates[0].write_action == "skip"
    assert updates[0].reason == "Derived workbook row is formula-owned and not writable."


def test_plan_updates_writes_single_salary_match_to_full_time_job_net(tmp_path):
    tracker = _workbook_path(tmp_path)
    _create_workbook(tracker)

    updates = plan_updates(
        tracker,
        [_categorized("tx-salary", "Full-time job (net)", "10000.00")],
        2026,
        "Apr",
        _config(),
    )

    assert len(updates) == 1
    assert updates[0].category == "Full-time job (net)"
    assert updates[0].amount == Decimal("10000.00")
    assert updates[0].target_cell == "C11"
    assert updates[0].write_action == "write"


def test_plan_updates_requires_review_for_multiple_salary_matches(tmp_path):
    tracker = _workbook_path(tmp_path)
    _create_workbook(tracker)

    updates = plan_updates(
        tracker,
        [
            _categorized("tx-salary-1", "Full-time job (net)", "10000.00"),
            _categorized("tx-salary-2", "Full-time job (net)", "2500.00"),
        ],
        2026,
        "Apr",
        _config(),
    )

    assert len(updates) == 1
    assert updates[0].category == "Full-time job (net)"
    assert updates[0].amount == Decimal("12500.00")
    assert updates[0].target_cell == "C11"
    assert updates[0].write_action == "review"
    assert updates[0].reason == "Multiple salary deposits matched; review before writing."


def test_plan_updates_writes_deterministic_expense_claims_as_existing_row(tmp_path):
    tracker = _workbook_path(tmp_path)
    _create_workbook(tracker)

    updates = plan_updates(
        tracker,
        [_categorized("tx-expense-claim", "Expense claims", "125.00")],
        2026,
        "Apr",
        _config(),
    )

    assert len(updates) == 1
    assert updates[0].category == "Expense claims"
    assert updates[0].amount == Decimal("125.00")
    assert updates[0].target_cell == "C22"
    assert updates[0].write_action == "write"


def test_commit_updates_writes_only_eligible_updates_to_copied_workbook(tmp_path):
    tracker = _workbook_path(tmp_path)
    _create_workbook(tracker)
    config = _config(fixed_rows=frozenset({"Fixed Category"}))
    updates = plan_updates(
        tracker,
        [
            _categorized("tx-food", "Food& Drinks (monthly)", "-30.00"),
            _categorized("tx-formula", "Formula Category", "-1.00"),
            _categorized("tx-manual", "Manual Category", "-2.00"),
            _categorized("tx-fixed", "Fixed Category", "-3.00"),
        ],
        2026,
        "Apr",
        config,
    )

    output_path = commit_updates(tracker, updates, config, tmp_path / "processed")

    original = load_workbook(tracker)
    copied = load_workbook(output_path, data_only=False)
    try:
        original_sheet = original["Net worth"]
        copied_sheet = copied["Net worth"]
        assert original_sheet["C5"].value is None
        assert copied_sheet["C5"].value == 30
        assert copied_sheet["C6"].value == "=1+1"
        assert copied_sheet["C7"].value == 123
        assert copied_sheet["C8"].value is None
    finally:
        original.close()
        copied.close()


def test_commit_updates_creates_missing_month_only_in_copied_workbook(tmp_path):
    tracker = _workbook_path(tmp_path)
    _create_workbook(tracker)
    workbook = load_workbook(tracker)
    try:
        sheet = workbook["Net worth"]
        sheet["C3"] = "Mar"
        workbook.save(tracker)
    finally:
        workbook.close()

    config = _config()
    plan = plan_workbook_changes(
        tracker,
        [_categorized("tx-food", "Food& Drinks (monthly)", "-30.00")],
        2026,
        "Apr",
        config,
    )

    output_path = commit_updates(
        tracker,
        plan.updates,
        config,
        tmp_path / "processed",
        structure_changes=plan.structure_changes,
    )

    original = load_workbook(tracker)
    copied = load_workbook(output_path)
    try:
        assert original["Net worth"]["D3"].value is None
        assert original["Net worth"]["D5"].value is None
        assert copied["Net worth"]["D2"].value == 2026
        assert copied["Net worth"]["D3"].value == "Apr"
        assert copied["Net worth"]["D5"].value == 30
    finally:
        original.close()
        copied.close()


def test_commit_updates_extends_merged_year_header_for_missing_month(tmp_path):
    tracker = _workbook_path(tmp_path)
    _create_workbook(tracker)
    workbook = load_workbook(tracker)
    try:
        sheet = workbook["Net worth"]
        sheet["C3"] = "Mar"
        sheet["D3"] = "Apr"
        sheet.merge_cells("C2:D2")
        sheet["C2"] = 2026
        workbook.save(tracker)
    finally:
        workbook.close()

    config = _config()
    plan = plan_workbook_changes(
        tracker,
        [_categorized("tx-food", "Food& Drinks (monthly)", "-30.00")],
        2026,
        "May",
        config,
    )

    output_path = commit_updates(
        tracker,
        plan.updates,
        config,
        tmp_path / "processed",
        structure_changes=plan.structure_changes,
    )

    copied = load_workbook(output_path)
    try:
        sheet = copied["Net worth"]
        assert "C2:E2" in {str(merged_range) for merged_range in sheet.merged_cells.ranges}
        assert sheet["C2"].value == 2026
        assert sheet["E3"].value == "May"
    finally:
        copied.close()


def test_commit_updates_creates_year_block_with_preserved_structure_and_cleared_values(
    tmp_path,
):
    tracker = _workbook_path(tmp_path)
    _create_year_workbook(tracker)
    workbook = load_workbook(tracker)
    try:
        sheet = workbook["Net worth"]
        sheet.merge_cells("C2:N2")
        sheet["C2"] = 2026
        sheet["B6"] = "Formula Category"
        sheet["C6"] = "=C5+1"
        sheet["B7"] = "Carry Forward Category"
        sheet["C7"] = 123
        sheet["C5"] = 99
        sheet.column_dimensions["C"].width = 18
        sheet["C5"].number_format = "#,##0.00"
        sheet["C5"].fill = PatternFill(fill_type="solid", fgColor="FFFF00")
        workbook.save(tracker)
    finally:
        workbook.close()

    config = _config(carry_forward_rows=frozenset({"Carry Forward Category"}))
    plan = plan_workbook_changes(
        tracker,
        [_categorized("tx-food", "Food& Drinks (monthly)", "-30.00")],
        2027,
        "Jan",
        config,
    )

    output_path = commit_updates(
        tracker,
        plan.updates,
        config,
        tmp_path / "processed",
        structure_changes=plan.structure_changes,
    )

    copied = load_workbook(output_path, data_only=False)
    try:
        sheet = copied["Net worth"]
        assert "O2:Z2" in {str(merged_range) for merged_range in sheet.merged_cells.ranges}
        assert sheet["O2"].value == 2027
        assert sheet["O3"].value == "Jan"
        assert sheet["Z3"].value == "Dec"
        assert sheet.column_dimensions["O"].width == 18
        assert sheet["O5"].number_format == "#,##0.00"
        assert sheet["O5"].fill.fgColor.rgb == "00FFFF00"
        assert sheet["O5"].value == 30
        assert sheet["P5"].value is None
        assert sheet["O6"].value == "=O5+1"
        assert sheet["O7"].value == 123
        assert sheet["P7"].value is None
    finally:
        copied.close()


def _create_workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Net worth"
    sheet["C2"] = 2026
    sheet["C3"] = "Apr"
    sheet["B5"] = "Food& Drinks (monthly)"
    sheet["B6"] = "Formula Category"
    sheet["C6"] = "=1+1"
    sheet["B7"] = "Manual Category"
    sheet["C7"] = 123
    sheet["B8"] = "Fixed Category"
    sheet["B9"] = "Review Category"
    sheet["B10"] = "Income (net)"
    sheet["B11"] = "Full-time job (net)"
    sheet["B12"] = "Cashflow"
    sheet["B13"] = "Recurring payments"
    sheet["B14"] = "Living expenses"
    sheet["B15"] = "Services"
    sheet["B16"] = "Insurance"
    sheet["B17"] = "Investments"
    sheet["B18"] = "Assets"
    sheet["B19"] = "Total net worth"
    sheet["B20"] = "Labour market contribution"
    sheet["B21"] = "Taxes"
    sheet["B22"] = "Expense claims"
    workbook.save(path)
    workbook.close()


def _create_year_workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Net worth"
    for index, month in enumerate(
        ("Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"),
        start=3,
    ):
        sheet.cell(row=2, column=index).value = 2026
        sheet.cell(row=3, column=index).value = month
    sheet["B5"] = "Food& Drinks (monthly)"
    workbook.save(path)
    workbook.close()


def _workbook_path(tmp_path) -> Path:
    return tmp_path / "tracker.xlsx"


def _config(
    fixed_rows: frozenset[str] = frozenset(),
    carry_forward_rows: frozenset[str] = frozenset(),
) -> AppConfig:
    return AppConfig(
        sheet_name="Net worth",
        tracker_currency="DKK",
        category_column=2,
        year_header_row=2,
        month_header_row=3,
        statement_currency="DKK",
        auto_write_threshold=0.85,
        review_threshold=0.60,
        reject_threshold=0.60,
        overwrite_fixed_rows=False,
        highlight_auto_filled_cells=False,
        categories=(),
        aliases={},
        historical_mappings={},
        rules=(),
        fixed_rows=fixed_rows,
        carry_forward_rows=carry_forward_rows,
    )


def _categorized(
    transaction_id: str,
    category: str,
    amount: str,
    review_required: bool = False,
) -> CategorizedTransaction:
    amount_value = Decimal(amount)
    return CategorizedTransaction(
        transaction=Transaction(
            transaction_id=transaction_id,
            date=date(2026, 4, 30),
            interest_date=None,
            description="REDACTED",
            amount=amount_value,
            currency="DKK",
            direction="income" if amount_value >= Decimal("0") else "expense",
        ),
        suggested_category=category,
        confidence=0.9,
        categorization_method="rule",
        review_required=review_required,
        reason="Test transaction",
    )
