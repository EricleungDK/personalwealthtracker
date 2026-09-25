import csv
import json
import sys
from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook

from personal_wealth_tracker.categorizer import categorize_transactions
from personal_wealth_tracker.category_memory import import_reviewed_decisions, load_category_memory
from personal_wealth_tracker.cli import main
from personal_wealth_tracker.config import AppConfig
from personal_wealth_tracker.models import Transaction
from personal_wealth_tracker.pipeline import run_pipeline
from personal_wealth_tracker.utils import TRANSACTION_ID_SCHEME_VERSION


FIXTURE = Path(__file__).parent / "fixtures" / "nordea_account_statement.redacted.pdf"


def test_import_confirmed_review_decision_into_category_memory(tmp_path):
    decisions_path = tmp_path / "reviewed_decisions.csv"
    memory_dir = tmp_path / "data" / "category_memory"
    _write_decisions(
        decisions_path,
        [
            {
                "transaction_id": "tx-netto-1",
                "date": "2026-04-12",
                "description": "NETTO 1234 KOBENHAVN",
                "amount": "-125.50",
                "direction": "expense",
                "confirmed_category": "Food& Drinks (monthly)",
                "confirmed": "yes",
            }
        ],
    )

    exit_code = main(
        [
            "learn-category-memory",
            "--decisions",
            str(decisions_path),
            "--memory-dir",
            str(memory_dir),
        ]
    )

    assert exit_code == 0
    memory_path = memory_dir / "category_memory.json"
    assert memory_path.exists()
    payload = json.loads(memory_path.read_text(encoding="utf-8"))
    assert payload["mappings"] == [
        {
            "merchant_identity": "NETTO KOBENHAVN",
            "category": "Food& Drinks (monthly)",
            "source_transaction_ids": ["tx-netto-1"],
        }
    ]


def test_import_learns_opted_in_review_workbook_decisions(tmp_path, capsys):
    decisions_path = tmp_path / "review_required_2026_apr.xlsx"
    memory_dir = tmp_path / "data" / "category_memory"
    _write_review_workbook_decisions(
        decisions_path,
        review_rows=[
            {
                "transaction_id": "tx-netto-1",
                "description": "NETTO 1234 KOBENHAVN",
                "amount": "-125.50",
                "manual_category": "Food& Drinks (monthly)",
                "learn_to_memory": "yes",
            },
            {
                "transaction_id": "tx-mobilepay-once",
                "description": "MOBILEPAY PRIVATE TRANSFER",
                "amount": "-300.00",
                "manual_category": "Parent B",
                "learn_to_memory": "",
            },
            {
                "transaction_id": "tx-explicit-no",
                "description": "ONE OFF SHOP",
                "amount": "-99.00",
                "manual_category": "Shopping (monthly)",
                "learn_to_memory": "no",
            },
            {
                "transaction_id": "tx-blank-category",
                "description": "UNREVIEWED",
                "amount": "-10.00",
                "manual_category": "",
                "learn_to_memory": "yes",
            },
        ],
        category_options=[
            ("Food& Drinks (monthly)", True),
            ("Parent B", True),
            ("Shopping (monthly)", True),
        ],
    )

    exit_code = main(
        [
            "learn-category-memory",
            "--decisions",
            str(decisions_path),
            "--memory-dir",
            str(memory_dir),
        ]
    )

    captured = capsys.readouterr()
    assert exit_code == 0
    assert "Imported category memory decisions: 1" in captured.out
    assert "Skipped unlearned decisions: 3" in captured.out
    assert "Skipped non-learnable decisions: 0" in captured.out
    payload = json.loads((memory_dir / "category_memory.json").read_text(encoding="utf-8"))
    assert payload["mappings"] == [
        {
            "merchant_identity": "NETTO KOBENHAVN",
            "category": "Food& Drinks (monthly)",
            "source_transaction_ids": ["tx-netto-1"],
        }
    ]


def test_import_review_workbook_skips_non_learnable_fields(tmp_path, capsys):
    decisions_path = tmp_path / "review_required_2026_apr.xlsx"
    memory_dir = tmp_path / "data" / "category_memory"
    _write_review_workbook_decisions(
        decisions_path,
        review_rows=[
            {
                "transaction_id": "tx-income-net",
                "description": "SALARY",
                "amount": "10000.00",
                "manual_category": "Income (net)",
                "learn_to_memory": "yes",
            },
            {
                "transaction_id": "tx-food",
                "description": "NETTO 1234 KOBENHAVN",
                "amount": "-125.50",
                "manual_category": "Food& Drinks (monthly)",
                "learn_to_memory": "yes",
            },
        ],
        category_options=[
            ("Income (net)", False),
            ("Food& Drinks (monthly)", True),
        ],
    )

    exit_code = main(
        [
            "learn-category-memory",
            "--decisions",
            str(decisions_path),
            "--memory-dir",
            str(memory_dir),
        ]
    )

    captured = capsys.readouterr()
    assert exit_code == 0
    assert "Imported category memory decisions: 1" in captured.out
    assert "Skipped non-learnable decisions: 1" in captured.out
    memory = load_category_memory(memory_dir)
    assert [mapping.category for mapping in memory.mappings] == ["Food& Drinks (monthly)"]


def test_import_review_workbook_skips_proxy_split_residual_decisions(tmp_path, capsys):
    decisions_path = tmp_path / "review_required_2026_apr.xlsx"
    memory_dir = tmp_path / "data" / "category_memory"
    _write_review_workbook_decisions(
        decisions_path,
        review_rows=[
            {
                "transaction_id": (
                    f"{TRANSACTION_ID_SCHEME_VERSION}:source:001"
                    ":split:example_transfer:residual"
                ),
                "description": "REVOLUT residual",
                "amount": "-2160.00",
                "manual_category": "Traveling",
                "learn_to_memory": "yes",
            }
        ],
        category_options=[("Traveling", True)],
    )

    exit_code = main(
        [
            "learn-category-memory",
            "--decisions",
            str(decisions_path),
            "--memory-dir",
            str(memory_dir),
        ]
    )

    captured = capsys.readouterr()
    assert exit_code == 0
    assert "Imported category memory decisions: 0" in captured.out
    memory = load_category_memory(memory_dir)
    assert memory.mappings == ()


def test_import_review_workbook_validates_learned_categories_against_registry_leaves(
    tmp_path, capsys
):
    decisions_path = tmp_path / "review_required_2026_apr.xlsx"
    memory_dir = tmp_path / "data" / "category_memory"
    config_dir = tmp_path / "config"
    _create_category_registry_config(config_dir)
    _write_review_workbook_decisions(
        decisions_path,
        review_rows=[
            {
                "transaction_id": "tx-existing-leaf",
                "description": "NETTO 1234 KOBENHAVN",
                "amount": "-125.50",
                "manual_category": "Food& Drinks (monthly)",
                "learn_to_memory": "yes",
            },
            {
                "transaction_id": "tx-new-leaf",
                "description": "PET SHOP",
                "amount": "-42.50",
                "manual_category": "Pet Supplies",
                "learn_to_memory": "yes",
            },
            {
                "transaction_id": "tx-parent",
                "description": "PARENT ROW",
                "amount": "-1.00",
                "manual_category": "Living expenses",
                "learn_to_memory": "yes",
            },
            {
                "transaction_id": "tx-missing",
                "description": "MISSING ROW",
                "amount": "-2.00",
                "manual_category": "Missing Category",
                "learn_to_memory": "yes",
            },
        ],
        category_options=[
            ("Food& Drinks (monthly)", True),
            ("Pet Supplies", True),
            ("Living expenses", True),
            ("Missing Category", True),
        ],
    )

    exit_code = main(
        [
            "learn-category-memory",
            "--decisions",
            str(decisions_path),
            "--memory-dir",
            str(memory_dir),
            "--config-dir",
            str(config_dir),
        ]
    )

    captured = capsys.readouterr()
    assert exit_code == 0
    assert "Imported category memory decisions: 2" in captured.out
    assert "Skipped non-learnable decisions: 2" in captured.out
    memory = load_category_memory(memory_dir)
    assert [mapping.category for mapping in memory.mappings] == [
        "Food& Drinks (monthly)",
        "Pet Supplies",
    ]


def test_import_review_workbook_learns_reviewed_new_leaf_after_registry_validation(
    tmp_path, capsys
):
    decisions_path = tmp_path / "review_required_2026_apr.xlsx"
    memory_dir = tmp_path / "data" / "category_memory"
    config_dir = tmp_path / "config"
    _create_category_registry_config(config_dir)
    _write_review_workbook_decisions(
        decisions_path,
        review_rows=[
            {
                "transaction_id": "tx-new-leaf",
                "description": "PET SHOP",
                "amount": "-42.50",
                "manual_category": "",
                "new_parent_category": "Living expenses",
                "new_leaf_category": "Pet Supplies",
                "learn_to_memory": "yes",
            }
        ],
        category_options=[
            ("Food& Drinks (monthly)", True),
        ],
    )

    exit_code = main(
        [
            "learn-category-memory",
            "--decisions",
            str(decisions_path),
            "--memory-dir",
            str(memory_dir),
            "--config-dir",
            str(config_dir),
        ]
    )

    captured = capsys.readouterr()
    assert exit_code == 0
    assert "Imported category memory decisions: 1" in captured.out
    memory = load_category_memory(memory_dir)
    assert [(mapping.merchant_identity, mapping.category) for mapping in memory.mappings] == [
        ("PET SHOP", "Pet Supplies"),
    ]


def test_import_review_workbook_does_not_learn_unconfirmed_local_llm_suggestion(
    tmp_path, capsys
):
    decisions_path = tmp_path / "review_required_2026_apr.xlsx"
    memory_dir = tmp_path / "data" / "category_memory"
    _write_review_workbook_decisions(
        decisions_path,
        review_rows=[
            {
                "transaction_id": "tx-local-llm",
                "description": "UNKNOWN TRAVEL",
                "amount": "-42.50",
                "suggested_category": "Traveling",
                "method": "local_llm_gemma",
                "manual_category": "",
                "learn_to_memory": "yes",
            }
        ],
        category_options=[("Traveling", True)],
    )

    exit_code = main(
        [
            "learn-category-memory",
            "--decisions",
            str(decisions_path),
            "--memory-dir",
            str(memory_dir),
        ]
    )

    captured = capsys.readouterr()
    assert exit_code == 0
    assert "Imported category memory decisions: 0" in captured.out
    assert "Skipped unlearned decisions: 1" in captured.out
    memory = load_category_memory(memory_dir)
    assert memory.mappings == ()


def test_import_review_workbook_rejects_unsupported_transaction_id_scheme(tmp_path, capsys):
    decisions_path = tmp_path / "review_required_2026_apr.xlsx"
    memory_dir = tmp_path / "data" / "category_memory"
    _write_review_workbook_decisions(
        decisions_path,
        review_rows=[
            {
                "transaction_id": "tx-food",
                "description": "NETTO 1234 KOBENHAVN",
                "amount": "-125.50",
                "manual_category": "Food& Drinks (monthly)",
                "learn_to_memory": "yes",
            },
        ],
        category_options=[
            ("Food& Drinks (monthly)", True),
        ],
        transaction_id_scheme="legacy-v0",
    )

    exit_code = main(
        [
            "learn-category-memory",
            "--decisions",
            str(decisions_path),
            "--memory-dir",
            str(memory_dir),
        ]
    )

    captured = capsys.readouterr()
    assert exit_code == 1
    assert "Unsupported reviewed workbook transaction ID scheme" in captured.err
    assert not (memory_dir / "category_memory.json").exists()


def test_import_review_workbook_preserves_memory_and_recurring_hints(tmp_path):
    memory_dir = tmp_path / "data" / "category_memory"
    first_decisions = tmp_path / "first_reviewed_decisions.csv"
    review_workbook = tmp_path / "review_required_2026_apr.xlsx"
    _write_decisions(
        first_decisions,
        [
            {
                "transaction_id": "tx-netto-1",
                "date": "2026-04-12",
                "description": "NETTO 1234 KOBENHAVN",
                "amount": "-125.50",
                "direction": "expense",
                "confirmed_category": "Food& Drinks (monthly)",
                "confirmed": "yes",
            }
        ],
    )
    import_reviewed_decisions(first_decisions, memory_dir)
    _write_review_workbook_decisions(
        review_workbook,
        review_rows=[
            {
                "transaction_id": "tx-mobile-1",
                "description": "PRIVATE TELECOM 4321",
                "amount": "-99.00",
                "manual_category": "Mobile phone (monthly)",
                "learn_to_memory": "yes",
                "recurring": "yes",
                "amount_tolerance": "2.00",
                "day_min": "25",
                "day_max": "31",
            }
        ],
        category_options=[
            ("Mobile phone (monthly)", True),
        ],
    )

    import_reviewed_decisions(review_workbook, memory_dir)

    memory = load_category_memory(memory_dir)
    assert [mapping.category for mapping in memory.mappings] == [
        "Food& Drinks (monthly)",
        "Mobile phone (monthly)",
    ]
    recurring = memory.mappings[1].recurring_hint
    assert recurring is not None
    assert recurring.amount == Decimal("99.00")
    assert recurring.amount_tolerance == Decimal("2.00")
    assert recurring.day_min == 25
    assert recurring.day_max == 31


def test_import_review_workbook_updates_reviewed_policy_from_learned_decisions(tmp_path):
    decisions_path = tmp_path / "review_required_2026_apr.xlsx"
    memory_dir = tmp_path / "data" / "category_memory"
    _write_review_workbook_decisions(
        decisions_path,
        review_rows=[
            {
                "transaction_id": "tx-mobilepay-rejsekort",
                "description": "MOBILEPAY REJSEKORT 1234",
                "amount": "-45.00",
                "manual_category": "Traveling",
                "learn_to_memory": "yes",
            },
            {
                "transaction_id": "tx-one-off",
                "description": "ONE OFF SHOP",
                "amount": "-20.00",
                "manual_category": "Shopping (monthly)",
                "learn_to_memory": "",
            },
        ],
        category_options=[
            ("Traveling", True),
            ("Shopping (monthly)", True),
        ],
    )

    import_reviewed_decisions(decisions_path, memory_dir)

    policy = (memory_dir / "reviewed_policy.local.md").read_text(encoding="utf-8")
    assert "- Merchant identity: `MOBILEPAY REJSEKORT` -> `Traveling`" in policy
    assert "ONE OFF SHOP" not in policy
    assert "## Manual Guidance" in policy


def test_import_review_workbook_preserves_reviewed_policy_manual_guidance(tmp_path):
    decisions_path = tmp_path / "review_required_2026_apr.xlsx"
    memory_dir = tmp_path / "data" / "category_memory"
    memory_dir.mkdir(parents=True)
    (memory_dir / "reviewed_policy.local.md").write_text(
        "\n".join(
            [
                "# Reviewed Policy",
                "",
                "<!-- AUTO-GENERATED REVIEWED EXAMPLES START -->",
                "- stale autogenerated entry",
                "<!-- AUTO-GENERATED REVIEWED EXAMPLES END -->",
                "",
                "## Manual Guidance",
                "",
                "- Prefer no_suggestion for unknown private transfers.",
                "",
            ]
        ),
        encoding="utf-8",
    )
    _write_review_workbook_decisions(
        decisions_path,
        review_rows=[
            {
                "transaction_id": "tx-apple",
                "description": "APPLE COM BILL 1234",
                "amount": "-25.00",
                "manual_category": "Apple Cloud",
                "learn_to_memory": "yes",
            },
        ],
        category_options=[("Apple Cloud", True)],
    )

    import_reviewed_decisions(decisions_path, memory_dir)

    policy = (memory_dir / "reviewed_policy.local.md").read_text(encoding="utf-8")
    assert "- stale autogenerated entry" not in policy
    assert "- Merchant identity: `APPLE COM BILL` -> `Apple Cloud`" in policy
    assert "- Prefer no_suggestion for unknown private transfers." in policy


def test_console_entrypoint_imports_category_memory_when_argv_is_none(
    tmp_path, monkeypatch
):
    decisions_path = tmp_path / "reviewed_decisions.csv"
    memory_dir = tmp_path / "data" / "category_memory"
    _write_decisions(
        decisions_path,
        [
            {
                "transaction_id": "tx-netto-1",
                "date": "2026-04-12",
                "description": "NETTO 1234 KOBENHAVN",
                "amount": "-125.50",
                "direction": "expense",
                "confirmed_category": "Food& Drinks (monthly)",
                "confirmed": "yes",
            }
        ],
    )
    monkeypatch.setattr(
        sys,
        "argv",
        [
            "wealth-tracker",
            "learn-category-memory",
            "--decisions",
            str(decisions_path),
            "--memory-dir",
            str(memory_dir),
        ],
    )

    assert main() == 0
    assert (memory_dir / "category_memory.json").exists()


def test_category_memory_matches_future_transaction_by_normalized_merchant_identity(tmp_path):
    decisions_path = tmp_path / "reviewed_decisions.csv"
    memory_dir = tmp_path / "data" / "category_memory"
    _write_decisions(
        decisions_path,
        [
            {
                "transaction_id": "tx-netto-1",
                "date": "2026-04-12",
                "description": "NETTO 1234 KOBENHAVN",
                "amount": "-125.50",
                "direction": "expense",
                "confirmed_category": "Food& Drinks (monthly)",
                "confirmed": "yes",
            }
        ],
    )
    import_reviewed_decisions(decisions_path, memory_dir)

    result = categorize_transactions(
        [_transaction("NETTO 5678 KOBENHAVN", "-89.95")],
        _config(),
        category_memory=load_category_memory(memory_dir),
    )[0]

    assert result.suggested_category == "Food& Drinks (monthly)"
    assert result.categorization_method == "category_memory"
    assert result.review_required is False


def test_recurring_category_memory_hints_constrain_future_matches(tmp_path):
    decisions_path = tmp_path / "reviewed_decisions.csv"
    memory_dir = tmp_path / "data" / "category_memory"
    _write_decisions(
        decisions_path,
        [
            {
                "transaction_id": "tx-mobile-1",
                "date": "2026-04-28",
                "description": "PRIVATE TELECOM 4321",
                "amount": "-99.00",
                "direction": "expense",
                "confirmed_category": "Mobile phone (monthly)",
                "confirmed": "yes",
                "recurring": "yes",
                "amount_tolerance": "2.00",
                "day_min": "25",
                "day_max": "31",
            }
        ],
    )
    import_reviewed_decisions(decisions_path, memory_dir)
    memory = load_category_memory(memory_dir)

    matching = categorize_transactions(
        [_transaction("PRIVATE TELECOM 5555", "-100.50", date(2026, 5, 29))],
        _config(categories=("Mobile phone (monthly)",)),
        category_memory=memory,
    )[0]
    outside_hint = categorize_transactions(
        [_transaction("PRIVATE TELECOM 5555", "-100.50", date(2026, 5, 10))],
        _config(categories=("Mobile phone (monthly)",)),
        category_memory=memory,
    )[0]

    assert matching.suggested_category == "Mobile phone (monthly)"
    assert matching.categorization_method == "category_memory"
    assert outside_hint.suggested_category is None
    assert outside_hint.review_required is True


def test_import_rejects_incomplete_confirmed_rows_with_useful_message(tmp_path, capsys):
    decisions_path = tmp_path / "reviewed_decisions.csv"
    memory_dir = tmp_path / "data" / "category_memory"
    _write_decisions(
        decisions_path,
        [
            {
                "transaction_id": "tx-netto-1",
                "date": "2026-04-12",
                "description": "NETTO 1234 KOBENHAVN",
                "amount": "-125.50",
                "direction": "expense",
                "confirmed_category": "",
                "confirmed": "yes",
            }
        ],
    )

    exit_code = main(
        [
            "learn-category-memory",
            "--decisions",
            str(decisions_path),
            "--memory-dir",
            str(memory_dir),
        ]
    )

    captured = capsys.readouterr()
    assert exit_code == 1
    assert "row 2" in captured.err
    assert "confirmed_category" in captured.err
    assert not (memory_dir / "category_memory.json").exists()


def test_import_rejects_invalid_recurring_hint_with_useful_message(tmp_path, capsys):
    decisions_path = tmp_path / "reviewed_decisions.csv"
    memory_dir = tmp_path / "data" / "category_memory"
    _write_decisions(
        decisions_path,
        [
            {
                "transaction_id": "tx-mobile-1",
                "date": "2026-04-28",
                "description": "PRIVATE TELECOM 4321",
                "amount": "-99.00",
                "direction": "expense",
                "confirmed_category": "Mobile phone (monthly)",
                "confirmed": "yes",
                "recurring": "yes",
                "amount_tolerance": "2.00",
                "day_min": "31",
                "day_max": "25",
            }
        ],
    )

    exit_code = main(
        [
            "learn-category-memory",
            "--decisions",
            str(decisions_path),
            "--memory-dir",
            str(memory_dir),
        ]
    )

    captured = capsys.readouterr()
    assert exit_code == 1
    assert "row 2" in captured.err
    assert "day_min" in captured.err
    assert "day_max" in captured.err
    assert not (memory_dir / "category_memory.json").exists()


def test_import_does_not_learn_from_unconfirmed_rows(tmp_path, capsys):
    decisions_path = tmp_path / "reviewed_decisions.csv"
    memory_dir = tmp_path / "data" / "category_memory"
    _write_decisions(
        decisions_path,
        [
            {
                "transaction_id": "tx-auto-guess",
                "date": "2026-04-12",
                "description": "GUESSY MERCHANT 9999",
                "amount": "-125.50",
                "direction": "expense",
                "confirmed_category": "Shopping (monthly)",
                "confirmed": "no",
            }
        ],
    )

    exit_code = main(
        [
            "learn-category-memory",
            "--decisions",
            str(decisions_path),
            "--memory-dir",
            str(memory_dir),
        ]
    )

    captured = capsys.readouterr()
    payload = json.loads((memory_dir / "category_memory.json").read_text(encoding="utf-8"))
    assert exit_code == 0
    assert "Skipped unconfirmed decisions: 1" in captured.out
    assert payload["mappings"] == []


def test_import_preserves_existing_category_memory_mappings(tmp_path):
    memory_dir = tmp_path / "data" / "category_memory"
    first_decisions = tmp_path / "first_reviewed_decisions.csv"
    second_decisions = tmp_path / "second_reviewed_decisions.csv"
    _write_decisions(
        first_decisions,
        [
            {
                "transaction_id": "tx-netto-1",
                "date": "2026-04-12",
                "description": "NETTO 1234 KOBENHAVN",
                "amount": "-125.50",
                "direction": "expense",
                "confirmed_category": "Food& Drinks (monthly)",
                "confirmed": "yes",
            }
        ],
    )
    _write_decisions(
        second_decisions,
        [
            {
                "transaction_id": "tx-travel-1",
                "date": "2026-04-28",
                "description": "REDACTED FOREIGN CARD",
                "amount": "-86.10",
                "direction": "expense",
                "confirmed_category": "Traveling",
                "confirmed": "yes",
            }
        ],
    )

    import_reviewed_decisions(first_decisions, memory_dir)
    import_reviewed_decisions(second_decisions, memory_dir)

    memory = load_category_memory(memory_dir)
    assert [mapping.category for mapping in memory.mappings] == [
        "Food& Drinks (monthly)",
        "Traveling",
    ]


def test_pipeline_uses_existing_category_memory_for_future_dry_run_matches(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    tracker = tmp_path / "tracker.xlsx"
    config_dir = tmp_path / "config"
    memory_dir = tmp_path / "data" / "category_memory"
    _create_tracker(tracker)
    _create_pipeline_config(config_dir)
    _write_decisions(
        tmp_path / "reviewed_decisions.csv",
        [
            {
                "transaction_id": "tx-travel-reviewed",
                "date": "2026-03-28",
                "description": "REDACTED FOREIGN CARD 6,9700 NOK 123,45",
                "amount": "-86.10",
                "direction": "expense",
                "confirmed_category": "Traveling",
                "confirmed": "yes",
            }
        ],
    )
    import_reviewed_decisions(tmp_path / "reviewed_decisions.csv", memory_dir)

    result = run_pipeline(
        tracker_path=tracker,
        statement_path=FIXTURE,
        config_dir=config_dir,
        year=2026,
        month="Apr",
        output_dir=tmp_path / "reports",
        category_memory_dir=memory_dir,
    )

    travel = [
        item for item in result.categorized_transactions if item.suggested_category == "Traveling"
    ]
    assert len(travel) == 1
    assert travel[0].categorization_method == "category_memory"
    assert travel[0].review_required is False


def _write_decisions(path, rows):
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=rows[0].keys())
        writer.writeheader()
        writer.writerows(rows)


def _write_review_workbook_decisions(
    path,
    review_rows,
    category_options,
    transaction_id_scheme=TRANSACTION_ID_SCHEME_VERSION,
):
    workbook = Workbook()
    review_sheet = workbook.active
    review_sheet.title = "Review Required"
    review_sheet.append(
        [
            "transaction_id",
            "date",
            "description",
            "amount",
            "direction",
            "merchant_identity",
            "suggested_category",
            "confidence",
            "method",
            "reason",
            "manual_category",
            "new_parent_category",
            "new_leaf_category",
            "learn_to_memory",
            "recurring",
            "amount_tolerance",
            "day_min",
            "day_max",
        ]
    )
    for row in review_rows:
        review_sheet.append(
            [
                row["transaction_id"],
                "2026-04-12",
                row["description"],
                row["amount"],
                "expense",
                "",
                row.get("suggested_category", ""),
                row.get("confidence", ""),
                row.get("method", ""),
                row.get("reason", ""),
                row["manual_category"],
                row.get("new_parent_category", ""),
                row.get("new_leaf_category", ""),
                row["learn_to_memory"],
                row.get("recurring", ""),
                row.get("amount_tolerance", ""),
                row.get("day_min", ""),
                row.get("day_max", ""),
            ]
        )

    options_sheet = workbook.create_sheet("Category Options")
    options_sheet.append(["row_number", "category", "learnable", "status"])
    for index, (category, learnable) in enumerate(category_options, start=2):
        options_sheet.append([index, category, learnable, "write eligible" if learnable else "derived"])

    metadata_sheet = workbook.create_sheet("Run Metadata")
    for key, value in [
        ("reporting_year", 2026),
        ("reporting_month", "Apr"),
        ("statement_parser", "nordea-csv"),
        ("generated_timestamp", "2026-05-16T10:00:00"),
        ("transaction_id_scheme", transaction_id_scheme),
    ]:
        metadata_sheet.append([key, value])
    workbook.save(path)
    workbook.close()


def _create_category_registry_config(config_dir: Path) -> None:
    config_dir.mkdir(parents=True)
    _write_text(
        config_dir / "settings.yaml",
        """
tracker:
  currency: "DKK"
statement:
  currency: "DKK"
""",
    )
    _write_text(
        config_dir / "categories.yaml",
        """
category_registry:
  - label: "Living expenses"
    type: "parent"
    allow_new_children: true
    children:
      - "Food& Drinks (monthly)"
      - "Pet Supplies"
  - label: "Income (net)"
    type: "derived"
aliases: {}
""",
    )
    _write_text(config_dir / "rules.yaml", "{}\n")


def _config(categories=("Food& Drinks (monthly)",)) -> AppConfig:
    return AppConfig(
        sheet_name="Net worth",
        tracker_currency="DKK",
        category_column=2,
        year_header_row=2,
        month_header_row=3,
        statement_currency="DKK",
        review_threshold=0.60,
        reject_threshold=0.60,
        overwrite_fixed_rows=False,
        highlight_auto_filled_cells=False,
        categories=categories,
        aliases={},
        historical_mappings={},
        rules=(),
        fixed_rows=frozenset(),
    )


def _transaction(
    description: str, amount: str, transaction_date: date = date(2026, 5, 2)
) -> Transaction:
    amount_value = Decimal(amount)
    return Transaction(
        transaction_id="future-tx",
        date=transaction_date,
        interest_date=None,
        description=description,
        amount=amount_value,
        currency="DKK",
        direction="income" if amount_value >= 0 else "expense",
    )


def _create_tracker(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Net worth"
    sheet["C2"] = 2026
    sheet["C3"] = "Apr"
    sheet["B5"] = "Apple Cloud"
    sheet["B6"] = "Full-time job (net)"
    sheet["B7"] = "Traveling"
    workbook.save(path)
    workbook.close()


def _create_pipeline_config(config_dir: Path) -> None:
    config_dir.mkdir(parents=True)
    _write_text(
        config_dir / "settings.yaml",
        """
tracker:
  sheet_name: "Net worth"
  currency: "DKK"
  category_column: 2
  year_header_row: 2
  month_header_row: 3
statement:
  currency: "DKK"
confidence_thresholds:
  auto_write: 0.85
  review_required: 0.60
  reject_below: 0.60
writer:
  overwrite_fixed_rows: false
  highlight_auto_filled_cells: false
""",
    )
    _write_text(
        config_dir / "categories.yaml",
        """
categories:
  - "Apple Cloud"
  - "Full-time job (net)"
  - "Traveling"
aliases: {}
""",
    )
    _write_text(
        config_dir / "rules.yaml",
        """
historical_mappings:
  "APPLE.COM/BILL": "Apple Cloud"
rules:
  - category: "Full-time job (net)"
    match_keywords: ["salary"]
    direction: "income"
    confidence: 0.95
fixed_rows: []
""",
    )


def _write_text(path: Path, content: str) -> None:
    path.write_text(content.lstrip(), encoding="utf-8")
