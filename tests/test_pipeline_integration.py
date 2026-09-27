import difflib
from pathlib import Path
from decimal import Decimal

from openpyxl import Workbook, load_workbook
import pytest

from personal_wealth_tracker.config import load_config
from personal_wealth_tracker.models import Authority
from personal_wealth_tracker.pipeline import run_pipeline
from personal_wealth_tracker.suggester import ConsensusSuggester, FakeSuggester, ScriptedVote
from personal_wealth_tracker.utils import TRANSACTION_ID_SCHEME_VERSION


FIXTURE = Path(__file__).parent / "fixtures" / "nordea_account_statement.redacted.pdf"


def test_pipeline_dry_run_writes_outputs_without_changing_workbook(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    tracker = tmp_path / "tracker.xlsx"
    config_dir = tmp_path / "config"
    _create_tracker(tracker)
    _create_config(config_dir)

    result = run_pipeline(
        tracker_path=tracker,
        statement_path=FIXTURE,
        config_dir=config_dir,
        year=2026,
        month="Apr",
        output_dir=tmp_path / "reports",
    )

    workbook = load_workbook(tracker)
    try:
        sheet = workbook["Net worth"]
        assert sheet["C5"].value is None
        assert sheet["C6"].value is None
        assert sheet["C7"].value is None
    finally:
        workbook.close()

    assert result.mode == "dry-run"
    assert result.output_workbook_path is None
    assert len(result.transactions) == 3
    assert {update.category for update in result.updates} == {
        "Apple Cloud",
        "Full-time job (net)",
        "Traveling",
    }
    assert result.report_path.exists()
    assert result.audit_path.exists()
    assert result.categorized_csv_path.exists()
    assert result.review_csv_path.exists()
    assert result.review_xlsx_path == tmp_path / "reports" / "review_required_2026_apr.xlsx"
    assert result.review_xlsx_path.exists()


def test_pipeline_applies_exact_proxy_split_rule_without_double_counting_source(
    tmp_path, monkeypatch
):
    monkeypatch.chdir(tmp_path)
    tracker = tmp_path / "tracker.xlsx"
    config_dir = tmp_path / "config"
    statement = tmp_path / "statement.csv"
    _create_proxy_split_tracker(tracker)
    _create_proxy_split_config(config_dir)
    _write_revolut_statement(statement, amount="-9840,00")

    result = run_pipeline(
        tracker_path=tracker,
        statement_path=statement,
        config_dir=config_dir,
        year=2026,
        month="Apr",
        output_dir=tmp_path / "reports",
    )

    assert len(result.transactions) == 1
    assert len(result.categorized_transactions) == 3
    source_transaction_id = result.transactions[0].transaction_id
    source_line = next(
        item
        for item in result.categorized_transactions
        if item.categorization_method == "proxy_split_source"
    )
    assert source_line.transaction.transaction_id == source_transaction_id
    assert source_line.suggested_category is None
    assert source_line.review_required is False
    assert source_line.split_rule == "example_transfer"
    assert source_line.split_role == "source"

    allocations = {
        item.suggested_category: item
        for item in result.categorized_transactions
        if item.categorization_method == "proxy_split_allocation"
    }
    assert set(allocations) == {"Parent A", "Parent B"}
    assert allocations["Parent A"].transaction.transaction_id == (
        f"{source_transaction_id}:split:example_transfer:parent_a"
    )
    assert allocations["Parent B"].transaction.transaction_id == (
        f"{source_transaction_id}:split:example_transfer:parent_b"
    )
    assert allocations["Parent A"].transaction.amount == Decimal("-6560.00")
    assert allocations["Parent B"].transaction.amount == Decimal("-3280.00")
    assert allocations["Parent A"].source_transaction_id == source_transaction_id
    assert allocations["Parent B"].source_transaction_id == source_transaction_id
    assert allocations["Parent A"].split_rule == "example_transfer"
    assert allocations["Parent B"].split_rule == "example_transfer"

    updates = {update.category: update for update in result.updates}
    assert set(updates) == {"Parent A", "Parent B"}
    assert updates["Parent A"].amount == Decimal("6560.00")
    assert updates["Parent B"].amount == Decimal("3280.00")
    assert source_transaction_id not in {
        transaction_id
        for update in result.updates
        for transaction_id in update.source_transactions
    }

    categorized_csv = result.categorized_csv_path.read_text(encoding="utf-8")
    assert "proxy_split_source" in categorized_csv
    assert "proxy_split_allocation" in categorized_csv
    assert f"{source_transaction_id}:split:example_transfer:parent_a" in categorized_csv
    assert f"{source_transaction_id}:split:example_transfer:parent_b" in categorized_csv

    review_csv = result.review_csv_path.read_text(encoding="utf-8")
    assert "proxy_split_source" in review_csv
    assert "proxy_split_allocation" in review_csv
    assert f"{source_transaction_id}:split:example_transfer:parent_a" in review_csv

    audit = result.audit_path.read_text(encoding="utf-8")
    assert '"categorization_method": "proxy_split_source"' in audit
    assert '"categorization_method": "proxy_split_allocation"' in audit
    assert f"{source_transaction_id}:split:example_transfer:parent_a" in audit
    assert "Parent A: 6560.00" in result.report_path.read_text(encoding="utf-8")

    workbook = load_workbook(result.review_xlsx_path, data_only=True)
    try:
        audit_sheet = workbook["Audit"]
        headers = {cell.value: index for index, cell in enumerate(audit_sheet[1], start=1)}
        methods = [
            audit_sheet.cell(row=row, column=headers["source"]).value
            for row in range(2, audit_sheet.max_row + 1)
        ]
    finally:
        workbook.close()
    assert methods == ["proxy_split_source"]


def test_pipeline_emits_residual_proxy_split_review_line_for_larger_transfer(
    tmp_path, monkeypatch
):
    monkeypatch.chdir(tmp_path)
    tracker = tmp_path / "tracker.xlsx"
    config_dir = tmp_path / "config"
    statement = tmp_path / "statement.csv"
    _create_proxy_split_tracker(tracker)
    _create_proxy_split_config(config_dir)
    _write_revolut_statement(statement, amount="-12000,00")

    result = run_pipeline(
        tracker_path=tracker,
        statement_path=statement,
        config_dir=config_dir,
        year=2026,
        month="Apr",
        output_dir=tmp_path / "reports",
    )

    source_transaction_id = result.transactions[0].transaction_id
    residual = next(
        item
        for item in result.categorized_transactions
        if item.categorization_method == "proxy_split_residual"
    )
    assert residual.transaction.transaction_id == (
        f"{source_transaction_id}:split:example_transfer:residual"
    )
    assert residual.transaction.amount == Decimal("-2160.00")
    assert residual.review_required is True
    assert residual.suggested_category is None
    assert residual.source_transaction_id == source_transaction_id
    assert residual.split_rule == "example_transfer"
    assert residual.split_role == "residual"
    assert residual.residual_amount == Decimal("2160.00")

    updates = {update.category: update for update in result.updates}
    assert set(updates) == {"Parent A", "Parent B"}
    assert "proxy_split_residual" in result.review_csv_path.read_text(encoding="utf-8")

    workbook = load_workbook(result.review_xlsx_path, data_only=True)
    try:
        review_sheet = workbook["Review Required"]
        headers = {cell.value: index for index, cell in enumerate(review_sheet[1], start=1)}
        review_ids = [
            review_sheet.cell(row=row, column=headers["transaction_id"]).value
            for row in range(2, review_sheet.max_row + 1)
        ]
    finally:
        workbook.close()
    assert residual.transaction.transaction_id in review_ids


def test_pipeline_applies_review_decision_to_residual_proxy_split_line_only(
    tmp_path, monkeypatch
):
    monkeypatch.chdir(tmp_path)
    tracker = tmp_path / "tracker.xlsx"
    config_dir = tmp_path / "config"
    statement = tmp_path / "statement.csv"
    _create_proxy_split_tracker(tracker)
    _create_proxy_split_config(config_dir)
    _write_revolut_statement(statement, amount="-12000,00")
    first_result = run_pipeline(
        tracker_path=tracker,
        statement_path=statement,
        config_dir=config_dir,
        year=2026,
        month="Apr",
        output_dir=tmp_path / "first_reports",
    )
    residual = next(
        item
        for item in first_result.categorized_transactions
        if item.categorization_method == "proxy_split_residual"
    )
    review_decisions = tmp_path / "review_required_2026_apr.xlsx"
    _write_review_decisions(
        review_decisions,
        year=2026,
        month="Apr",
        rows=[
            {
                "transaction_id": residual.transaction.transaction_id,
                "description": residual.transaction.description,
                "amount": str(residual.transaction.amount),
                "manual_category": "Traveling",
            }
        ],
    )

    result = run_pipeline(
        tracker_path=tracker,
        statement_path=statement,
        config_dir=config_dir,
        year=2026,
        month="Apr",
        output_dir=tmp_path / "reports",
        review_decisions_path=review_decisions,
    )

    reviewed_residual = next(
        item
        for item in result.categorized_transactions
        if item.transaction.transaction_id == residual.transaction.transaction_id
    )
    assert reviewed_residual.suggested_category == "Traveling"
    assert reviewed_residual.categorization_method == "monthly_review_decision"
    assert reviewed_residual.review_required is False
    assert reviewed_residual.transaction.amount == Decimal("-2160.00")
    assert reviewed_residual.source_transaction_id == first_result.transactions[0].transaction_id

    updates = {update.category: update for update in result.updates}
    assert updates["Parent A"].amount == Decimal("6560.00")
    assert updates["Parent B"].amount == Decimal("3280.00")
    assert updates["Traveling"].amount == Decimal("2160.00")
    assert updates["Traveling"].source_transactions == (residual.transaction.transaction_id,)


def test_pipeline_keeps_underfunded_proxy_split_candidate_review_only(
    tmp_path, monkeypatch
):
    monkeypatch.chdir(tmp_path)
    tracker = tmp_path / "tracker.xlsx"
    config_dir = tmp_path / "config"
    statement = tmp_path / "statement.csv"
    _create_proxy_split_tracker(tracker)
    _create_proxy_split_config(config_dir)
    _write_revolut_statement(statement, amount="-9000,00")

    result = run_pipeline(
        tracker_path=tracker,
        statement_path=statement,
        config_dir=config_dir,
        year=2026,
        month="Apr",
        output_dir=tmp_path / "reports",
    )

    assert len(result.categorized_transactions) == 1
    blocked = result.categorized_transactions[0]
    assert blocked.categorization_method == "proxy_split_blocked"
    assert blocked.review_required is True
    assert blocked.suggested_category is None
    assert blocked.split_rule == "example_transfer"
    assert blocked.split_role == "source"
    assert "below configured proxy split allocation total" in blocked.reason
    assert result.updates == []
    assert "proxy_split_blocked" in result.review_csv_path.read_text(encoding="utf-8")


def test_pipeline_blocks_multiple_proxy_split_candidates_in_one_reporting_month(
    tmp_path, monkeypatch
):
    monkeypatch.chdir(tmp_path)
    tracker = tmp_path / "tracker.xlsx"
    config_dir = tmp_path / "config"
    statement = tmp_path / "statement.csv"
    _create_proxy_split_tracker(tracker)
    _create_proxy_split_config(config_dir)
    _write(
        statement,
        "\ufeffBooking date;Amount;Balance;Currency;Name;Title;Sender;Recipient;Reconciled\n"
        "2026/04/04;-9840,00;50000,00;DKK;REVOLUT;Transfer;1111;2222;Yes\n"
        "2026/04/18;-12000,00;38000,00;DKK;REVOLUT;Transfer;1111;2222;Yes\n",
    )

    result = run_pipeline(
        tracker_path=tracker,
        statement_path=statement,
        config_dir=config_dir,
        year=2026,
        month="Apr",
        output_dir=tmp_path / "reports",
    )

    assert len(result.categorized_transactions) == 2
    assert {item.categorization_method for item in result.categorized_transactions} == {
        "proxy_split_blocked"
    }
    assert all(item.review_required for item in result.categorized_transactions)
    assert all(
        "Multiple proxy split candidates" in item.reason
        for item in result.categorized_transactions
    )
    assert result.updates == []
    report = result.report_path.read_text(encoding="utf-8")
    assert "Multiple proxy split candidates" in report
    assert "Parent A: 6560.00" not in report


def test_pipeline_allows_multiple_proxy_split_candidates_up_to_configured_limit(
    tmp_path, monkeypatch
):
    monkeypatch.chdir(tmp_path)
    tracker = tmp_path / "tracker.xlsx"
    config_dir = tmp_path / "config"
    statement = tmp_path / "statement.csv"
    _create_proxy_split_tracker(tracker)
    _create_proxy_split_config(config_dir, monthly_limit=2)
    _write(
        statement,
        "\ufeffBooking date;Amount;Balance;Currency;Name;Title;Sender;Recipient;Reconciled\n"
        "2026/04/04;-9840,00;50000,00;DKK;REVOLUT;Transfer;1111;2222;Yes\n"
        "2026/04/18;-15000,00;35000,00;DKK;REVOLUT;Transfer;1111;2222;Yes\n",
    )

    result = run_pipeline(
        tracker_path=tracker,
        statement_path=statement,
        config_dir=config_dir,
        year=2026,
        month="Apr",
        output_dir=tmp_path / "reports",
    )

    assert len(result.transactions) == 2
    assert sum(
        1
        for item in result.categorized_transactions
        if item.categorization_method == "proxy_split_source"
    ) == 2
    assert sum(
        1
        for item in result.categorized_transactions
        if item.categorization_method == "proxy_split_allocation"
        and item.suggested_category == "Parent A"
    ) == 2
    assert sum(
        1
        for item in result.categorized_transactions
        if item.categorization_method == "proxy_split_allocation"
        and item.suggested_category == "Parent B"
    ) == 2
    assert not any(
        item.categorization_method == "proxy_split_blocked"
        for item in result.categorized_transactions
    )

    residuals = [
        item
        for item in result.categorized_transactions
        if item.categorization_method == "proxy_split_residual"
    ]
    assert len(residuals) == 1
    assert residuals[0].transaction.amount == Decimal("-5160.00")

    updates = {update.category: update for update in result.updates}
    assert updates["Parent A"].amount == Decimal("13120.00")
    assert updates["Parent B"].amount == Decimal("6560.00")
    assert "Multiple proxy split candidates" not in result.report_path.read_text(
        encoding="utf-8"
    )


def test_pipeline_dry_run_reports_missing_period_creation_without_changing_workbook(
    tmp_path, monkeypatch
):
    monkeypatch.chdir(tmp_path)
    tracker = tmp_path / "tracker.xlsx"
    config_dir = tmp_path / "config"
    _create_tracker(tracker)
    _create_config(config_dir)
    workbook = load_workbook(tracker)
    try:
        sheet = workbook["Net worth"]
        sheet["C3"] = "Mar"
        workbook.save(tracker)
    finally:
        workbook.close()

    result = run_pipeline(
        tracker_path=tracker,
        statement_path=FIXTURE,
        config_dir=config_dir,
        year=2026,
        month="Apr",
        output_dir=tmp_path / "reports",
    )

    original = load_workbook(tracker)
    try:
        assert original["Net worth"]["D3"].value is None
    finally:
        original.close()

    assert len(result.structure_changes) == 1
    assert result.structure_changes[0].change_type == "create_period"
    report = result.report_path.read_text(encoding="utf-8")
    assert "## Planned Structure Changes" in report
    assert "create_period: Apr 2026" in report


def test_pipeline_commit_updates_tracker_in_place_after_backup(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    tracker = tmp_path / "tracker.xlsx"
    config_dir = tmp_path / "config"
    _create_tracker(tracker)
    _create_config(config_dir, trust_policy=RAISED_TRUST_POLICY)

    result = run_pipeline(
        tracker_path=tracker,
        statement_path=FIXTURE,
        config_dir=config_dir,
        year=2026,
        month="Apr",
        output_dir=tmp_path / "reports",
        commit=True,
    )

    assert result.output_workbook_path == tracker
    [backup_path] = (tmp_path / "data" / "backups").glob("tracker_backup_*.xlsx")
    updated = load_workbook(tracker)
    backup = load_workbook(backup_path)
    try:
        assert [updated["Net worth"][cell].value for cell in ("C5", "C6", "C7")] == [
            25,
            10000,
            86.1,
        ]
        assert [backup["Net worth"][cell].value for cell in ("C5", "C6", "C7")] == [
            None,
            None,
            None,
        ]
    finally:
        updated.close()
        backup.close()

    assert result.mode == "commit"
    assert result.review_count == 0
    assert not (tmp_path / "data" / "processed").exists()


def test_consecutive_month_commits_accumulate_in_one_tracker(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    tracker = tmp_path / "tracker.xlsx"
    config_dir = tmp_path / "config"
    _create_multi_month_tracker(tracker, ("Apr", "May"))
    _create_category_registry_config(config_dir)

    def consensus():
        votes = {"NOODLE BAR": ScriptedVote("Traveling", 0.9)}
        return ConsensusSuggester(FakeSuggester(votes, "a"), FakeSuggester(votes, "b"))

    for month in ("Apr", "May"):
        committed = _run_month(
            tmp_path,
            tracker,
            config_dir,
            month,
            ["NOODLE BAR"],
            commit=True,
            local_llm_suggestions=True,
            suggester=consensus(),
        )
        assert committed.output_workbook_path == tracker

    workbook = load_workbook(tracker)
    try:
        assert workbook["Net worth"]["C7"].value == 40
        assert workbook["Net worth"]["D7"].value == 40
    finally:
        workbook.close()


def _noodle_consensus(category: str):
    votes = {"NOODLE BAR": ScriptedVote(category, 0.9)}
    return ConsensusSuggester(FakeSuggester(votes, "a"), FakeSuggester(votes, "b"))


def _commit_noodle_april(tmp_path, tracker, config_dir, category: str):
    return _run_month(
        tmp_path,
        tracker,
        config_dir,
        "Apr",
        ["NOODLE BAR"],
        commit=True,
        local_llm_suggestions=True,
        suggester=_noodle_consensus(category),
    )


def test_recommitting_a_month_overwrites_its_own_values(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    tracker = tmp_path / "tracker.xlsx"
    config_dir = tmp_path / "config"
    _create_multi_month_tracker(tracker, ("Apr",))
    _create_category_registry_config(config_dir)
    _commit_noodle_april(tmp_path, tracker, config_dir, "Traveling")

    rerun = _commit_noodle_april(tmp_path, tracker, config_dir, "Traveling")

    assert rerun.review_count == 0
    assert rerun.output_workbook_path == tracker
    workbook = load_workbook(tracker)
    try:
        assert workbook["Net worth"]["C7"].value == 40
    finally:
        workbook.close()


def test_recommit_with_changed_category_clears_stale_committed_value(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    tracker = tmp_path / "tracker.xlsx"
    config_dir = tmp_path / "config"
    _create_multi_month_tracker(tracker, ("Apr",))
    _create_category_registry_config(config_dir)
    _commit_noodle_april(tmp_path, tracker, config_dir, "Traveling")

    rerun = _commit_noodle_april(tmp_path, tracker, config_dir, "Apple Cloud")

    assert rerun.output_workbook_path == tracker
    workbook = load_workbook(tracker)
    try:
        assert workbook["Net worth"]["C5"].value == 40
        assert workbook["Net worth"]["C7"].value is None
    finally:
        workbook.close()


def test_recommit_keeps_cell_the_user_edited_after_commit_in_review(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    tracker = tmp_path / "tracker.xlsx"
    config_dir = tmp_path / "config"
    _create_multi_month_tracker(tracker, ("Apr",))
    _create_category_registry_config(config_dir)
    _commit_noodle_april(tmp_path, tracker, config_dir, "Traveling")
    _set_workbook_label(tracker, "C7", 55)

    rerun = _commit_noodle_april(tmp_path, tracker, config_dir, "Traveling")

    assert rerun.review_count == 1
    assert rerun.output_workbook_path is None
    workbook = load_workbook(tracker)
    try:
        assert workbook["Net worth"]["C7"].value == 55
    finally:
        workbook.close()


def test_commit_never_writes_partial_month(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    tracker = tmp_path / "tracker.xlsx"
    config_dir = tmp_path / "config"
    _create_tracker(tracker)
    _create_config(config_dir)

    result = run_pipeline(
        tracker_path=tracker,
        statement_path=FIXTURE,
        config_dir=config_dir,
        year=2026,
        month="Apr",
        output_dir=tmp_path / "reports",
        commit=True,
    )

    salary = next(
        item
        for item in result.categorized_transactions
        if item.suggested_category == "Full-time job (net)"
    )
    assert salary.review_required is True
    assert {update.category: update.write_action for update in result.updates} == {
        "Apple Cloud": "write",
        "Full-time job (net)": "write",
        "Traveling": "write",
    }
    assert result.review_count == 1
    assert result.output_workbook_path is None
    assert not (tmp_path / "data" / "processed").exists()
    original = load_workbook(tracker)
    try:
        assert [original["Net worth"][cell].value for cell in ("C5", "C6", "C7")] == [
            None,
            None,
            None,
        ]
    finally:
        original.close()


def test_commit_with_review_rows_writes_exception_sheet_instead_of_workbook(
    tmp_path, monkeypatch
):
    monkeypatch.chdir(tmp_path)
    tracker = tmp_path / "tracker.xlsx"
    config_dir = tmp_path / "config"
    statement = tmp_path / "statement.csv"
    _create_tracker(tracker)
    _create_category_registry_config(config_dir)
    _write_unknown_shop_statement(statement)
    suggester = FakeSuggester({"UNKNOWN SHOP": ScriptedVote("Traveling", 0.9)})

    result = run_pipeline(
        tracker_path=tracker,
        statement_path=statement,
        config_dir=config_dir,
        year=2026,
        month="Apr",
        output_dir=tmp_path / "reports",
        commit=True,
        local_llm_suggestions=True,
        suggester=suggester,
    )

    assert result.output_workbook_path is None
    assert not (tmp_path / "data" / "processed").exists()
    assert not (tmp_path / "data" / "backups").exists()
    assert result.review_count == 1
    workbook = load_workbook(result.review_xlsx_path, data_only=True)
    try:
        rows = _sheet_rows(workbook["Review Required"])
    finally:
        workbook.close()
    assert [row["transaction_id"] for row in rows] == [
        result.transactions[0].transaction_id
    ]
    assert rows[0]["suggested_category"] == "Traveling"
    assert rows[0]["manual_category"] is None
    assert "Workbook not written: 1 row(s) in review." in result.report_path.read_text(
        encoding="utf-8"
    )


def test_exception_sheet_dropdown_offers_suggester_alternatives(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    tracker = tmp_path / "tracker.xlsx"
    config_dir = tmp_path / "config"
    statement = tmp_path / "statement.csv"
    _create_tracker(tracker)
    _create_category_registry_config(config_dir)
    _write_unknown_shop_statement(statement)
    suggester = FakeSuggester(
        {"UNKNOWN SHOP": ScriptedVote("Traveling", 0.9, alternatives=("Apple Cloud",))}
    )

    result = run_pipeline(
        tracker_path=tracker,
        statement_path=statement,
        config_dir=config_dir,
        year=2026,
        month="Apr",
        output_dir=tmp_path / "reports",
        local_llm_suggestions=True,
        suggester=suggester,
    )

    assert result.categorized_transactions[0].alternatives == ("Apple Cloud",)
    workbook = load_workbook(result.review_xlsx_path)
    try:
        options = _dropdown_options(workbook, "E2")
    finally:
        workbook.close()
    assert options[:3] == ["Traveling", "Apple Cloud", "NONE"]


def test_rerun_with_filled_exception_sheet_commits_month_totals(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    tracker = tmp_path / "tracker.xlsx"
    config_dir = tmp_path / "config"
    statement = tmp_path / "statement.csv"
    output_dir = tmp_path / "reports"
    _create_tracker(tracker)
    _create_category_registry_config(config_dir)
    _write(
        statement,
        "\ufeffBooking date;Amount;Balance;Currency;Name;Title;Sender;Recipient;Reconciled\n"
        "2026/04/01;-40,00;960,00;DKK;SHOP ONE;Card purchase;1111;2222;Yes\n"
        "2026/04/02;-15,00;945,00;DKK;SHOP TWO;Card purchase;1111;2222;Yes\n"
        "2026/04/03;-7,50;937,50;DKK;SHOP THREE;Card purchase;1111;2222;Yes\n",
    )
    suggester = FakeSuggester(
        {
            merchant: ScriptedVote("Traveling", 0.9)
            for merchant in ("SHOP ONE", "SHOP TWO", "SHOP THREE")
        }
    )
    pipeline_args = {
        "tracker_path": tracker,
        "statement_path": statement,
        "config_dir": config_dir,
        "year": 2026,
        "month": "Apr",
        "output_dir": output_dir,
        "commit": True,
        "local_llm_suggestions": True,
        "suggester": suggester,
    }
    first = run_pipeline(**pipeline_args)
    assert first.output_workbook_path is None
    decisions = {"SHOP ONE": None, "SHOP TWO": "NONE", "SHOP THREE": "Apple Cloud"}
    workbook = load_workbook(first.review_xlsx_path)
    try:
        sheet = workbook["Review Required"]
        headers = {cell.value: index for index, cell in enumerate(sheet[1], start=1)}
        for row in range(2, sheet.max_row + 1):
            description = sheet.cell(row=row, column=headers["description"]).value
            sheet.cell(row=row, column=headers["manual_category"]).value = next(
                decision for merchant, decision in decisions.items() if merchant in description
            )
        workbook.save(first.review_xlsx_path)
    finally:
        workbook.close()

    result = run_pipeline(**pipeline_args, review_decisions_path=first.review_xlsx_path)

    assert result.review_count == 0
    rejected = next(
        item for item in result.categorized_transactions if "SHOP TWO" in item.transaction.description
    )
    assert rejected.suggested_category is None
    assert rejected.review_required is False
    assert result.output_workbook_path == tracker
    updated = load_workbook(tracker)
    try:
        assert updated["Net worth"]["C5"].value == 7.5
        assert updated["Net worth"]["C7"].value == 40
    finally:
        updated.close()
    assert list((tmp_path / "data" / "backups").glob("tracker_backup_*.xlsx"))


def test_audit_sheet_lists_auto_rows_with_source_votes_and_reason(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    tracker = tmp_path / "tracker.xlsx"
    config_dir = tmp_path / "config"
    statement = tmp_path / "statement.csv"
    _create_tracker(tracker)
    _create_category_registry_config(config_dir)
    with (config_dir / "settings.yaml").open("a", encoding="utf-8") as handle:
        handle.write("trust_policy:\n  min_agreement: 1\n")
    _write(
        statement,
        "\ufeffBooking date;Amount;Balance;Currency;Name;Title;Sender;Recipient;Reconciled\n"
        "2026/04/01;-42,50;957,50;DKK;UNKNOWN SHOP;Card purchase;1111;2222;Yes\n"
        "2026/04/02;-2500,00;-1542,50;DKK;BIG SHOP;Card purchase;1111;2222;Yes\n",
    )
    suggester = FakeSuggester(
        {merchant: ScriptedVote("Traveling", 0.9) for merchant in ("UNKNOWN SHOP", "BIG SHOP")}
    )

    result = run_pipeline(
        tracker_path=tracker,
        statement_path=statement,
        config_dir=config_dir,
        year=2026,
        month="Apr",
        output_dir=tmp_path / "reports",
        local_llm_suggestions=True,
        suggester=suggester,
    )

    workbook = load_workbook(result.review_xlsx_path, data_only=True)
    try:
        audit_rows = _sheet_rows(workbook["Audit"])
    finally:
        workbook.close()
    assert len(audit_rows) == 1
    audit_row = audit_rows[0]
    assert "UNKNOWN SHOP" in audit_row["description"]
    assert audit_row["category"] == "Traveling"
    assert audit_row["source"] == "local_llm_gemma"
    assert audit_row["votes"] == "Traveling (fake, 0.90)"
    assert audit_row["reason"] == "1 model votes agree."


def test_committed_human_decision_is_auto_next_month_via_memory(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    tracker = tmp_path / "tracker.xlsx"
    config_dir = tmp_path / "config"
    _create_multi_month_tracker(tracker, ("Apr", "May"))
    _create_category_registry_config(config_dir)
    first = _run_month(tmp_path, tracker, config_dir, "Apr", ["COFFEE HUT"], commit=True)
    assert first.output_workbook_path is None
    _fill_exception_sheet(first.review_xlsx_path, {"COFFEE HUT": "Traveling"})
    committed = _run_month(
        tmp_path,
        tracker,
        config_dir,
        "Apr",
        ["COFFEE HUT"],
        commit=True,
        review_decisions_path=first.review_xlsx_path,
    )
    assert committed.output_workbook_path is not None

    next_month = _run_month(tmp_path, tracker, config_dir, "May", ["COFFEE HUT"])

    row = next_month.categorized_transactions[0]
    assert row.suggested_category == "Traveling"
    assert row.categorization_method == "category_memory"
    assert row.authority is Authority.auto


def test_all_digit_description_is_never_learned_as_empty_merchant(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    tracker = tmp_path / "tracker.xlsx"
    config_dir = tmp_path / "config"
    _create_multi_month_tracker(tracker, ("Apr", "May"))
    _create_category_registry_config(config_dir)
    first = _run_month(tmp_path, tracker, config_dir, "Apr", ["3084302062152494"], commit=True)
    _fill_exception_sheet(first.review_xlsx_path, {"3084302062152494": "Traveling"})
    committed = _run_month(
        tmp_path,
        tracker,
        config_dir,
        "Apr",
        ["3084302062152494"],
        commit=True,
        review_decisions_path=first.review_xlsx_path,
    )
    assert committed.output_workbook_path is not None

    next_month = _run_month(tmp_path, tracker, config_dir, "May", ["9999888877776666"])

    assert next_month.categorized_transactions[0].categorization_method != "category_memory"
    memory = tmp_path / "data" / "category_memory" / "category_memory.json"
    assert '"merchant_identity": ""' not in memory.read_text(encoding="utf-8")


def test_legacy_empty_merchant_memory_entry_never_matches(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    tracker = tmp_path / "tracker.xlsx"
    config_dir = tmp_path / "config"
    _create_multi_month_tracker(tracker, ("May",))
    _create_category_registry_config(config_dir)
    memory_dir = tmp_path / "data" / "category_memory"
    memory_dir.mkdir(parents=True)
    _write(
        memory_dir / "category_memory.json",
        '{"mappings": [{"category": "Traveling", "merchant_identity": "", '
        '"provenance": "human", "source_transaction_ids": ["legacy"]}]}',
    )

    result = _run_month(tmp_path, tracker, config_dir, "May", ["9999888877776666"])

    assert result.categorized_transactions[0].categorization_method != "category_memory"


def test_accepted_subscription_proposal_registers_leaf_under_services_on_commit_only(
    tmp_path, monkeypatch
):
    monkeypatch.chdir(tmp_path)
    tracker = tmp_path / "tracker.xlsx"
    config_dir = tmp_path / "config"
    _create_services_tracker(tracker)
    _create_services_config(config_dir)
    suggester = ConsensusSuggester(
        FakeSuggester({"CLAUDE.AI": ScriptedVote(None, 0.9, new_subscription="Claude")}),
        FakeSuggester({"CLAUDE.AI": ScriptedVote(None, 0.8, new_subscription="claude")}),
    )
    month_args = {"commit": True, "local_llm_suggestions": True, "suggester": suggester}
    categories_yaml = config_dir / "categories.yaml"
    registry_before = categories_yaml.read_text(encoding="utf-8")

    first = _run_month(tmp_path, tracker, config_dir, "Apr", ["CLAUDE.AI"], **month_args)

    assert first.output_workbook_path is None
    workbook = load_workbook(first.review_xlsx_path)
    try:
        sheet = workbook["Review Required"]
        (row,) = _sheet_rows(sheet)
        options = _dropdown_options(workbook, "F2")
    finally:
        workbook.close()
    assert list(row)[3:6] == ["suggested_category", "suggested_parent_category", "manual_category"]
    assert row["suggested_category"] == "Claude subscription"
    assert row["suggested_parent_category"] == "Services"
    assert row["manual_category"] is None
    assert options == ["NONE", "Disney+"]
    assert "- New-leaf proposals: 1" in first.report_path.read_text(encoding="utf-8")
    assert '"new_leaf_proposal_count": 1' in first.audit_path.read_text(encoding="utf-8")

    preview = _run_month(
        tmp_path,
        tracker,
        config_dir,
        "Apr",
        ["CLAUDE.AI"],
        review_decisions_path=first.review_xlsx_path,
    )
    assert preview.category_registry_additions[0].leaf_category == "Claude subscription"
    assert categories_yaml.read_text(encoding="utf-8") == registry_before

    committed = _run_month(
        tmp_path,
        tracker,
        config_dir,
        "Apr",
        ["CLAUDE.AI"],
        review_decisions_path=first.review_xlsx_path,
        **month_args,
    )

    assert committed.output_workbook_path is not None
    assert _added_lines(registry_before, categories_yaml.read_text(encoding="utf-8")) == [
        "      - label: Claude subscription",
        '        description: "Claude subscription billing."',
    ]
    registry = load_config(config_dir).category_registry
    assert registry.children_by_parent["Services"] == ("Disney+", "Claude subscription")
    assert registry.leaf_glossary["Claude subscription"] == "Claude subscription billing."
    copied = load_workbook(committed.output_workbook_path)
    try:
        assert copied["Net worth"]["B7"].value == "Claude subscription"
        assert copied["Net worth"]["C7"].value == 40
    finally:
        copied.close()

    registry_after_commit = categories_yaml.read_bytes()
    rerun = _run_month(
        tmp_path,
        tracker,
        config_dir,
        "Apr",
        ["CLAUDE.AI"],
        review_decisions_path=first.review_xlsx_path,
        **month_args,
    )

    assert rerun.output_workbook_path is not None
    assert rerun.category_registry_additions == ()
    assert categories_yaml.read_bytes() == registry_after_commit

    next_month = _run_month(tmp_path, tracker, config_dir, "May", ["CLAUDE.AI"])

    assert next_month.categorized_transactions[0].suggested_category == "Claude subscription"
    assert next_month.categorized_transactions[0].authority is Authority.auto


def test_blocked_commit_with_accepted_proposal_leaves_category_registry_untouched(
    tmp_path, monkeypatch
):
    monkeypatch.chdir(tmp_path)
    tracker = tmp_path / "tracker.xlsx"
    config_dir = tmp_path / "config"
    _create_services_tracker(tracker)
    _create_services_config(config_dir)
    categories_yaml = config_dir / "categories.yaml"
    registry_before = categories_yaml.read_bytes()
    month_args = {
        "commit": True,
        "local_llm_suggestions": True,
        "suggester": FakeSuggester(
            {"CLAUDE.AI": ScriptedVote(None, 0.9, new_subscription="Claude")}
        ),
    }
    merchants = ["CLAUDE.AI", "UNKNOWN SHOP"]
    first = _run_month(tmp_path, tracker, config_dir, "Apr", merchants, **month_args)

    blocked = _run_month(
        tmp_path,
        tracker,
        config_dir,
        "Apr",
        merchants,
        review_decisions_path=first.review_xlsx_path,
        **month_args,
    )

    assert blocked.review_count == 1
    assert blocked.output_workbook_path is None
    assert [addition.leaf_category for addition in blocked.category_registry_additions] == [
        "Claude subscription"
    ]
    assert "Claude subscription" in {update.category for update in blocked.updates}
    assert categories_yaml.read_bytes() == registry_before


def test_subscription_proposal_can_be_accepted_explicitly_with_edited_name(
    tmp_path, monkeypatch
):
    monkeypatch.chdir(tmp_path)
    tracker = tmp_path / "tracker.xlsx"
    config_dir = tmp_path / "config"
    _create_services_tracker(tracker)
    _create_services_config(config_dir)
    month_args = {
        "commit": True,
        "local_llm_suggestions": True,
        "suggester": FakeSuggester(
            {"CLAUDE.AI": ScriptedVote(None, 0.9, new_subscription="Claude")}
        ),
    }
    first = _run_month(tmp_path, tracker, config_dir, "Apr", ["CLAUDE.AI"], **month_args)
    _fill_exception_sheet(first.review_xlsx_path, {"CLAUDE.AI": "Services"}, "new_parent_category")
    _fill_exception_sheet(
        first.review_xlsx_path, {"CLAUDE.AI": "Claude Pro subscription"}, "new_leaf_category"
    )

    committed = _run_month(
        tmp_path,
        tracker,
        config_dir,
        "Apr",
        ["CLAUDE.AI"],
        review_decisions_path=first.review_xlsx_path,
        **month_args,
    )

    assert committed.output_workbook_path is not None
    registry = load_config(config_dir).category_registry
    assert registry.children_by_parent["Services"] == ("Disney+", "Claude Pro subscription")
    assert registry.leaf_glossary["Claude Pro subscription"] == (
        "Added in monthly review Apr 2026."
    )


def test_subscription_proposal_can_be_overridden_with_existing_leaf(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    tracker = tmp_path / "tracker.xlsx"
    config_dir = tmp_path / "config"
    _create_services_tracker(tracker)
    _create_services_config(config_dir)
    month_args = {
        "commit": True,
        "local_llm_suggestions": True,
        "suggester": FakeSuggester(
            {"DISNEYPLUS": ScriptedVote(None, 0.9, new_subscription="Disney Plus")}
        ),
    }
    first = _run_month(tmp_path, tracker, config_dir, "Apr", ["DISNEYPLUS"], **month_args)
    _fill_exception_sheet(first.review_xlsx_path, {"DISNEYPLUS": "Disney+"})

    committed = _run_month(
        tmp_path,
        tracker,
        config_dir,
        "Apr",
        ["DISNEYPLUS"],
        review_decisions_path=first.review_xlsx_path,
        **month_args,
    )

    assert committed.categorized_transactions[0].suggested_category == "Disney+"
    assert committed.category_registry_additions == ()
    assert load_config(config_dir).category_registry.children_by_parent["Services"] == (
        "Disney+",
    )


def test_exception_sheet_learn_to_memory_no_keeps_decision_out_of_memory(
    tmp_path, monkeypatch
):
    monkeypatch.chdir(tmp_path)
    tracker = tmp_path / "tracker.xlsx"
    config_dir = tmp_path / "config"
    _create_multi_month_tracker(tracker, ("Apr", "May"))
    _create_category_registry_config(config_dir)
    first = _run_month(tmp_path, tracker, config_dir, "Apr", ["COFFEE HUT"], commit=True)
    _fill_exception_sheet(first.review_xlsx_path, {"COFFEE HUT": "Traveling"})
    _fill_exception_sheet(first.review_xlsx_path, {"COFFEE HUT": "no"}, "learn_to_memory")
    committed = _run_month(
        tmp_path,
        tracker,
        config_dir,
        "Apr",
        ["COFFEE HUT"],
        commit=True,
        review_decisions_path=first.review_xlsx_path,
    )
    assert committed.output_workbook_path is not None

    next_month = _run_month(tmp_path, tracker, config_dir, "May", ["COFFEE HUT"])

    assert next_month.categorized_transactions[0].categorization_method == "unmatched"


def test_dry_run_writes_no_category_memory(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    tracker = tmp_path / "tracker.xlsx"
    config_dir = tmp_path / "config"
    _create_multi_month_tracker(tracker, ("Apr",))
    _create_category_registry_config(config_dir)
    first = _run_month(tmp_path, tracker, config_dir, "Apr", ["COFFEE HUT"])
    _fill_exception_sheet(first.review_xlsx_path, {"COFFEE HUT": "Traveling"})

    result = _run_month(
        tmp_path,
        tracker,
        config_dir,
        "Apr",
        ["COFFEE HUT"],
        review_decisions_path=first.review_xlsx_path,
    )

    assert result.review_count == 0
    assert not (tmp_path / "data" / "category_memory").exists()


def test_consensus_result_is_hint_after_one_month_and_memory_after_two(
    tmp_path, monkeypatch
):
    monkeypatch.chdir(tmp_path)
    tracker = tmp_path / "tracker.xlsx"
    config_dir = tmp_path / "config"
    _create_multi_month_tracker(tracker, ("Apr", "May", "Jun"))
    _create_category_registry_config(config_dir)

    def consensus():
        votes = {"NOODLE BAR": ScriptedVote("Traveling", 0.9)}
        return ConsensusSuggester(FakeSuggester(votes, "a"), FakeSuggester(votes, "b"))

    april = _run_month(
        tmp_path,
        tracker,
        config_dir,
        "Apr",
        ["NOODLE BAR"],
        commit=True,
        local_llm_suggestions=True,
        suggester=consensus(),
    )
    assert april.output_workbook_path is not None
    policy = tmp_path / "data" / "category_memory" / "reviewed_policy.local.md"
    assert "NOODLE BAR" not in policy.read_text(encoding="utf-8")
    hint_probe = FakeSuggester({})

    may_dry_run = _run_month(
        tmp_path,
        tracker,
        config_dir,
        "May",
        ["NOODLE BAR"],
        local_llm_suggestions=True,
        suggester=hint_probe,
    )

    assert may_dry_run.categorized_transactions[0].categorization_method != "category_memory"
    _, context = hint_probe.calls[0]
    assert context.memory_neighbours("NOODLE BAR") == (("NOODLE BAR", "Traveling"),)

    _run_month(
        tmp_path,
        tracker,
        config_dir,
        "May",
        ["NOODLE BAR"],
        commit=True,
        local_llm_suggestions=True,
        suggester=consensus(),
    )
    june = _run_month(tmp_path, tracker, config_dir, "Jun", ["NOODLE BAR"])

    row = june.categorized_transactions[0]
    assert row.suggested_category == "Traveling"
    assert row.categorization_method == "category_memory"


def test_recommitting_same_month_does_not_trust_auto_memory(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    tracker = tmp_path / "tracker.xlsx"
    config_dir = tmp_path / "config"
    _create_multi_month_tracker(tracker, ("Apr", "May"))
    _create_category_registry_config(config_dir)
    votes = {"NOODLE BAR": ScriptedVote("Traveling", 0.9)}
    for _ in range(2):
        _run_month(
            tmp_path,
            tracker,
            config_dir,
            "Apr",
            ["NOODLE BAR"],
            commit=True,
            local_llm_suggestions=True,
            suggester=ConsensusSuggester(FakeSuggester(votes, "a"), FakeSuggester(votes, "b")),
        )

    may = _run_month(tmp_path, tracker, config_dir, "May", ["NOODLE BAR"])

    assert may.categorized_transactions[0].categorization_method == "unmatched"


def test_audit_correction_reverses_learned_memory_and_records_corrected_category(
    tmp_path, monkeypatch
):
    monkeypatch.chdir(tmp_path)
    tracker = tmp_path / "tracker.xlsx"
    config_dir = tmp_path / "config"
    _create_multi_month_tracker(tracker, ("Apr", "May"))
    _create_category_registry_config(config_dir)
    votes = {"NOODLE BAR": ScriptedVote("Traveling", 0.9)}
    april_args = {
        "commit": True,
        "local_llm_suggestions": True,
        "suggester": ConsensusSuggester(FakeSuggester(votes, "a"), FakeSuggester(votes, "b")),
    }
    april = _run_month(tmp_path, tracker, config_dir, "Apr", ["NOODLE BAR"], **april_args)
    _fill_audit_corrections(april.review_xlsx_path, {"NOODLE BAR": "Apple Cloud"})

    corrected = _run_month(
        tmp_path,
        tracker,
        config_dir,
        "Apr",
        ["NOODLE BAR"],
        review_decisions_path=april.review_xlsx_path,
        **april_args,
    )

    copied = load_workbook(corrected.output_workbook_path)
    try:
        assert copied["Net worth"]["C5"].value == 40
        assert copied["Net worth"]["C7"].value is None
    finally:
        copied.close()
    may = _run_month(tmp_path, tracker, config_dir, "May", ["NOODLE BAR"])
    row = may.categorized_transactions[0]
    assert row.suggested_category == "Apple Cloud"
    assert row.categorization_method == "category_memory"


def test_audit_correction_to_none_forgets_learned_memory(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    tracker = tmp_path / "tracker.xlsx"
    config_dir = tmp_path / "config"
    _create_multi_month_tracker(tracker, ("Apr", "May"))
    _create_category_registry_config(config_dir)
    votes = {"NOODLE BAR": ScriptedVote("Traveling", 0.9)}

    def commit_month(month: str, **options):
        return _run_month(
            tmp_path,
            tracker,
            config_dir,
            month,
            ["NOODLE BAR"],
            commit=True,
            local_llm_suggestions=True,
            suggester=ConsensusSuggester(FakeSuggester(votes, "a"), FakeSuggester(votes, "b")),
            **options,
        )

    april = commit_month("Apr")
    _fill_audit_corrections(april.review_xlsx_path, {"NOODLE BAR": "NONE"})
    commit_month("Apr", review_decisions_path=april.review_xlsx_path)
    commit_month("May")

    probe = _run_month(tmp_path, tracker, config_dir, "May", ["NOODLE BAR"])
    assert probe.categorized_transactions[0].categorization_method != "category_memory"


def test_monthly_commit_does_not_perform_currency_label_cleanup(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    tracker = tmp_path / "tracker.xlsx"
    config_dir = tmp_path / "config"
    _create_tracker(tracker)
    _create_config(config_dir, trust_policy=RAISED_TRUST_POLICY)
    _set_workbook_label(tracker, "A1", "Tracker currency: EUR")

    result = run_pipeline(
        tracker_path=tracker,
        statement_path=FIXTURE,
        config_dir=config_dir,
        year=2026,
        month="Apr",
        output_dir=tmp_path / "reports",
        commit=True,
    )

    original = load_workbook(tracker)
    copied = load_workbook(result.output_workbook_path)
    try:
        assert original["Net worth"]["A1"].value == "Tracker currency: EUR"
        assert copied["Net worth"]["A1"].value == "Tracker currency: EUR"
    finally:
        original.close()
        copied.close()

    report = result.report_path.read_text(encoding="utf-8")
    assert "Workbook cleanup tasks: not run during monthly update." in report


def test_pipeline_review_csv_includes_skipped_derived_workbook_rows(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    tracker = tmp_path / "tracker.xlsx"
    config_dir = tmp_path / "config"
    _create_tracker(tracker)
    _create_config(config_dir, salary_category="Income (net)", trust_policy=RAISED_TRUST_POLICY)

    result = run_pipeline(
        tracker_path=tracker,
        statement_path=FIXTURE,
        config_dir=config_dir,
        year=2026,
        month="Apr",
        output_dir=tmp_path / "reports",
    )

    review_csv = result.review_csv_path.read_text(encoding="utf-8")
    assert "workbook_update" in review_csv
    assert "Income (net)" in review_csv
    assert "Derived workbook row is formula-owned and not writable." in review_csv


def test_pipeline_applies_explicit_monthly_review_decisions(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    tracker = tmp_path / "tracker.xlsx"
    config_dir = tmp_path / "config"
    _create_tracker(tracker)
    _create_config(config_dir)
    first_result = run_pipeline(
        tracker_path=tracker,
        statement_path=FIXTURE,
        config_dir=config_dir,
        year=2026,
        month="Apr",
        output_dir=tmp_path / "first_reports",
    )
    reviewed_transaction = next(
        item
        for item in first_result.categorized_transactions
        if item.suggested_category == "Traveling"
    )
    review_decisions = tmp_path / "review_required_2026_apr.xlsx"
    _write_review_decisions(
        review_decisions,
        year=2026,
        month="Apr",
        rows=[
            {
                "transaction_id": reviewed_transaction.transaction.transaction_id,
                "description": reviewed_transaction.transaction.description,
                "amount": str(reviewed_transaction.transaction.amount),
                "manual_category": "Apple Cloud",
            },
            {
                "transaction_id": first_result.categorized_transactions[0].transaction.transaction_id,
                "description": "ignored because blank manual category",
                "amount": "0",
                "manual_category": "",
            },
        ],
    )

    result = run_pipeline(
        tracker_path=tracker,
        statement_path=FIXTURE,
        config_dir=config_dir,
        year=2026,
        month="Apr",
        output_dir=tmp_path / "reports",
        review_decisions_path=review_decisions,
    )

    reviewed = next(
        item
        for item in result.categorized_transactions
        if item.transaction.transaction_id == reviewed_transaction.transaction.transaction_id
    )
    assert reviewed.suggested_category == "Apple Cloud"
    assert reviewed.categorization_method == "monthly_review_decision"
    assert reviewed.review_required is False
    assert {update.category for update in result.updates} == {"Apple Cloud", "Full-time job (net)"}
    assert "monthly_review_decision" in result.categorized_csv_path.read_text(encoding="utf-8")
    assert "monthly_review_decision" in result.audit_path.read_text(encoding="utf-8")
    assert "- monthly_review_decision: 1" in result.report_path.read_text(encoding="utf-8")


def test_pipeline_preserves_review_decisions_when_input_is_default_review_output(
    tmp_path, monkeypatch
):
    monkeypatch.chdir(tmp_path)
    tracker = tmp_path / "tracker.xlsx"
    config_dir = tmp_path / "config"
    output_dir = tmp_path / "reports"
    statement = tmp_path / "statement.csv"
    _create_tracker(tracker)
    _create_config(config_dir)
    _write(
        statement,
        "\ufeffBooking date;Amount;Balance;Currency;Name;Title;Sender;Recipient;Reconciled\n"
        "2026/04/01;-42,50;957,50;DKK;UNKNOWN SHOP;Card purchase;1111;2222;Yes\n",
    )
    first_result = run_pipeline(
        tracker_path=tracker,
        statement_path=statement,
        config_dir=config_dir,
        year=2026,
        month="Apr",
        output_dir=output_dir,
    )
    review_decisions = first_result.review_xlsx_path
    reviewed_transaction = first_result.categorized_transactions[0]
    workbook = load_workbook(review_decisions)
    try:
        review_sheet = workbook["Review Required"]
        headers = {cell.value: index for index, cell in enumerate(review_sheet[1], start=1)}
        for row in range(2, review_sheet.max_row + 1):
            if (
                review_sheet.cell(row=row, column=headers["transaction_id"]).value
                == reviewed_transaction.transaction.transaction_id
            ):
                review_sheet.cell(row=row, column=headers["manual_category"]).value = "Apple Cloud"
                break
        else:
            raise AssertionError("review row not found")
        workbook.save(review_decisions)
    finally:
        workbook.close()

    result = run_pipeline(
        tracker_path=tracker,
        statement_path=statement,
        config_dir=config_dir,
        year=2026,
        month="Apr",
        output_dir=output_dir,
        review_decisions_path=review_decisions,
    )

    preserved = load_workbook(review_decisions, data_only=True)
    try:
        review_sheet = preserved["Review Required"]
        headers = {cell.value: index for index, cell in enumerate(review_sheet[1], start=1)}
        manual_values = [
            review_sheet.cell(row=row, column=headers["manual_category"]).value
            for row in range(2, review_sheet.max_row + 1)
        ]
    finally:
        preserved.close()

    assert "Apple Cloud" in manual_values
    assert result.review_xlsx_path == output_dir / "review_required_2026_apr_after_decisions.xlsx"
    assert result.review_xlsx_path.exists()
    assert "- monthly_review_decision: 1" in result.report_path.read_text(encoding="utf-8")


def test_pipeline_rejects_review_decisions_for_wrong_period(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    tracker = tmp_path / "tracker.xlsx"
    config_dir = tmp_path / "config"
    _create_tracker(tracker)
    _create_config(config_dir)
    review_decisions = tmp_path / "review_required_2026_may.xlsx"
    _write_review_decisions(
        review_decisions,
        year=2026,
        month="May",
        rows=[],
    )

    with pytest.raises(ValueError, match="review decisions.*May 2026.*Apr 2026"):
        run_pipeline(
            tracker_path=tracker,
            statement_path=FIXTURE,
            config_dir=config_dir,
            year=2026,
            month="Apr",
            output_dir=tmp_path / "reports",
            review_decisions_path=review_decisions,
        )


def test_pipeline_rejects_review_decisions_for_unknown_transaction_id(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    tracker = tmp_path / "tracker.xlsx"
    config_dir = tmp_path / "config"
    _create_tracker(tracker)
    _create_config(config_dir)
    review_decisions = tmp_path / "review_required_2026_apr.xlsx"
    _write_review_decisions(
        review_decisions,
        year=2026,
        month="Apr",
        rows=[
            {
                "transaction_id": f"{TRANSACTION_ID_SCHEME_VERSION}:missing:001",
                "description": "stale row",
                "amount": "-1",
                "manual_category": "Apple Cloud",
            }
        ],
    )

    with pytest.raises(ValueError, match="not present in current statement"):
        run_pipeline(
            tracker_path=tracker,
            statement_path=FIXTURE,
            config_dir=config_dir,
            year=2026,
            month="Apr",
            output_dir=tmp_path / "reports",
            review_decisions_path=review_decisions,
        )


def test_monthly_review_decision_preserves_workbook_write_safety(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    tracker = tmp_path / "tracker.xlsx"
    config_dir = tmp_path / "config"
    _create_tracker(tracker)
    _create_config(config_dir)
    first_result = run_pipeline(
        tracker_path=tracker,
        statement_path=FIXTURE,
        config_dir=config_dir,
        year=2026,
        month="Apr",
        output_dir=tmp_path / "first_reports",
    )
    transaction_id = first_result.categorized_transactions[0].transaction.transaction_id
    review_decisions = tmp_path / "review_required_2026_apr.xlsx"
    _write_review_decisions(
        review_decisions,
        year=2026,
        month="Apr",
        rows=[
            {
                "transaction_id": transaction_id,
                "description": "manual derived row",
                "amount": "-1",
                "manual_category": "Income (net)",
            }
        ],
    )

    result = run_pipeline(
        tracker_path=tracker,
        statement_path=FIXTURE,
        config_dir=config_dir,
        year=2026,
        month="Apr",
        output_dir=tmp_path / "reports",
        review_decisions_path=review_decisions,
    )

    income_update = next(update for update in result.updates if update.category == "Income (net)")
    assert income_update.write_action == "skip"
    assert income_update.reason == "Derived workbook row is formula-owned and not writable."


def test_pipeline_rejects_manual_review_category_missing_from_current_tracker(
    tmp_path, monkeypatch
):
    monkeypatch.chdir(tmp_path)
    tracker = tmp_path / "tracker.xlsx"
    config_dir = tmp_path / "config"
    _create_tracker(tracker)
    _create_config(config_dir)
    first_result = run_pipeline(
        tracker_path=tracker,
        statement_path=FIXTURE,
        config_dir=config_dir,
        year=2026,
        month="Apr",
        output_dir=tmp_path / "first_reports",
    )
    transaction_id = first_result.categorized_transactions[0].transaction.transaction_id
    review_decisions = tmp_path / "review_required_2026_apr.xlsx"
    _write_review_decisions(
        review_decisions,
        year=2026,
        month="Apr",
        rows=[
            {
                "transaction_id": transaction_id,
                "description": "manual category that is not in the tracker",
                "amount": "-1",
                "manual_category": "Missing Tracker Category",
            }
        ],
    )

    with pytest.raises(ValueError, match="manual_category.*Missing Tracker Category"):
        run_pipeline(
            tracker_path=tracker,
            statement_path=FIXTURE,
            config_dir=config_dir,
            year=2026,
            month="Apr",
            output_dir=tmp_path / "reports",
            review_decisions_path=review_decisions,
        )


def test_pipeline_accepts_manual_review_category_registered_as_missing_leaf(
    tmp_path, monkeypatch
):
    monkeypatch.chdir(tmp_path)
    tracker = tmp_path / "tracker.xlsx"
    config_dir = tmp_path / "config"
    statement = tmp_path / "statement.csv"
    _create_tracker(tracker)
    _create_category_registry_config(config_dir)
    _write_unknown_shop_statement(statement)
    first_result = run_pipeline(
        tracker_path=tracker,
        statement_path=statement,
        config_dir=config_dir,
        year=2026,
        month="Apr",
        output_dir=tmp_path / "first_reports",
    )
    transaction = first_result.categorized_transactions[0].transaction
    config_path = config_dir / "categories.yaml"
    config_path.write_text(
        config_path.read_text(encoding="utf-8").replace(
            '      - "Traveling"\n',
            '      - "Traveling"\n      - "Pet Supplies"\n',
        ),
        encoding="utf-8",
    )
    review_decisions = tmp_path / "review_required_2026_apr.xlsx"
    _write_review_decisions(
        review_decisions,
        year=2026,
        month="Apr",
        rows=[
            {
                "transaction_id": transaction.transaction_id,
                "description": transaction.description,
                "amount": str(transaction.amount),
                "manual_category": "Pet Supplies",
            }
        ],
    )

    result = run_pipeline(
        tracker_path=tracker,
        statement_path=statement,
        config_dir=config_dir,
        year=2026,
        month="Apr",
        output_dir=tmp_path / "reports",
        review_decisions_path=review_decisions,
    )

    reviewed = result.categorized_transactions[0]
    assert reviewed.suggested_category == "Pet Supplies"
    assert reviewed.categorization_method == "monthly_review_decision"
    assert result.updates[0].category == "Pet Supplies"
    assert result.updates[0].reason == "Target category row not found."


def test_pipeline_local_llm_suggestion_is_review_only_and_reuses_review_fields(
    tmp_path, monkeypatch
):
    monkeypatch.chdir(tmp_path)
    tracker = tmp_path / "tracker.xlsx"
    config_dir = tmp_path / "config"
    statement = tmp_path / "statement.csv"
    _create_tracker(tracker)
    _create_category_registry_config(config_dir)
    _write_unknown_shop_statement(statement)
    suggester = FakeSuggester(
        {"UNKNOWN SHOP": ScriptedVote("Traveling", 0.68, "Merchant looks travel related.")}
    )

    result = run_pipeline(
        tracker_path=tracker,
        statement_path=statement,
        config_dir=config_dir,
        year=2026,
        month="Apr",
        output_dir=tmp_path / "reports",
        local_llm_suggestions=True,
        suggester=suggester,
    )

    item = result.categorized_transactions[0]
    assert item.suggested_category == "Traveling"
    assert item.categorization_method == "local_llm_gemma"
    assert item.review_required is True
    assert result.review_count == 1

    report = result.report_path.read_text(encoding="utf-8")
    assert "## Local LLM Mode" in report
    assert "- Status: enabled" in report
    assert "- Second model: gemma4:12b" in report
    assert "- Eligible rows: 1" in report
    assert "- Provider calls attempted: 1" in report
    assert "- Existing-leaf suggestions: 1" in report
    assert "- local_llm_gemma: 1" in report

    audit = result.audit_path.read_text(encoding="utf-8")
    assert '"record_type": "local_llm_summary"' in audit
    assert '"existing_leaf_suggestions": 1' in audit
    assert '"second_model": "gemma4:12b"' in audit
    assert '"categorization_method": "local_llm_gemma"' in audit

    workbook = load_workbook(result.review_xlsx_path, data_only=True)
    try:
        review_sheet = workbook["Review Required"]
        headers = {cell.value: index for index, cell in enumerate(review_sheet[1], start=1)}
        assert "llm_suggested_category" not in headers
        assert "llm_reason" not in headers
        assert review_sheet.cell(row=2, column=headers["suggested_category"]).value == "Traveling"
        reason = review_sheet.cell(row=2, column=headers["reason"]).value
        assert "Merchant looks travel related." in reason
        assert reason.endswith("(local_llm_gemma)")
    finally:
        workbook.close()


def test_pipeline_includes_reviewed_policy_in_local_llm_prompt(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    tracker = tmp_path / "tracker.xlsx"
    config_dir = tmp_path / "config"
    statement = tmp_path / "statement.csv"
    memory_dir = tmp_path / "data" / "category_memory"
    memory_dir.mkdir(parents=True)
    (memory_dir / "reviewed_policy.local.md").write_text(
        "\n".join(
            [
                "# Reviewed Policy",
                "",
                "<!-- AUTO-GENERATED REVIEWED EXAMPLES START -->",
                "- Merchant identity: `MOBILEPAY REJSEKORT` -> `Traveling`",
                "<!-- AUTO-GENERATED REVIEWED EXAMPLES END -->",
                "",
                "## Manual Guidance",
                "",
                "- Prefer no_suggestion for unknown private transfers.",
                "",
            ]
        ),
        encoding="utf-8",
    )
    _create_tracker(tracker)
    _create_category_registry_config(config_dir)
    _write_unknown_shop_statement(statement)
    suggester = FakeSuggester({})

    run_pipeline(
        tracker_path=tracker,
        statement_path=statement,
        config_dir=config_dir,
        year=2026,
        month="Apr",
        output_dir=tmp_path / "reports",
        category_memory_dir=memory_dir,
        local_llm_suggestions=True,
        suggester=suggester,
    )

    _, context = suggester.calls[0]
    assert "MOBILEPAY REJSEKORT" in context.reviewed_policy
    assert "unknown private transfers" in context.reviewed_policy
    assert set(context.leaf_glossary) == {"Apple Cloud", "Traveling", "Full-time job (net)"}


def test_pipeline_passes_guidance_aliases_to_suggester_context(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    tracker = tmp_path / "tracker.xlsx"
    config_dir = tmp_path / "config"
    statement = tmp_path / "statement.csv"
    _create_tracker(tracker)
    _create_category_registry_config(config_dir)
    _write(config_dir / "guidance_aliases.local.yaml", '"TRAIN EXAMPLE": "Traveling"\n')
    _write_unknown_shop_statement(statement)
    suggester = FakeSuggester({})

    run_pipeline(
        tracker_path=tracker,
        statement_path=statement,
        config_dir=config_dir,
        year=2026,
        month="Apr",
        output_dir=tmp_path / "reports",
        local_llm_suggestions=True,
        suggester=suggester,
    )

    _, context = suggester.calls[0]
    assert context.guidance_aliases == {"TRAIN EXAMPLE": "Traveling"}


def test_pipeline_dry_run_plans_new_leaf_category_without_writing_registry(
    tmp_path, monkeypatch
):
    monkeypatch.chdir(tmp_path)
    tracker = tmp_path / "tracker.xlsx"
    config_dir = tmp_path / "config"
    statement = tmp_path / "statement.csv"
    _create_tracker(tracker)
    _create_category_registry_config(config_dir)
    _write_unknown_shop_statement(statement)

    first_result = run_pipeline(
        tracker_path=tracker,
        statement_path=statement,
        config_dir=config_dir,
        year=2026,
        month="Apr",
        output_dir=tmp_path / "first_reports",
    )
    transaction = first_result.categorized_transactions[0].transaction
    review_decisions = tmp_path / "review_required_2026_apr.xlsx"
    _write_review_decisions(
        review_decisions,
        year=2026,
        month="Apr",
        include_new_category_columns=True,
        rows=[
            {
                "transaction_id": transaction.transaction_id,
                "description": transaction.description,
                "amount": str(transaction.amount),
                "manual_category": "",
                "new_parent_category": "Living expenses",
                "new_leaf_category": "Pet Supplies",
            }
        ],
    )

    registry_before = (config_dir / "categories.yaml").read_bytes()

    result = run_pipeline(
        tracker_path=tracker,
        statement_path=statement,
        config_dir=config_dir,
        year=2026,
        month="Apr",
        output_dir=tmp_path / "reports",
        review_decisions_path=review_decisions,
    )

    reviewed = result.categorized_transactions[0]
    assert reviewed.suggested_category == "Pet Supplies"
    assert reviewed.categorization_method == "monthly_review_decision"
    assert reviewed.review_required is False
    assert result.output_workbook_path is None
    update = result.updates[0]
    assert update.category == "Pet Supplies"
    assert update.write_action == "review"
    assert update.reason == "Target category row not found."

    assert (config_dir / "categories.yaml").read_bytes() == registry_before
    assert result.category_registry_additions[0].description == (
        "Added in monthly review Apr 2026."
    )
    report = result.report_path.read_text(encoding="utf-8")
    assert "## Category Registry Updates" in report
    assert "Pet Supplies under Living expenses" in report
    audit = result.audit_path.read_text(encoding="utf-8")
    assert '"record_type": "category_registry_addition"' in audit
    assert '"leaf_category": "Pet Supplies"' in audit


def test_pipeline_rejects_new_leaf_category_duplicate(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    tracker = tmp_path / "tracker.xlsx"
    config_dir = tmp_path / "config"
    statement = tmp_path / "statement.csv"
    _create_tracker(tracker)
    _create_category_registry_config(config_dir)
    _write_unknown_shop_statement(statement)
    first_result = run_pipeline(
        tracker_path=tracker,
        statement_path=statement,
        config_dir=config_dir,
        year=2026,
        month="Apr",
        output_dir=tmp_path / "first_reports",
    )
    review_decisions = tmp_path / "review_required_2026_apr.xlsx"
    _write_review_decisions(
        review_decisions,
        year=2026,
        month="Apr",
        include_new_category_columns=True,
        rows=[
            {
                "transaction_id": first_result.transactions[0].transaction_id,
                "description": "duplicate leaf",
                "amount": "-1",
                "manual_category": "",
                "new_parent_category": "Living expenses",
                "new_leaf_category": " apple cloud ",
            }
        ],
    )

    with pytest.raises(ValueError, match="Duplicate category label.*Apple Cloud.*apple cloud"):
        run_pipeline(
            tracker_path=tracker,
            statement_path=statement,
            config_dir=config_dir,
            year=2026,
            month="Apr",
            output_dir=tmp_path / "reports",
            review_decisions_path=review_decisions,
        )


def test_pipeline_rejects_new_leaf_category_for_invalid_parent(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    tracker = tmp_path / "tracker.xlsx"
    config_dir = tmp_path / "config"
    statement = tmp_path / "statement.csv"
    _create_tracker(tracker)
    _create_category_registry_config(config_dir)
    _write_unknown_shop_statement(statement)
    first_result = run_pipeline(
        tracker_path=tracker,
        statement_path=statement,
        config_dir=config_dir,
        year=2026,
        month="Apr",
        output_dir=tmp_path / "first_reports",
    )
    review_decisions = tmp_path / "review_required_2026_apr.xlsx"
    _write_review_decisions(
        review_decisions,
        year=2026,
        month="Apr",
        include_new_category_columns=True,
        rows=[
            {
                "transaction_id": first_result.transactions[0].transaction_id,
                "description": "invalid parent",
                "amount": "-1",
                "manual_category": "",
                "new_parent_category": "Income (net)",
                "new_leaf_category": "Pet Supplies",
            }
        ],
    )

    with pytest.raises(ValueError, match="new_parent_category.*Income \\(net\\).*not allowed"):
        run_pipeline(
            tracker_path=tracker,
            statement_path=statement,
            config_dir=config_dir,
            year=2026,
            month="Apr",
            output_dir=tmp_path / "reports",
            review_decisions_path=review_decisions,
        )


def test_pipeline_rejects_review_row_with_manual_and_new_leaf_category(
    tmp_path, monkeypatch
):
    monkeypatch.chdir(tmp_path)
    tracker = tmp_path / "tracker.xlsx"
    config_dir = tmp_path / "config"
    statement = tmp_path / "statement.csv"
    _create_tracker(tracker)
    _create_category_registry_config(config_dir)
    _write_unknown_shop_statement(statement)
    first_result = run_pipeline(
        tracker_path=tracker,
        statement_path=statement,
        config_dir=config_dir,
        year=2026,
        month="Apr",
        output_dir=tmp_path / "first_reports",
    )
    review_decisions = tmp_path / "review_required_2026_apr.xlsx"
    _write_review_decisions(
        review_decisions,
        year=2026,
        month="Apr",
        include_new_category_columns=True,
        rows=[
            {
                "transaction_id": first_result.transactions[0].transaction_id,
                "description": "conflicting row",
                "amount": "-1",
                "manual_category": "Apple Cloud",
                "new_parent_category": "Living expenses",
                "new_leaf_category": "Pet Supplies",
            }
        ],
    )

    with pytest.raises(ValueError, match="manual_category.*new_leaf_category"):
        run_pipeline(
            tracker_path=tracker,
            statement_path=statement,
            config_dir=config_dir,
            year=2026,
            month="Apr",
            output_dir=tmp_path / "reports",
            review_decisions_path=review_decisions,
        )


def test_pipeline_rejects_review_decisions_with_unsupported_transaction_id_scheme(
    tmp_path, monkeypatch
):
    monkeypatch.chdir(tmp_path)
    tracker = tmp_path / "tracker.xlsx"
    config_dir = tmp_path / "config"
    _create_tracker(tracker)
    _create_config(config_dir)
    review_decisions = tmp_path / "review_required_2026_apr.xlsx"
    _write_review_decisions(
        review_decisions,
        year=2026,
        month="Apr",
        transaction_id_scheme="legacy-v0",
        rows=[],
    )

    with pytest.raises(ValueError, match="Unsupported review decision transaction ID scheme"):
        run_pipeline(
            tracker_path=tracker,
            statement_path=FIXTURE,
            config_dir=config_dir,
            year=2026,
            month="Apr",
            output_dir=tmp_path / "reports",
            review_decisions_path=review_decisions,
        )


def _create_proxy_split_tracker(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Net worth"
    sheet["C2"] = 2026
    sheet["C3"] = "Apr"
    sheet["B5"] = "Parent A"
    sheet["B6"] = "Parent B"
    sheet["B7"] = "Traveling"
    workbook.save(path)
    workbook.close()


def _create_proxy_split_config(config_dir: Path, monthly_limit: int = 1) -> None:
    config_dir.mkdir(parents=True)
    _write(
        config_dir / "settings.yaml",
        """
tracker:
  sheet_name: "Net worth"
  currency: "DKK"
  category_column: 2
  year_header_row: 2
  month_header_row: 3
statement:
  currency: "DKK"
confidence_thresholds:
  auto_write: 0.85
  review_required: 0.60
  reject_below: 0.60
writer:
  overwrite_fixed_rows: false
  highlight_auto_filled_cells: false
""",
    )
    _write(
        config_dir / "categories.yaml",
        """
category_registry:
  - label: "Living expenses"
    type: "parent"
    allow_new_children: true
    children:
      - "Parent A"
      - "Parent B"
      - "Traveling"
aliases: {}
""",
    )
    _write(config_dir / "rules.yaml", "historical_mappings: {}\nrules: []\nfixed_rows: []\n")
    _write(
        config_dir / "rules.local.yaml",
        f"""
proxy_split_rules:
  - name: example_transfer
    match_keywords: [revolut]
    direction: expense
    conversion_rate: "0.82"
    monthly_limit: {monthly_limit}
    allocations:
      - role: parent_a
        category: Parent A
        base_amount: "8000"
      - role: parent_b
        category: Parent B
        base_amount: "4000"
""",
    )


def _write_revolut_statement(path: Path, amount: str) -> None:
    _write(
        path,
        "\ufeffBooking date;Amount;Balance;Currency;Name;Title;Sender;Recipient;Reconciled\n"
        f"2026/04/04;{amount};50000,00;DKK;REVOLUT;Transfer;1111;2222;Yes\n",
    )


def _create_tracker(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Net worth"
    sheet["C2"] = 2026
    sheet["C3"] = "Apr"
    sheet["B5"] = "Apple Cloud"
    sheet["B6"] = "Full-time job (net)"
    sheet["B7"] = "Traveling"
    sheet["B8"] = "Income (net)"
    workbook.save(path)
    workbook.close()


RAISED_TRUST_POLICY = """
trust_policy:
  auto_max_amount: 20000
  never_auto_categories: []
"""


def _create_config(
    config_dir: Path,
    salary_category: str = "Full-time job (net)",
    trust_policy: str = "",
) -> None:
    config_dir.mkdir(parents=True)
    _write(
        config_dir / "settings.yaml",
        """
tracker:
  sheet_name: "Net worth"
  currency: "DKK"
  category_column: 2
  year_header_row: 2
  month_header_row: 3
statement:
  currency: "DKK"
confidence_thresholds:
  auto_write: 0.85
  review_required: 0.60
  reject_below: 0.60
writer:
  overwrite_fixed_rows: false
  highlight_auto_filled_cells: false
"""
        + trust_policy,
    )
    _write(
        config_dir / "categories.yaml",
        f"""
categories:
  - "Apple Cloud"
  - "{salary_category}"
  - "Traveling"
aliases: {{}}
""",
    )
    _write(
        config_dir / "rules.yaml",
        f"""
historical_mappings:
  "APPLE.COM/BILL": "Apple Cloud"
rules:
  - category: "{salary_category}"
    match_keywords: ["salary"]
    direction: "income"
    confidence: 0.95
  - category: "Traveling"
    match_keywords: ["foreign card"]
    direction: "expense"
    confidence: 0.95
fixed_rows: []
""",
    )


def _create_category_registry_config(config_dir: Path) -> None:
    config_dir.mkdir(parents=True)
    _write(
        config_dir / "settings.yaml",
        """
tracker:
  sheet_name: "Net worth"
  currency: "DKK"
  category_column: 2
  year_header_row: 2
  month_header_row: 3
statement:
  currency: "DKK"
confidence_thresholds:
  auto_write: 0.85
  review_required: 0.60
  reject_below: 0.60
writer:
  overwrite_fixed_rows: false
  highlight_auto_filled_cells: false
""",
    )
    _write(
        config_dir / "categories.yaml",
        """
category_registry:
  - label: "Living expenses"
    type: "parent"
    allow_new_children: true
    children:
      - "Apple Cloud"
      - "Traveling"
  - label: "Income (net)"
    type: "derived"
  - "Full-time job (net)"
aliases: {}
""",
    )
    _write(
        config_dir / "rules.yaml",
        """
historical_mappings: {}
rules: []
fixed_rows: []
""",
    )


def _write_unknown_shop_statement(path: Path) -> None:
    _write(
        path,
        "\ufeffBooking date;Amount;Balance;Currency;Name;Title;Sender;Recipient;Reconciled\n"
        "2026/04/01;-42,50;957,50;DKK;UNKNOWN SHOP;Card purchase;1111;2222;Yes\n",
    )


def _sheet_rows(sheet) -> list[dict[str, object]]:
    headers = [cell.value for cell in sheet[1]]
    return [
        dict(zip(headers, (cell.value for cell in row)))
        for row in sheet.iter_rows(min_row=2)
    ]


def _dropdown_options(workbook, cell: str) -> list[object]:
    (validation,) = [
        validation
        for validation in workbook["Review Required"].data_validations.dataValidation
        if cell in validation.sqref
    ]
    sheet_name, cells = validation.formula1.rsplit("!", 1)
    sheet = workbook[sheet_name.strip("'")]
    return [row[0].value for row in sheet[cells.replace("$", "")]]


def _write(path: Path, content: str) -> None:
    path.write_text(content.lstrip(), encoding="utf-8")


def _set_workbook_label(path: Path, cell: str, value: str) -> None:
    workbook = load_workbook(path)
    try:
        workbook["Net worth"][cell] = value
        workbook.save(path)
    finally:
        workbook.close()


def _write_review_decisions(
    path: Path,
    year: int,
    month: str,
    rows: list[dict[str, str]],
    transaction_id_scheme: str = TRANSACTION_ID_SCHEME_VERSION,
    include_new_category_columns: bool = False,
) -> None:
    workbook = Workbook()
    review_sheet = workbook.active
    review_sheet.title = "Review Required"
    headers = [
        "transaction_id",
        "date",
        "description",
        "amount",
        "direction",
        "merchant_identity",
        "suggested_category",
        "confidence",
        "method",
        "reason",
        "manual_category",
    ]
    if include_new_category_columns:
        headers.extend(["new_parent_category", "new_leaf_category"])
    headers.append("learn_to_memory")
    review_sheet.append(headers)
    for row in rows:
        payload = {
            "transaction_id": row["transaction_id"],
            "date": "2026-04-01",
            "description": row["description"],
            "amount": row["amount"],
            "direction": "expense",
            "merchant_identity": "",
            "suggested_category": "",
            "confidence": "",
            "method": "",
            "reason": "",
            "manual_category": row["manual_category"],
            "new_parent_category": row.get("new_parent_category", ""),
            "new_leaf_category": row.get("new_leaf_category", ""),
            "learn_to_memory": "",
        }
        review_sheet.append([payload[header] for header in headers])
    metadata_sheet = workbook.create_sheet("Run Metadata")
    for key, value in [
        ("reporting_year", year),
        ("reporting_month", month),
        ("statement_parser", "nordea-pdf"),
        ("generated_timestamp", "2026-05-16T10:00:00"),
        ("transaction_id_scheme", transaction_id_scheme),
    ]:
        metadata_sheet.append([key, value])
    workbook.save(path)
    workbook.close()


MONTH_NUMBER = {"Apr": 4, "May": 5, "Jun": 6}


def _create_multi_month_tracker(path: Path, months: tuple[str, ...]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Net worth"
    for column, month in enumerate(months, start=3):
        sheet.cell(row=2, column=column, value=2026)
        sheet.cell(row=3, column=column, value=month)
    sheet["B5"] = "Apple Cloud"
    sheet["B6"] = "Full-time job (net)"
    sheet["B7"] = "Traveling"
    sheet["B8"] = "Income (net)"
    workbook.save(path)
    workbook.close()


def _create_services_tracker(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Net worth"
    for column, month in enumerate(("Apr", "May"), start=3):
        sheet.cell(row=2, column=column, value=2026)
        sheet.cell(row=3, column=column, value=month)
    sheet["B5"] = "Services"
    sheet["B6"] = "Disney+"
    sheet["B7"] = "Insurance"
    workbook.save(path)
    workbook.close()


def _create_services_config(config_dir: Path) -> None:
    _create_category_registry_config(config_dir)
    _write(
        config_dir / "categories.yaml",
        """
category_registry:
  - label: "Services"
    type: "parent"
    allow_new_children: true
    children:
      - "Disney+"
  - label: "Insurance"
    type: "parent"
    children: []
aliases: {}
""",
    )


def _added_lines(before: str, after: str) -> list[str]:
    """Lines `after` adds to `before`; fails if any line of `before` changed or moved."""
    diff = list(difflib.ndiff(before.splitlines(), after.splitlines()))
    assert [line for line in diff if line.startswith("- ")] == []
    return [line[2:] for line in diff if line.startswith("+ ")]


def _run_month(
    tmp_path: Path,
    tracker: Path,
    config_dir: Path,
    month: str,
    merchants: list[str],
    **options,
):
    statement = tmp_path / f"statement_{month}.csv"
    _write(
        statement,
        "\ufeffBooking date;Amount;Balance;Currency;Name;Title;Sender;Recipient;Reconciled\n"
        + "".join(
            f"2026/{MONTH_NUMBER[month]:02d}/{day:02d};-40,00;900,00;DKK;{merchant};"
            "Card purchase;1111;2222;Yes\n"
            for day, merchant in enumerate(merchants, start=1)
        ),
    )
    return run_pipeline(
        tracker_path=tracker,
        statement_path=statement,
        config_dir=config_dir,
        year=2026,
        month=month,
        output_dir=tmp_path / "reports",
        **options,
    )


def _fill_exception_sheet(
    path: Path, decisions: dict[str, str | None], column: str = "manual_category"
) -> None:
    workbook = load_workbook(path)
    try:
        sheet = workbook["Review Required"]
        headers = {cell.value: index for index, cell in enumerate(sheet[1], start=1)}
        for row in range(2, sheet.max_row + 1):
            description = sheet.cell(row=row, column=headers["description"]).value
            for merchant, decision in decisions.items():
                if merchant in description:
                    sheet.cell(row=row, column=headers[column]).value = decision
        workbook.save(path)
    finally:
        workbook.close()


def _fill_audit_corrections(path: Path, corrections: dict[str, str]) -> None:
    workbook = load_workbook(path)
    try:
        sheet = workbook["Audit"]
        headers = {cell.value: index for index, cell in enumerate(sheet[1], start=1)}
        for row in range(2, sheet.max_row + 1):
            description = sheet.cell(row=row, column=headers["description"]).value
            for merchant, correction in corrections.items():
                if merchant in description:
                    sheet.cell(row=row, column=headers["corrected_category"]).value = correction
        workbook.save(path)
    finally:
        workbook.close()
